"""Compare synthetic first-download workbooks with a saved pre-change backend.

Save legal_platform.py, analytics.py, main.py, indicator_reporting.py and
indicator_reconciliation.py in --baseline-dir before changing the implementation.
No source case files or exported responses are cached by this benchmark.
"""
from __future__ import annotations

import argparse
import importlib
import io
import json
import os
import statistics
import sys
import types
from copy import copy
from pathlib import Path
from time import perf_counter
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["INTERSOS_DEFER_LEGAL_LOAD"] = "1"
from scripts.benchmark_indicator_performance import synthetic_frames


def modules(prefix):
    return {name: importlib.import_module(f"{prefix}.{name}") for name in
            ("legal_platform", "analytics", "main", "indicator_reporting", "indicator_reconciliation")}


def load_baseline(directory):
    name = "backend.excel_before"
    package = types.ModuleType(name)
    package.__path__ = [str(directory), str(ROOT / "backend")]
    sys.modules[name] = package
    return modules(name)


def fixtures(implementation, count):
    legal, analytics, main, reporting, checking = [implementation[name] for name in
        ("legal_platform", "analytics", "main", "indicator_reporting", "indicator_reconciliation")]
    frames = synthetic_frames(count)
    frames["beneficiaries"]["Date of Identification"] = "01/10/2026"
    frames["beneficiaries"]["Created On"] = "2026-10-01"
    frames["beneficiaries"]["Notes"] = "Synthetic notes only"
    frames["legalservices"]["Type of Document"] = ""
    frames["legalhotlines"] = pd.DataFrame({"Hotline ID": [f"{index:06}" for index in range(count)],
        "Contact Date": ["10/01/2026"] * count, "Follow-up Date": ["2026-10-02"] * count})
    store = legal.LegalStore(frames, "synthetic", [])
    flags = []
    for index, row in frames["legalservices"].iterrows():
        store._flag(flags, "legalservices", "Missing Type of Document", "Medium", index, row, "Missing document")
    store.flags["legalservices"] = flags
    raw = object.__new__(analytics.DataStore)
    raw.raw_sheets = {"synthetic": frames["beneficiaries"]}
    _, prepared = store.indicator_preparation()
    table_rows = [{"Date": "2026-10-01", "Count": index, "Identifier": f"{index:06}", "Notes": "Synthetic"}
                  for index in range(count)]
    exclusions = [{"identifierValue": f"{index:06}", "excludedAt": "2026-10-01T12:30:00+03:00"} for index in range(count)]
    comparison = {"month": "2026-10", "project": "All", "filename": "synthetic.xlsx", "rows": [
        {"note": "Date differs", "lawyer": "Synthetic", "beneficiaryId": f"B{index}", "name": "Synthetic",
         "differences": [{"field": "Date", "assessment": "01/10/2026", "excel": "02/10/2026"}]} for index in range(count)]}
    store.detention_reconciliation = lambda *args: comparison
    check = {"metadata": {"filename": "synthetic.xlsx", "sheet": "Reporting Tool"}, "months": ["2026-10"],
             "summary": {}, "warnings": [], "rows": [{"indicator": "Synthetic", "month": "2026-10",
             "project": "UNHCR 2026 - Gov", "location": "Anbar", "population": "idp", "sex": "male",
             "ageGroup": "18-39", "platformValue": index, "workbookValue": index, "variance": 0,
             "status": "matched"} for index in range(count)]}
    def table(style="default"):
        return main.table_workbook(main.TableWorkbookRequest(columns=list(table_rows[0]), rows=table_rows, style=style)).body
    def exclusion_export():
        with patch.object(main, "duplicate_exclusions", types.SimpleNamespace(entries=lambda: exclusions)):
            return main.export_duplicate_exclusions().body
    def indicator(narrative=False):
        report = reporting.build_indicator_report(frames, months=["2026-01", "2026-02", "2026-03"], prepared=prepared)
        return reporting.build_narrative_workbook(report) if narrative else reporting.build_indicator_workbook(report)
    return {
        "raw_dataset": lambda: store.export("beneficiaries"),
        "legal_explorer": lambda: store.explorer_export("beneficiaries"),
        "hotline_explorer": lambda: store.explorer_export("legalhotlines"),
        "workbook_explorer": lambda: raw.explorer_export("synthetic", "", [], None, "asc", [], "xlsx"),
        "review_findings": lambda: store.review_export("legalservices", selected_rules=["Missing Type of Document"]),
        "bulk_cases": lambda: store.case_export("", {}),
        "selected_cases": lambda: store.case_export("", {}, ["B0"]),
        "indicators": indicator,
        "narrative": lambda: indicator(True),
        "reporting_check_writer": lambda: checking.build_reconciliation_workbook(check),
        "detention_comparison_writer": lambda: store.detention_reconciliation_export(b"", "synthetic.xlsx", "2026-10"),
        "table": table,
        "pivot": lambda: table("interactive-detail"),
        "exclusions": exclusion_export,
    }


def assert_workbooks_equal(before, after):
    left, right = load_workbook(io.BytesIO(before)), load_workbook(io.BytesIO(after))
    assert left.sheetnames == right.sheetnames
    checked_styles = set()
    for old, new in zip(left, right):
        assert old.max_row == new.max_row and old.max_column == new.max_column, old.title
        assert str(old.merged_cells) == str(new.merged_cells), old.title
        assert old.freeze_panes == new.freeze_panes and old.auto_filter == new.auto_filter, old.title
        assert old.sheet_view == new.sheet_view, old.title
        assert dict(old.tables) == dict(new.tables), old.title
        assert {key: value.height for key, value in old.row_dimensions.items()} == {key: value.height for key, value in new.row_dimensions.items()}, old.title
        assert {key: value.width for key, value in old.column_dimensions.items()} == {key: value.width for key, value in new.column_dimensions.items()}, old.title
        assert len(old._charts) == len(new._charts), old.title
        for old_row, new_row in zip(old, new):
            for a, b in zip(old_row, new_row):
                location = f"{old.title}!{a.coordinate}"
                assert a.value == b.value, (location, a.value, b.value)
                assert a.number_format == b.number_format, (location, a.number_format, b.number_format)
                style_pair = (a.style_id, b.style_id)
                if style_pair not in checked_styles:
                    for attribute in ("font", "fill", "border", "alignment", "protection"):
                        assert copy(getattr(a, attribute)) == copy(getattr(b, attribute)), (location, attribute)
                    checked_styles.add(style_pair)
                if a.value is not None:
                    assert a.data_type == b.data_type, location


def measure(implementation, run):
    from backend import excel_export
    for name in ("Alignment", "Border", "Font", "PatternFill", "Side"):
        getattr(excel_export, name).cache_clear()
    stages = {"calculation": 0.0, "date_formatting": 0.0, "serialization": 0.0}
    def timed(function, stage):
        def wrapped(*args, **kwargs):
            start = perf_counter()
            try: return function(*args, **kwargs)
            finally: stages[stage] += perf_counter() - start
        return wrapped
    from contextlib import ExitStack
    with ExitStack() as stack:
        original_dates = implementation["legal_platform"].format_excel_dates
        for module in implementation.values():
            stack.enter_context(patch.object(module, "format_excel_dates", timed(original_dates, "date_formatting")))
        stack.enter_context(patch.object(Workbook, "save", timed(Workbook.save, "serialization")))
        reporting = implementation["indicator_reporting"]
        stack.enter_context(patch.object(reporting, "build_indicator_report", timed(reporting.build_indicator_report, "calculation")))
        start = perf_counter()
        payload = run()
        total = perf_counter() - start
    stages["workbook_and_filtering"] = max(0, total - sum(stages.values()))
    return payload, {"total": total, **stages}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--baseline-dir", type=Path, required=True)
    parser.add_argument("--sizes", type=int, nargs="+", default=[100, 3000])
    parser.add_argument("--repeats", type=int, default=3)
    parser.add_argument("--families", nargs="+", help="Only measure these export families")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    before, after = load_baseline(args.baseline_dir), modules("backend")
    results = []
    for size in args.sizes:
        original, optimized = fixtures(before, size), fixtures(after, size)
        for name in original:
            if args.families and name not in args.families:
                continue
            old_times, new_times = [], []
            for _ in range(args.repeats):
                old, old_stages = measure(before, original[name])
                new, new_stages = measure(after, optimized[name])
                assert_workbooks_equal(old, new)
                old_times.append(old_stages); new_times.append(new_stages)
            median = lambda samples: {key: round(statistics.median(item[key] for item in samples), 4) for key in samples[0]}
            old, new = median(old_times), median(new_times)
            result = {"family": name, "source_rows": size, "baseline_seconds": old, "optimized_seconds": new,
                      "reduction_percent": round((1-new["total"]/old["total"])*100, 1), "parity": True}
            print(json.dumps(result), flush=True)
            results.append(result)
    payload = {"synthetic_only": True, "repeats": args.repeats,
               "notes": "No response caches. Style caches are cleared before every download. Reporting Check and detention writer inputs are prepared synthetic results; their comparison calculations are not timed. Hotline parsing is included in workbook_and_filtering. Stage medians are independent and may not sum to the median total.",
               "results": results}
    if args.output: args.output.write_text(json.dumps(payload, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
