import copy
import io
from datetime import datetime

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.custom_analysis import make_custom_router, parse_date, extract_time, infer, read_tables, Configuration, initial_columns, table_data, build
from backend.kobo import ProjectStore
from backend.kobo_api import make_router


@pytest.fixture
def env(tmp_path):
    store = ProjectStore(tmp_path / 'custom.sqlite3')
    app = FastAPI()
    app.include_router(make_custom_router(store))
    shared = make_router(store, prefix='/api/custom')
    shared.routes = [r for r in shared.routes if '/imports/' not in r.path and not r.path.endswith('/review')]
    app.include_router(shared)
    app.include_router(make_router(ProjectStore(tmp_path / 'kobo.sqlite3')))
    return TestClient(app), store


def stage(client, raw=b'Code,Age,Date,City\n001,20,04/05/2026,Mosul\n002,30,15/05/2026,Baghdad\n', filename='data.csv', pid=''):
    r = client.post('/api/custom/columns/preview', files={'data': (filename, raw)}, data={'projectId': pid})
    assert r.status_code == 200, r.text
    return r.json()


def confirm(client, p, expected=200):
    r = client.post('/api/custom/columns/confirm', json={'token': p['token'], 'name': 'Custom dataset', 'config': p['config']})
    assert r.status_code == expected, r.text
    return r.json()


def test_dates():
    assert parse_date('04/05/2026', 'uk') == '2026-05-04'
    assert parse_date('04/05/2026', 'us') == '2026-04-05'
    assert parse_date('29/02/2024 13:30:25', 'uk') == '2024-02-29'
    assert parse_date('2026-10-25T13:45:00Z') == '2026-10-25'
    assert parse_date(datetime(2026, 10, 25)) == '2026-10-25'
    for value in ['04/05/2026', '29/02/2025', '31/04/2026', '15/01/2026 99:20', '2026-10-25T99:20:00Z']:
        with pytest.raises(ValueError): parse_date(value)


def test_inference():
    assert infer('Code', ['001', '002']) == 'text'
    assert infer('person_id', [123, 124]) == 'text'
    assert infer('age', [0, 12, 25]) == 'number'
    assert infer('day', ['04/05/2026', '15/05/2026']) == 'date'
    assert infer('mixed', ['hello', '25']) == 'text'
    assert infer('empty', [None, '']) == 'text'


def test_large_text_defaults_preserve_numbers_dates_and_saved_choices():
    from datetime import date, timedelta
    rows = [(i + 2, [f'Response {i}', i, (date(2026, 1, 1) + timedelta(days=i)).isoformat(), 'Repeated']) for i in range(101)]
    columns = initial_columns(['Comment', 'Amount', 'Date', 'City'], rows)
    assert [(c.kind, c.included) for c in columns] == [('text', False), ('number', True), ('date', True), ('category', True)]
    assert initial_columns(['Comment'], rows[:100])[0].included
    previous = {'columns': [c.model_dump() for c in columns]}
    previous['columns'][0]['included'] = True
    previous['columns'][1]['included'] = False
    restored = initial_columns(['Comment', 'Amount', 'Date', 'City'], rows, previous)
    assert [c.included for c in restored] == [True, False, True, True]


def test_all_large_text_fields_can_be_previewed_and_opted_in(env):
    client, store = env
    raw = ('Comment\n' + '\n'.join(f'Response {i}' for i in range(101)) + '\n').encode()
    p = stage(client, raw)
    assert not p['config']['columns'][0]['included']
    assert p['profiles'][0]['distinct'] == 101
    confirm(client, p, 400)  # Saving still requires a deliberate column selection.
    p['config']['columns'][0]['included'] = True
    configured = client.post(f'/api/custom/columns/{p["token"]}/configure', json=p['config'])
    assert configured.status_code == 200, configured.text
    saved = confirm(client, configured.json())
    assert len(store.get(saved['id'])['records']) == 101
    reopened = client.post(f'/api/custom/projects/{saved["id"]}/columns').json()
    assert reopened['config']['columns'][0]['included']
    replacement = stage(client, raw, pid=saved['id'])
    assert replacement['config']['columns'][0]['included']


def test_reopen_reads_project_once(env, monkeypatch):
    client, store = env
    p = stage(client); p['config']['dateFormat'] = 'uk'
    saved = confirm(client, p)
    original = store.get
    reads = []

    def tracked_get(pid):
        reads.append(pid)
        return original(pid)

    monkeypatch.setattr(store, 'get', tracked_get)
    opened = client.post('/api/custom/projects/' + saved['id'] + '/columns')
    assert opened.status_code == 200, opened.text
    assert reads == [saved['id']]
    assert opened.json()['config'] == saved['customConfig']
    assert opened.json()['profiles']


def test_import_review_and_isolation(env):
    client, store = env
    p = stage(client)
    assert p['issues'] and p['profiles'][2]['ambiguous']
    p['config']['columns'][2]['invalidAsMissing'] = True
    confirm(client, p, 422)  # Ambiguity always needs a convention.
    p['config']['dateFormat'] = 'uk'
    saved = confirm(client, p)
    assert saved['customConfig']['dateFormat'] == 'uk'
    assert client.get('/api/kobo/projects').json() == []
    assert client.get('/api/kobo/projects/' + saved['id']).status_code == 404
    assert store.get(saved['id'])['records'][0]['values'][p['config']['columns'][0]['id']] == '001'
    assert store.get(saved['id'])['records'][0]['values'][p['config']['columns'][2]['id']] == '2026-05-04'
    analysis = client.post(f'/api/custom/projects/{saved["id"]}/analysis', json={'settings': saved['settings']})
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()['total'] == 2
    assert next(q for q in analysis.json()['questions'] if q['name'] == 'Age')['statistics']['mean'] == 25
    assert client.post('/api/custom/columns/confirm', json={'token': p['token'], 'name': 'Again', 'config': p['config']}).status_code == 410


def test_invalid_values_and_categories(env):
    client, store = env
    p = stage(client, b'Age,Needs\n12,"water; food"\noops,food\n')
    age, needs = p['config']['columns']
    age['kind'] = 'number'
    needs.update(kind='multiple', separator=';', categories={'food': 'Food', 'water': 'Water'})
    confirm(client, p, 422)
    age['invalidAsMissing'] = True
    saved = confirm(client, p)
    records = store.get(saved['id'])['records']
    assert records[0]['values'][needs['id']] == ['Water', 'Food']
    assert records[1]['values'][age['id']] is None
    assert saved['warnings']


def test_xlsx_header_rows_native_dates_duplicates_and_blank_headers(env):
    client, store = env
    book = openpyxl.Workbook()
    book.active.title = 'Empty'
    sheet = book.create_sheet('Data')
    sheet.append(['Report title'])
    sheet.append(['Name', 'Name', None, 'Date'])
    sheet.append(['أحمد', 'Ali', '001', datetime(2026, 10, 25)])
    out = io.BytesIO(); book.save(out)
    p = stage(client, out.getvalue(), 'data.xlsx')
    assert p['config']['columns'] == []
    p['config'].update(worksheet='Data', headerRow=2)
    r = client.post(f'/api/custom/columns/{p["token"]}/configure', json=p['config'])
    assert r.status_code == 200, r.text
    p = r.json()
    assert len({c['id'] for c in p['config']['columns']}) == 4
    saved = confirm(client, p)
    record = store.get(saved['id'])['records'][0]
    assert record['row'] == 3
    assert record['values'][p['config']['columns'][3]['id']] == '2026-10-25'
    edit = client.post(f'/api/custom/projects/{saved["id"]}/columns').json()
    assert edit['config'] == p['config']


def test_reupload_preserves_ids_formats_and_reports_changes(env):
    client, store = env
    p = stage(client); p['config']['dateFormat'] = 'uk'; saved = confirm(client, p)
    changed = stage(client, b'Date,Code,Age,New\n04/05/2026,003,40,x\n', pid=saved['id'])
    assert changed['config']['dateFormat'] == 'uk'
    ids = {c['header']: c['id'] for c in p['config']['columns']}
    assert all(c['id'] == ids[c['header']] for c in changed['config']['columns'] if c['header'] in ids)
    assert any(c['change'] == 'removed' and c['label'] == 'City' for c in changed['changes'])
    assert any(c['change'] == 'added' and c['label'] == 'New' for c in changed['changes'])
    updated = confirm(client, changed)
    assert updated['id'] == saved['id']
    assert len(store.list()) == 1


def test_backup_settings_data_geography_and_rename(env):
    client, store = env
    p = stage(client); p['config']['dateFormat'] = 'uk'; saved = confirm(client, p)
    root = '/api/custom/projects/' + saved['id']
    assert client.patch(root, json={'name': 'Renamed'}).status_code == 200
    assert client.put(root+'/settings', json=saved['settings']).status_code == 200
    settings_file = client.post(root+'/settings/export', json=saved['settings'])
    assert settings_file.status_code == 200
    assert client.post(root+'/settings/import-preview', json={'settings': saved['settings'], 'package': settings_file.json()}).status_code == 200
    columns = [c['id'] for c in p['config']['columns']]
    data = client.post(root+'/data', json={'settings': saved['settings'], 'columns': columns})
    assert data.status_code == 200, data.text
    export = client.post(root+'/data/export', json={'settings': saved['settings'], 'columns': columns})
    assert export.status_code == 200 and export.content.startswith(b'PK')
    geo = client.post(root+'/geography', json={'settings': saved['settings'], 'field': columns[3]})
    assert geo.status_code == 200, geo.text
    backup = client.get(root+'/backup')
    assert backup.status_code == 200
    assert client.post('/api/kobo/backups/restore', files={'file': ('backup.zip', backup.content)}).status_code == 400
    restored = client.post('/api/custom/backups/restore', files={'file': ('backup.zip', backup.content)})
    assert restored.status_code == 200, restored.text
    assert restored.json()['customConfig'] == p['config']
    assert restored.json()['name'] == 'Renamed'
    assert client.post('/api/custom/projects/'+restored.json()['id']+'/columns').status_code == 200


def test_project_backup_keeps_names_both_contacts_and_original_upload(env):
    client, store = env
    raw = 'Name,Primary Contact,Secondary Contact,Document ID\nاسم تجريبي,07501234567,07701234567,00001234\nاسم آخر,07507654321,,00005678\n'.encode('utf-8-sig')
    preview = stage(client, raw)
    for column in preview['config']['columns']:
        column['kind'] = 'text'
    saved = confirm(client, preview)
    original = store.get(saved['id'])
    source = store.sources(saved['id'])
    assert raw in source
    root = '/api/custom/projects/' + saved['id']
    backup = client.get(root + '/backup')
    assert backup.status_code == 200
    restored = client.post('/api/custom/backups/restore', files={'file': ('project.zip', backup.content)})
    assert restored.status_code == 200, restored.text
    imported_id = restored.json()['id']
    assert store.get(imported_id)['records'] == original['records']
    assert store.sources(imported_id) == source
    assert store.get(saved['id']) == original
    assert len(store.list()) == 2


def test_type_changes_clear_incompatible_filters_and_charts(env):
    client, _ = env
    p = stage(client); p['config']['dateFormat'] = 'uk'; saved = confirm(client, p)
    age = p['config']['columns'][1]['id']
    settings = saved['settings']
    settings['filters'] = [{'field': age, 'kind': 'number', 'minimum': 10}]
    settings['questions'][age]['chart'] = 'histogram'
    root = '/api/custom/projects/'+saved['id']
    assert client.put(root+'/settings', json=settings).status_code == 200
    edit = client.post(root+'/columns').json()
    edit['config']['columns'][1]['kind'] = 'text'
    updated = confirm(client, edit)
    assert updated['settings']['filters'] == []
    assert updated['settings']['questions'][age]['chart'] is None
    assert client.post(root+'/analysis', json={'settings': updated['settings']}).status_code == 200


def test_validation_cancellation_and_source_identity(env):
    client, _ = env
    assert client.post('/api/custom/columns/preview', files={'data': ('data.csv', b'')}).status_code == 400
    assert client.post('/api/custom/columns/preview', files={'data': ('data.xls', b'not excel')}).status_code == 400
    p = stage(client)
    original = copy.deepcopy(p)
    p['config']['columns'][0]['id'] = 'changed'
    confirm(client, p, 400)
    p = original
    assert client.delete('/api/custom/columns/'+p['token']).status_code == 200
    confirm(client, p, 410)


def test_saved_views_flag_type_changes_and_broken_breakdowns_are_repaired(env):
    client, _ = env
    p = stage(client); p['config']['dateFormat'] = 'uk'; saved = confirm(client, p)
    age, city = p['config']['columns'][1]['id'], p['config']['columns'][3]['id']
    settings = saved['settings']
    settings['questions'][age].update(chart='heatmap', breakdown=city, breakdowns=[city])
    settings['savedViews'] = [{'id': 'view', 'name': 'Original', 'snapshot': copy.deepcopy(settings), 'questionKinds': {age: 'number', city: 'category'}}]
    root = '/api/custom/projects/'+saved['id']
    assert client.put(root+'/settings', json=settings).status_code == 200
    edit = client.post(root+'/columns').json()
    edit['config']['columns'][3]['kind'] = 'multiple'
    edit['config']['columns'][3]['separator'] = ';'
    updated = confirm(client, edit)
    assert updated['settings']['savedViews'][0]['reviewIssues']
    assert updated['settings']['questions'][age]['chart'] is None
    assert client.post(root+'/analysis', json={'settings': updated['settings']}).status_code == 200


def test_export_recovery_is_bound_to_workspace(env):
    client, _ = env
    p = stage(client); p['config']['dateFormat'] = 'uk'; saved = confirm(client, p)
    task_id = '12345678-1234-1234-1234-123456789012'
    task = {'filename': 'analysis.xlsx', 'input': {'projectId': saved['id'], 'projectName': saved['name'], 'settings': saved['settings'], 'format': 'excel'}}
    assert client.put('/api/custom/export-recovery/'+task_id, json=task).status_code == 200
    assert client.get('/api/custom/export-recovery').json()[0]['input']['workspace'] == 'custom'
    assert client.get('/api/kobo/export-recovery').json() == []


@pytest.mark.parametrize('mode', ['native', 'binary'])
def test_custom_excel_preparation(env, tmp_path, mode):
    from queue import Queue
    from backend.kobo import KoboSettings, analyze
    from backend.kobo_excel_source import response_layout, query_config
    from backend.kobo_excel_slicers import slicer_specs
    from backend.kobo_excel_prepare import prepare_workbook
    client, store = env
    p = stage(client, b'Code,Age,Date,Needs\n001,20,04/05/2026 09:45,water;food\n002,30,15/05/2026,food\n')
    p['config']['dateFormat'] = 'uk'
    p['config']['columns'][2]['includeTime'] = True
    p['config']['columns'][3].update(kind='multiple', separator=';')
    saved = confirm(client, p)
    project = store.get(saved['id'])
    settings = KoboSettings.model_validate(saved['settings'])
    snapshot = analyze(project, settings)
    layouts = response_layout(project, settings)
    config = query_config(project, snapshot, settings, slicer_specs(project, settings), layouts)
    config['excelMode'] = mode
    target = tmp_path / f'custom-{mode}.xlsx'
    prepare_workbook(target, layouts, config, settings, Queue())
    book = openpyxl.load_workbook(target, data_only=True)
    try:
        assert book.sheetnames
        assert any('001' in row for sheet in book for row in sheet.iter_rows(values_only=True))
        assert any('09:45:00' in row for sheet in book for row in sheet.iter_rows(values_only=True))
    finally:
        book.close()


def test_time_extraction_formats_and_date_only_values():
    assert extract_time('2026-10-25T09:30:45.125+03:00') == '09:30:45.125000'
    assert extract_time('25/10/2026 9:30', 'uk') == '09:30:00'
    assert extract_time('10/25/2026 00:00:00', 'us') == '00:00:00'
    assert extract_time(datetime(2026, 10, 25, 15, 42, 12)) == '15:42:12'
    assert extract_time('2026-10-25') is None
    with pytest.raises(ValueError): extract_time('25/10/2026 25:00', 'uk')
    with pytest.raises(ValueError): extract_time('04/05/2026 12:00')


def test_optional_time_column_and_filter_defaults_persist(env):
    client, store = env
    p = stage(client, b'Date,City\n2026-10-25T09:30:45+03:00,Mosul\n2026-10-26,Baghdad\n2026-10-27T00:00:00Z,Mosul\n')
    column = p['config']['columns'][0]
    assert p['config']['filterFields'] == []
    assert column['includeTime'] is False
    assert p['profiles'][0]['timeCount'] == 2
    column.update(includeTime=True, timeLabel='Start time')
    time_id = column['id'] + '__time'
    p['config']['filterFields'] = [time_id]
    saved = confirm(client, p)
    assert saved['settings']['filterFields'] == [time_id]
    project = store.get(saved['id'])
    assert [r['values'][time_id] for r in project['records']] == ['09:30:45', None, '00:00:00']
    assert project['records'][0]['values'][column['id']] == '2026-10-25'
    assert next(q for q in project['form']['questions'] if q['id'] == time_id)['name'] == 'Start time'
    root = '/api/custom/projects/' + saved['id']
    assert client.post(root+'/analysis', json={'settings': saved['settings']}).status_code == 200
    edit = client.post(root+'/columns').json()
    assert edit['config']['columns'][0]['includeTime'] is True
    assert edit['config']['filterFields'] == [time_id]
    edit['config']['columns'][0]['includeTime'] = False
    confirm(client, edit, 400)  # Stale time-filter references cannot be saved.
    edit['config']['filterFields'] = []
    updated = confirm(client, edit)
    assert not any(q['id'] == time_id for q in updated['form']['questions'])
    assert updated['settings']['filterFields'] == []


def test_no_filters_selected_by_default_and_deselection_clears_active_filters(env):
    client, _ = env
    p = stage(client); p['config']['dateFormat'] = 'uk'
    saved = confirm(client, p)
    assert saved['settings']['filterFields'] == []
    root = '/api/custom/projects/' + saved['id']
    age = p['config']['columns'][1]['id']
    settings = saved['settings']
    settings['filterFields'] = [age]
    settings['filters'] = [{'field': age, 'kind': 'number', 'minimum': 25}]
    assert client.put(root+'/settings', json=settings).status_code == 200
    edit = client.post(root+'/columns').json()
    assert edit['config']['filterFields'] == [age]
    edit['config']['filterFields'] = []
    updated = confirm(client, edit)
    assert updated['settings']['filters'] == []
    assert updated['settings']['filterFields'] == []


def test_invalid_timestamp_acceptance_and_native_excel_time(env):
    client, store = env
    p = stage(client, b'Date\n2026-10-25T99:00:00Z\n2026-10-26T14:00:00Z\n')
    p['config']['columns'][0].update(kind='date', includeTime=True, invalidAsMissing=True)
    saved = confirm(client, p)
    column_id = p['config']['columns'][0]['id']
    assert store.get(saved['id'])['records'][0]['values'][column_id + '__time'] is None
    book = openpyxl.Workbook(); book.active.append(['Date']); book.active.append([datetime(2026, 10, 25, 14, 30)])
    out = io.BytesIO(); book.save(out)
    p = stage(client, out.getvalue(), 'native.xlsx')
    p['config']['columns'][0]['includeTime'] = True
    saved = confirm(client, p)
    assert store.get(saved['id'])['records'][0]['values'][p['config']['columns'][0]['id'] + '__time'] == '14:30:00'


def test_time_columns_survive_reupload_and_backup(env):
    client, store = env
    p = stage(client, b'Date,City\n2026-10-25T14:30:00Z,Mosul\n')
    column = p['config']['columns'][0]
    column.update(includeTime=True, timeLabel='Visit time')
    time_id = column['id'] + '__time'
    p['config']['filterFields'] = [time_id]
    saved = confirm(client, p)
    replacement = stage(client, b'City,Date\nBaghdad,2026-10-26T18:45:00+03:00\n', pid=saved['id'])
    assert replacement['config']['columns'][1]['includeTime'] is True
    assert replacement['config']['filterFields'] == [time_id]
    updated = confirm(client, replacement)
    assert store.get(updated['id'])['records'][0]['values'][time_id] == '18:45:00'
    backup = client.get('/api/custom/projects/' + updated['id'] + '/backup')
    restored = client.post('/api/custom/backups/restore', files={'file': ('time.zip', backup.content)})
    assert restored.status_code == 200, restored.text
    edit = client.post('/api/custom/projects/' + restored.json()['id'] + '/columns').json()
    assert edit['config']['columns'][1]['timeLabel'] == 'Visit time'
    assert edit['config']['filterFields'] == [time_id]


@pytest.mark.parametrize('separator', [',', ';', '|', '\t', '\n', ' ', '::'])
def test_multiple_choice_delimiters_split_trim_deduplicate_and_persist(env, separator):
    import csv
    client, store = env
    cell = separator.join([' water ', '', 'food', 'water', ' '])
    raw = io.StringIO(newline='')
    writer = csv.writer(raw); writer.writerow(['Needs']); writer.writerow([cell]); writer.writerow(['food'])
    p = stage(client, raw.getvalue().encode())
    column = p['config']['columns'][0]
    column.update(kind='multiple', separator=separator)
    check = client.post('/api/custom/columns/'+p['token']+'/configure', json=p['config'])
    assert check.status_code == 200, check.text
    assert check.json()['profiles'][0]['preview'][0]['value'] == ['water', 'food']
    assert check.json()['profiles'][0]['values'] == ['water', 'food']
    saved = confirm(client, p)
    assert store.get(saved['id'])['records'][0]['values'][column['id']] == ['water', 'food']
    edit = client.post('/api/custom/projects/'+saved['id']+'/columns').json()
    assert edit['config']['columns'][0]['separator'] == separator


def test_multiple_choice_without_delimiter_can_be_reviewed_but_not_saved(env):
    client, _ = env
    p = stage(client, b'Needs\nwater;food\nfood\n')
    p['config']['columns'][0].update(kind='multiple', separator='')
    review = client.post('/api/custom/columns/'+p['token']+'/configure', json=p['config'])
    assert review.status_code == 200
    assert review.json()['issues'][0]['message'] == 'Choose a multiple-choice separator.'
    confirm(client, p, 422)
    p['config']['columns'][0]['separator'] = ';'
    assert confirm(client, p)['id']
