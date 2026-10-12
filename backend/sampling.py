"""Read-only legal-services sampling view. Table and export share one pipeline."""
from __future__ import annotations

import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .excel_export import Font, PatternFill, append_frame
from .legal_platform import (
    DATE_HINT, EXPLORER_DATE_FORMAT, _safe_export, clean_id, display_value,
    YEAR_MONTH, distinct_transform, filter_rows, format_excel_dates, phone_digits,
)
from .filter_selection import selection_mask, selection_value
from .linked_filters import available_values, search_mask


LEADING = (
    ("project", "Projects", "Project"),
    ("location", "Project Location"),
    ("name", "Name"),
    ("gender", "Gender"),
    ("dob", "DoB", "Date of Birth"),
    ("age", "Age"),
    ("phone", "Contact Number", "Primary Contact Number"),
    ("secondary", "Secondary Contact Number"),
    ("nationality", "Nationality"),
    ("community", "Community Type"),
    ("type", "Type of Service Provided"),
    ("document", "Type of Document"),
)


def text(value):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip().casefold().replace("ک", "ك").replace("ی", "ي")


def sampling_dates(series):
    # Match the workbook parser: ISO is year-first; local source dates are day-first.
    parsed = distinct_transform(series, lambda value: pd.to_datetime(
        value, errors="coerce", dayfirst=not bool(re.match(r"^\d{4}-", str(value).strip())), utc=True))
    return pd.to_datetime(parsed, utc=True).dt.tz_localize(None)


def find_column(columns, *names):
    # Match complete English field names before bilingual suffixes, never substrings
    # such as Secondary Contact Number, Age Gender Group, or Nationality Stateless.
    for name in names:
        for column in columns:
            if text(column) == text(name):
                return column
    for name in names:
        for column in columns:
            suffix = text(column)[len(text(name)):]
            if text(column).startswith(text(name)) and suffix and re.fullmatch(r"[\s/\-:]*[\u0600-\u06ff\s/\-:]+", suffix):
                return column
    return None


def prepare(frame):
    if frame is None:
        return None, {}, "Sampling is unavailable because Legal Services has not been imported."
    fields = {key: find_column(frame.columns, *names) for key, *names in LEADING}
    fields.update(status=find_column(frame.columns, "Service Status"),
                  completed=find_column(frame.columns, "Date Service Completed"),
                  beneficiary=find_column(frame.columns, "Beneficiary ID"))
    required = {"status": "Service Status", "type": "Type of Service Provided", "phone": "primary Contact Number",
                "gender": "Gender", "completed": "Date Service Completed", "beneficiary": "Beneficiary ID"}
    missing = [label for key, label in required.items() if not fields[key]]
    if missing:
        return None, fields, "Sampling is unavailable. Missing source fields: " + ", ".join(missing) + "."
    status = frame[fields["status"]].fillna("").map(text)
    service = frame[fields["type"]].fillna("").map(text)
    completed = status.str.match(r"^completed(?:\b|$)") | status.isin(["اكتملت", "مكتملة", "مكتمل"])
    representation = service.str.match(r"^legal representation(?:\b|$)") | service.isin(["تمثيل", "التمثيل القانوني", "تمثيل قانوني"])
    phones = frame[fields["phone"]].map(phone_digits)
    eligible = frame.loc[completed & representation & phones.str.fullmatch(r"(?:77|78|75)[0-9]{8}")].copy().reset_index(drop=True)
    leading = [fields[key] for key, *_ in LEADING if fields[key]]
    columns = list(dict.fromkeys(leading + list(frame.columns)))
    return eligible[columns], fields, ""


def pipeline(frame, filters=None, search="", removeDuplicateBeneficiaryIds=False, removeDuplicateContacts=False, facets=False):
    base, fields, reason = prepare(frame)
    if reason:
        return None, fields, {}, reason
    filters = filters or {}
    searched = search_mask(base, search)
    selected = base.loc[searched]
    masks = {}
    for column, values in filters.items():
        if column in base and values:
            if DATE_HINT.search(column) and all(YEAR_MONTH.fullmatch(selection_value(str(value))) for value in values):
                matching = base.index[selection_mask(sampling_dates(base[column]).dt.to_period("M").astype(str), values)]
            else:
                matching = filter_rows(base, column, values).index
            masks[column] = pd.Series(base.index.isin(matching), index=base.index)
            selected = selected.loc[selected.index.isin(matching)]
    options = {}
    if facets:
        series = {column: (sampling_dates(base[column]).dt.to_period("M").astype(str) if DATE_HINT.search(column)
                           else base[column].fillna("").astype(str).str.strip()) for column in base.columns}
        options = available_values(series, filters, searched, masks)
    dates = sampling_dates(selected[fields["completed"]])
    priority = selected.loc[dates.sort_values(ascending=False, kind="stable", na_position="last").index]
    seen_ids, seen_phones, retained = set(), set(), []
    if removeDuplicateBeneficiaryIds or removeDuplicateContacts:
        for index, beneficiary, phone in zip(priority.index, priority[fields["beneficiary"]].map(clean_id), priority[fields["phone"]].map(phone_digits)):
            if removeDuplicateBeneficiaryIds and beneficiary and beneficiary in seen_ids:
                continue
            if removeDuplicateContacts and phone in seen_phones:
                continue
            retained.append(index)
            if beneficiary:
                seen_ids.add(beneficiary)
            seen_phones.add(phone)
    else:
        retained = list(priority.index)
    male, female, other = [], [], []
    for index, gender in selected.loc[retained, fields["gender"]].map(text).items():
        group = male if re.match(r"^male\b", gender) or gender == "ذكر" else female if re.match(r"^female\b", gender) or gender in {"أنثى", "انثى"} else other
        group.append(index)
    ordered = []
    for position in range(max(len(male), len(female))):
        if position < len(male): ordered.append(male[position])
        if position < len(female): ordered.append(female[position])
    return selected.loc[ordered + other], fields, options, ""


def table(frame, *, page=1, pageSize=100, **query):
    selected, fields, options, reason = pipeline(frame, facets=True, **query)
    total = 0 if selected is None else len(selected)
    page = min(page, max(1, (total + pageSize - 1) // pageSize))
    rows = [] if selected is None else [{column: display_value(value) for column, value in row.items()}
                                      for row in selected.iloc[(page - 1) * pageSize:page * pageSize].to_dict("records")]
    return {"available": not bool(reason), "reason": reason, "columns": [] if selected is None else list(selected.columns),
            "rows": rows, "total": total, "page": page, "pageSize": pageSize, "filterOptions": options,
            "quickFilters": [fields[key] for key in ("project", "location", "completed") if fields.get(key)]}


def export(frame, **query):
    selected, _, _, reason = pipeline(frame, **query)
    if reason:
        raise ValueError(reason)
    selected = _safe_export(selected)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sampling"
    append_frame(sheet, selected)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for index, column in enumerate(selected.columns, 1):
        width = max([len(str(column))] + [len(str(value)) for value in selected[column].head(250)]) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(45, max(12, width))
    format_excel_dates(workbook, EXPLORER_DATE_FORMAT)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
