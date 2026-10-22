"""XLSForm ingestion and respondent-based analysis, shared by all Kobo outputs."""
from __future__ import annotations

import csv
import ast
import io
import json
import math
import os
import re
import sqlite3
import statistics
import unicodedata
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal

import openpyxl
from pydantic import BaseModel, Field, field_validator, model_validator

from .kobo_geography import GeographyConfig
from .kobo_indicators import IndicatorDefinition
from .file_security import validate_xlsx_archive
from .kobo_percentages import apply_display_percentages


class KoboFilter(BaseModel):
    field: str
    kind: Literal['category', 'number', 'date'] = 'category'
    values: list[str] = Field(default_factory=list)
    minimum: float | str | None = None
    maximum: float | str | None = None

    @model_validator(mode='after')
    def validate_bounds(self):
        if self.kind == 'category': return self
        bounds = []
        for value in [self.minimum, self.maximum]:
            if value is None: bounds.append(None); continue
            try:
                parsed = float(value) if self.kind == 'number' else date.fromisoformat(str(value))
                if isinstance(parsed, float) and not math.isfinite(parsed): raise ValueError()
                bounds.append(parsed)
            except (ValueError, TypeError) as exc: raise ValueError(f'Invalid {self.kind} filter bound.') from exc
        if all(v is not None for v in bounds) and bounds[0] > bounds[1]: raise ValueError('Filter minimum must not exceed its maximum.')
        return self


class ChartAppearance(BaseModel):
    metric: Literal['auto', 'count', 'percent'] = 'auto'
    labels: Literal['auto', 'hidden', 'count', 'percent', 'both'] = 'auto'
    legend: Literal['auto', 'hidden', 'top', 'bottom', 'left', 'right'] = 'auto'
    fontSize: Literal[0, 10, 12, 14, 16] = 0
    height: Literal[0, 320, 420, 560, 720] = 0
    grid: Literal['auto', 'show', 'hide'] = 'auto'
    bars: Literal['group', 'stack'] = 'group'


class QuestionSettings(BaseModel):
    appearance: ChartAppearance = Field(default_factory=ChartAppearance)
    palette: Literal['Balanced', 'Blue', 'Teal', 'Warm', 'Solid blue', 'Ocean', 'Botanical', 'Jewel', 'Contrast'] | None = None
    chart: Literal['bar', 'column', 'pie', 'donut', 'line', 'histogram', 'heatmap'] | None = None
    colors: dict[str, str] = Field(default_factory=dict)
    pointColors: dict[str, str] = Field(default_factory=dict)
    bins: int = Field(10, ge=1, le=100)
    numericMode: Literal['auto', 'bins', 'values'] = 'auto'
    dateUnit: Literal['day', 'month', 'year'] = 'month'
    trendDateField: str = ''
    trendUnit: Literal['day', 'month', 'year'] = 'month'
    breakdown: str = ''
    breakdowns: list[str] = Field(default_factory=list, max_length=3)
    breakdownDisplay: Literal['count', 'percent', 'both'] = 'both'
    answerSort: Literal['form', 'count-desc', 'count-asc', 'label-asc', 'label-desc'] = 'form'

    @field_validator('pointColors')
    @classmethod
    def validate_point_colors(cls, value):
        for key, color in value.items():
            try: parts = json.loads(key)
            except (ValueError, TypeError): raise ValueError('Invalid chart mark identifier.')
            if not (isinstance(parts, list) and len(parts) == 3 and isinstance(parts[0], list)
                    and all(isinstance(field, str) for field in parts[0])
                    and (parts[1] is None or isinstance(parts[1], str)) and isinstance(parts[2], str)):
                raise ValueError('Invalid chart mark identifier.')
            if not re.fullmatch(r'#[0-9a-fA-F]{6}', color): raise ValueError('Chart colors must use #RRGGBB format.')
        return value

    @field_validator('colors')
    @classmethod
    def validate_colors(cls, value):
        if any(not re.fullmatch(r'#[0-9a-fA-F]{6}', color) for color in value.values()): raise ValueError('Chart colors must use #RRGGBB format.')
        return value


class AnalysisViewSettings(BaseModel):
    language: Literal['en', 'ar'] = 'en'
    filters: list[KoboFilter] = Field(default_factory=list)
    filterFields: list[str] = Field(default_factory=list)
    hidden: list[str] = Field(default_factory=list)
    metadata: bool = False
    hideBlanks: bool = True
    hideZero: bool = True
    hideEmpty: bool = True
    percentageDecimals: Literal[0, 1, 2] = 0
    questionSort: Literal['form', 'name-asc', 'name-desc', 'answered-desc', 'answered-asc'] = 'form'
    questions: dict[str, QuestionSettings] = Field(default_factory=dict)
    viewMode: Literal['charts', 'table'] = 'charts'


class SavedAnalysisView(BaseModel):
    id: str = Field(min_length=1, max_length=100)
    name: str = Field(min_length=1, max_length=160)
    description: str = Field('', max_length=2000)
    snapshot: AnalysisViewSettings
    updatedAt: str = Field('', max_length=60)
    reviewIssues: list[str] = Field(default_factory=list)
    questionKinds: dict[str, str] = Field(default_factory=dict)

    @field_validator('name')
    @classmethod
    def meaningful_name(cls, value):
        if not value.strip(): raise ValueError('Enter a saved view name.')
        return value.strip()


class WorkspaceReport(BaseModel):
    title: str = Field('', max_length=300)
    findings: str = Field('', max_length=30000)
    methodology: str = Field('', max_length=30000)
    recommendations: str = Field('', max_length=30000)
    questions: list[str] = Field(default_factory=list, max_length=500)


class KoboSettings(AnalysisViewSettings):
    version: int = 1
    excelGroups: list[str] | None = None
    indicators: list[IndicatorDefinition] = Field(default_factory=list, max_length=100)
    savedViews: list[SavedAnalysisView] = Field(default_factory=list, max_length=100)
    activeViewId: str = ''
    dashboardQuestions: list[str] = Field(default_factory=list, max_length=500)
    geographyField: str = ''
    geographyConfigs: dict[str, GeographyConfig] = Field(default_factory=dict)
    report: WorkspaceReport = Field(default_factory=WorkspaceReport)

    @model_validator(mode='after')
    def unique_saved_views(self):
        if len({v.id for v in self.savedViews}) != len(self.savedViews): raise ValueError('Saved view identifiers must be unique.')
        if len({v.name.casefold() for v in self.savedViews}) != len(self.savedViews): raise ValueError('Each saved view must have a different name.')
        if self.activeViewId and not any(v.id == self.activeViewId for v in self.savedViews): raise ValueError('Active saved view no longer exists.')
        return self

    @field_validator('indicators')
    @classmethod
    def unique_indicator_ids(cls, value):
        if len({item.id for item in value}) != len(value): raise ValueError('Indicator identifiers must be unique.')
        return value


def new_project_settings(form: dict) -> KoboSettings:
    """Initialize new projects only; stored and explicitly supplied settings keep their colors."""
    return KoboSettings(language='en', hidden=[q['id'] for q in form['questions'] if q['identifying']],
                        questions={q['id']: QuestionSettings(palette='Blue') for q in form['questions']})


def norm(value: Any) -> str:
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFC', str(value or '')).strip()).casefold()


def scalar(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def labels(record: dict) -> dict[str, str]:
    result = {}
    for key, value in record.items():
        if not value or not key or not str(key).startswith('label'):
            continue
        if key == 'label': result['default'] = str(value)
        elif re.search(r'(\(ar\)|::ar$|arabic|العربية)', key, re.I): result['ar'] = str(value)
        elif re.search(r'(\(en\)|::en$|english)', key, re.I): result['en'] = str(value)
        else: result.setdefault('default', str(value))
    return result


def label(item: dict, language: str) -> str:
    ls = item.get('labels', {})
    return ls.get(language) or ls.get('default') or ls.get('en') or ls.get('ar') or item.get('name') or item.get('id', '')


def parse_form(raw: bytes) -> dict:
    validate_xlsx_archive(raw)
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        if 'survey' not in book.sheetnames: raise ValueError('XLSForm requires a survey sheet.')
        def records(sheet):
            if sheet not in book.sheetnames: return []
            rows = book[sheet].iter_rows(values_only=True)
            header = [str(v).strip() if v is not None else '' for v in next(rows, [])]
            return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]
        choices = defaultdict(list)
        for row in records('choices'):
            if missing(row.get('list_name')) or missing(row.get('name')): continue
            choices[str(row['list_name'])].append({'name': str(row['name']), 'labels': labels(row), 'attributes': {k: scalar(v) for k, v in row.items() if k and not k.startswith('label')}})
        settings = next(iter(records('settings')), {})
        questions, groups, stack, warnings, identifiers, notes = [], [], [], [], set(), []
        metadata_types = {'start', 'end', 'today', 'deviceid', 'subscriberid', 'simserial', 'phonenumber', 'username', 'email', 'audit'}
        supported = {'text', 'integer', 'decimal', 'date', 'datetime', 'time', 'select_one', 'select_multiple', 'calculate', *metadata_types}
        for row in records('survey'):
            kind = str(row.get('type') or '').strip().replace('begin group', 'begin_group').replace('end group', 'end_group').replace('begin repeat', 'begin_repeat').replace('end repeat', 'end_repeat')
            if not kind: continue
            base = kind.split()[0]
            if base in {'end_group', 'end_repeat'}:
                if not stack or stack[-1]['repeat'] != (base == 'end_repeat'): raise ValueError('Unbalanced XLSForm group/repeat rows.')
                stack.pop(); continue
            name = str(row.get('name') or '').strip()
            if base == 'note' and not name:
                # Display-only instructions do not need a response field.
                continue
            if not name: raise ValueError(f'Form row of type {kind} has no question/group name.')
            path = '/'.join([g['name'] for g in stack] + [name])
            if path in identifiers: raise ValueError(f'Duplicate form path: {path}')
            identifiers.add(path)
            if base in {'begin_group', 'begin_repeat'}:
                group = {'id': path, 'name': name, 'labels': labels(row), 'parent': stack[-1]['id'] if stack else '', 'repeat': base == 'begin_repeat', 'order': len(groups)}
                groups.append(group); stack.append(group); continue
            if base == 'note':
                notes.append({'id': path, 'name': name, 'labels': labels(row), 'group': stack[-1]['id'] if stack else ''}); continue
            if base not in supported:
                warnings.append(f'{path}: unsupported type {kind}; retained in source only.'); continue
            repeat = next((g['id'] for g in reversed(stack) if g['repeat']), '')
            choice_list = kind.split()[1] if base in {'select_one', 'select_multiple'} and len(kind.split()) > 1 else ''
            if choice_list and not choices[choice_list]: raise ValueError(f'{path}: choice list {choice_list} is missing or empty in choices.')
            numeric = base in {'integer', 'decimal'}
            q = {'id': path, 'name': name, 'type': base, 'kind': 'number' if numeric else 'date' if base in {'date', 'datetime', 'today', 'start', 'end'} else 'multiple' if base == 'select_multiple' else 'category' if base == 'select_one' else 'text', 'labels': labels(row), 'group': stack[-1]['id'] if stack else '', 'repeat': repeat, 'metadata': base in metadata_types or base == 'calculate', 'choices': choices[choice_list] if choice_list else [], 'order': len(questions), 'relevant': str(row.get('relevant') or ''), 'choiceFilter': str(row.get('choice_filter') or ''), 'identifying': base == 'text' and bool(re.search(r'(respondent.*name|name.*respondent|enumerator.*name|name.*enumerator|full.?name|phone|telephone|email|اسم.*(المستجيب|المستجوب)|name_of)', ' '.join([name, *labels(row).values()]), re.I))}
            questions.append(q)
        if stack: raise ValueError('XLSForm has an unclosed group or repeat.')
        default = str(settings.get('default_language') or 'English (en)')
        if 'settings' not in book.sheetnames: warnings.append('No settings sheet; English labels are the default.')
        return {'questions': questions, 'groups': groups, 'notes': notes, 'settings': {k: scalar(v) for k, v in settings.items() if k}, 'language': 'ar' if re.search(r'arabic|العربية|\(ar\)|^ar$', default, re.I) else 'en', 'warnings': warnings}
    finally: book.close()


def data_worksheets(raw: bytes, filename: str) -> list[str]:
    if not filename.lower().endswith('.xlsx'): raise ValueError('Choose an .xlsx response workbook.')
    validate_xlsx_archive(raw)
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        return [sheet.title for sheet in book.worksheets]
    finally: book.close()


def parse_data(raw: bytes, filename: str, worksheet: str = '') -> list[dict]:
    if filename.lower().endswith('.csv'):
        if worksheet: raise ValueError('Worksheet selection is only available for Excel workbooks.')
        try: text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text = raw.decode('utf-16')
            except UnicodeDecodeError as exc: raise ValueError('CSV must use UTF-8 or UTF-16 encoding.') from exc
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        all_rows = list(csv.reader(io.StringIO(text), dialect))
        if not all_rows: raise ValueError('Data CSV is empty.')
        if any(len(row) > len(all_rows[0]) for row in all_rows[1:]): raise ValueError('CSV has rows wider than its header. Check quoting and delimiter.')
        return [{'name': 'submissions', 'headers': all_rows[0], 'rows': [row + [None] * (len(all_rows[0]) - len(row)) for row in all_rows[1:] if any(not missing(v) for v in row)]}]
    if not filename.lower().endswith('.xlsx'): raise ValueError('Data must be an .xlsx or .csv file.')
    validate_xlsx_archive(raw)
    book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheets = []
        if worksheet and worksheet not in [sheet.title for sheet in book.worksheets]:
            raise ValueError(f'The worksheet "{worksheet}" is not available in this workbook.')
        for sheet in ([book[worksheet]] if worksheet else book.worksheets):
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            if not header or (worksheet and not any(not missing(v) for v in header)): continue
            sheets.append({'name': sheet.title, 'headers': [str(v) if v is not None else '' for v in header], 'rows': [[scalar(v) for v in row] for row in rows if any(not missing(v) for v in row)]})
        if worksheet and (not sheets or not sheets[0]['rows']):
            raise ValueError(f'The worksheet "{worksheet}" is empty. Choose a sheet with headers and response rows.')
        if not sheets: raise ValueError('Data workbook contains no header rows.')
        return sheets
    finally: book.close()


def aliases(q: dict) -> set[str]:
    return {norm(v) for v in [q['id'], q['name'], *q['labels'].values()] if v}


def propose_mapping(form: dict, sheets: list[dict]) -> tuple[list[dict], list[str]]:
    mappings, warnings = [], []
    form_groups = {g['id']: g for g in form['groups']}
    def question_aliases(q):
        result = aliases(q)
        chain = []; gid = q['group']
        while gid:
            g = form_groups[gid]; chain.insert(0, g); gid = g['parent']
        if chain:
            for language in ['name', 'en', 'ar', 'default']:
                prefixes = [g['name'] if language == 'name' else label(g, language) for g in chain]
                for suffix in [q['name'], *q['labels'].values()]:
                    for separator in ['/', ' / ', '.']:
                        result.add(norm(separator.join([*prefixes, suffix])))
        return result
    note_headers = set()
    for note in form.get('notes', []):
        note_headers.update(question_aliases({**note, 'group': note.get('group', note['id'].rpartition('/')[0])}))
    for si, sheet in enumerate(sheets):
        index = defaultdict(list)
        for q in form['questions']:
            if q['repeat']:
                group = next(g for g in form['groups'] if g['id'] == q['repeat'])
                if not any(norm(sheet['name']) == a or norm(sheet['name']).endswith('/' + a) for a in aliases(group)): continue
            elif si != 0: continue
            prefixes = question_aliases(q)
            for prefix in prefixes: index[prefix].append({'question': q['id'], 'option': ''})
            if q['kind'] == 'multiple':
                for choice in q['choices']:
                    suffixes = aliases({'id': choice['name'], **choice})
                    for prefix in prefixes:
                        for separator in ['/', ' / ', '.']:
                            for suffix in suffixes: index[prefix + separator + suffix].append({'question': q['id'], 'option': choice['name']})
        for ci, header in enumerate(sheet['headers']):
            candidates = index.get(norm(header), [])
            candidates = list({(c['question'], c['option']): c for c in candidates}.values())
            metadata = header.startswith('_') or header.startswith('meta/') or norm(header) in note_headers
            mappings.append({'sheet': si, 'column': ci, 'header': header, 'question': candidates[0]['question'] if len(candidates) == 1 else '', 'option': candidates[0]['option'] if len(candidates) == 1 else '', 'candidates': candidates, 'status': 'matched' if len(candidates) == 1 else 'ignored' if metadata else 'ambiguous' if candidates else 'unmatched'})
    duplicates = Counter((m['sheet'], m['question'], m['option']) for m in mappings if m['question'])
    for m in mappings:
        if m['question'] and duplicates[(m['sheet'], m['question'], m['option'])] > 1: m['status'] = 'ambiguous'
    if any(g['repeat'] for g in form['groups']) and len(sheets) == 1: warnings.append('Repeat data may be absent. Kobo CSV exports do not include separate repeat sheets.')
    versions = set()
    for sheet in sheets:
        if '__version__' in sheet['headers']:
            ci = sheet['headers'].index('__version__'); versions.update(str(r[ci]) for r in sheet['rows'] if not missing(r[ci]))
    expected = str(form['settings'].get('version') or '')
    if versions and expected and versions != {expected}: warnings.append('Data includes form versions different from this XLSForm. Review unmatched columns and choices.')
    return mappings, warnings


def decode_choice(value: Any, q: dict, multiple=False) -> tuple[list[str], str | None]:
    if missing(value): return [], None
    value = str(value).strip()
    lookup = defaultdict(set)
    for choice in q['choices']:
        for a in aliases({'id': choice['name'], **choice}): lookup[a].add(choice['name'])
    if norm(value) in lookup:
        codes = lookup[norm(value)]
        return (list(codes), None) if len(codes) == 1 else ([], 'ambiguous choice label')
    if not multiple: return [], 'unknown choice value'
    # Kobo codes are whitespace-separated. Labels may contain whitespace and commas.
    tokens = value.split()
    codes = {c['name'] for c in q['choices']}
    if all(t in codes for t in tokens): return sorted(set(tokens)), None
    for delimiter in [';', ',']:
        parts = [p.strip() for p in value.split(delimiter) if p.strip()]
        if len(parts) > 1 and all(norm(p) in lookup and len(lookup[norm(p)]) == 1 for p in parts):
            return sorted({next(iter(lookup[norm(p)])) for p in parts}), None
    # Segment using the complete choice dictionary; reject ambiguous segmentations.
    text = norm(value)
    memo = {}
    def segment(offset):
        if offset in memo: return memo[offset]
        while offset < len(text) and text[offset] in ' ,;': offset += 1
        if offset == len(text): return [()]
        out = []
        for alias, names in lookup.items():
            end = offset + len(alias)
            if alias and text.startswith(alias, offset) and (end == len(text) or text[end] in ' ,;'):
                for tail in segment(end):
                    for name in names: out.append((name, *tail))
                    if len(out) > 1: break
            if len(out) > 1: break
        memo[offset] = out[:2]; return memo[offset]
    solutions = segment(0)
    if len(solutions) == 1: return sorted(set(solutions[0])), None
    return [], 'ambiguous combined choices' if solutions else 'unknown combined choices'


def filtered_choices(q: dict, values: dict, questions: dict) -> dict:
    """Interpret only simple XLSForm choice filters, never execute form code."""
    expression = q.get('choiceFilter', '')
    if not expression: return q
    expression = re.sub(r'\$\{([^}]+)\}', lambda m: '__q_' + m[1], expression)
    expression = re.sub(r'(?<![<>!=])=(?!=)', '==', expression)
    context = {'__q_' + question['name']: values.get(question['id']) for question in questions.values()}
    try:
        tree = ast.parse(expression, mode='eval')
        def evaluate(node, attributes):
            if isinstance(node, ast.Expression): return evaluate(node.body, attributes)
            if isinstance(node, ast.Constant): return node.value
            if isinstance(node, ast.Name): return context.get(node.id) if node.id.startswith('__q_') else attributes.get(node.id)
            if isinstance(node, ast.BoolOp):
                results = [evaluate(v, attributes) for v in node.values]
                return all(results) if isinstance(node.op, ast.And) else any(results)
            if isinstance(node, ast.Compare) and len(node.ops) == 1:
                left, right = evaluate(node.left, attributes), evaluate(node.comparators[0], attributes)
                if isinstance(node.ops[0], ast.Eq): return str(left) == str(right)
                if isinstance(node.ops[0], ast.NotEq): return str(left) != str(right)
            raise ValueError('Unsupported choice filter expression')
        choices = [choice for choice in q['choices'] if evaluate(tree, choice.get('attributes', {}))]
        return {**q, 'choices': choices}
    except (SyntaxError, ValueError, TypeError): return q


def normalize(form: dict, sheets: list[dict], mappings: list[dict], policies: dict | None = None, *, review: dict | None = None) -> tuple[list[dict], list[dict]]:
    from .kobo_import_links import validate_links, reconcile_values
    validate_links(form, sheets, mappings)
    policies = policies or {}
    questions = {q['id']: q for q in form['questions']}
    by_sheet = defaultdict(lambda: defaultdict(list))
    for m in mappings:
        if not m.get('question'): continue
        if m['question'] not in questions: raise ValueError('Mapping references an unknown question.')
        if not 0 <= m['sheet'] < len(sheets) or not 0 <= m['column'] < len(sheets[m['sheet']]['headers']): raise ValueError('Mapping column is outside its worksheet.')
        if m.get('option') and m['option'] not in {c['name'] for c in questions[m['question']]['choices']}: raise ValueError('Mapping references an unknown answer option.')
        by_sheet[m['sheet']][m['question']].append(m)
    issues = Counter()
    records = []
    parent_maps = {}
    for si, sheet in enumerate(sheets):
        hi = {v: i for i, v in enumerate(sheet['headers'])}
        for ri, row in enumerate(sheet['rows']):
            def get(key): return row[hi[key]] if key in hi and hi[key] < len(row) else None
            rid = f'{si}:{ri}'
            parent_key = str(get('_parent_index')) if get('_parent_index') is not None else ''
            parent_sheet = str(get('_parent_table_name') or '')
            parent = parent_maps.get((parent_sheet, parent_key)) or parent_maps.get(('', parent_key), '') if si else ''
            values = {}
            scope = next((g['id'] for g in form['groups'] if g['repeat'] and any(norm(sheet['name']) == a or norm(sheet['name']).endswith('/' + a) for a in aliases(g))), '') if si else ''
            for qid, ms in by_sheet[si].items():
                q = questions[qid]; scope = q['repeat'] or scope
                effective_question = filtered_choices(q, values, questions)
                policy = policies.get(qid) or ('split' if q['kind'] == 'multiple' and any(m.get('option') for m in ms) else 'combined')
                value = None if q['kind'] == 'multiple' and policy == 'split' else reconcile_values(effective_question, ms, row, issues, review, sheet['name'], ri + 2)
                if q['kind'] == 'multiple':
                    combined = value or []
                    split = set(); split_seen = False
                    for m in ms:
                        if not m.get('option'): continue
                        v = row[m['column']] if m['column'] < len(row) else None
                        if missing(v): continue
                        split_seen = True
                        if norm(str(v)) in {'1', '1.0', 'true', 'yes', 'selected'}: split.add(m['option'])
                        elif norm(str(v)) not in {'0', '0.0', 'false', 'no'}: issues[(qid, 'invalid split-option flag')] += 1
                    selected = split if policy == 'split' or (missing(value) and split_seen) else set(combined)
                    values[qid] = sorted(selected) or None
                else: values[qid] = value
            records.append({'id': rid, 'parent': parent, 'scope': scope, 'sheet': si, 'row': ri + 2, 'values': values})
            index = get('_index')
            if index is not None:
                parent_maps[(sheet['name'], str(index))] = rid
                if si == 0: parent_maps[('', str(index))] = rid
    for record in records:
        if record['scope'] and not record['parent']: issues[(record['scope'], 'repeat entry has no matching parent submission')] += 1
    return records, [{'question': qid, 'message': message, 'count': count} for (qid, message), count in issues.items()]


class ProjectStore:
    def __init__(self, path: Path | None = None):
        self.path = path or (Path(os.environ['INTERSOS_KOBO_DB']) if os.getenv('INTERSOS_KOBO_DB') else Path(os.getenv('LOCALAPPDATA', str(Path.home()))) / 'INTERSOS Legal Platform' / 'kobo-projects.sqlite3')

    def connect(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(self.path, timeout=30)
        db.execute('PRAGMA foreign_keys=ON')
        db.execute('CREATE TABLE IF NOT EXISTS projects (id TEXT PRIMARY KEY, name TEXT NOT NULL, updated TEXT NOT NULL, payload TEXT NOT NULL, settings TEXT NOT NULL, form BLOB NOT NULL, data BLOB NOT NULL)')
        return db

    def list(self):
        with self.connect() as db:
            return [{'id': r[0], 'name': r[1], 'updated': r[2]} for r in db.execute('SELECT id,name,updated FROM projects ORDER BY updated DESC')]

    def create(self, name, payload, form_raw, data_raw, settings: KoboSettings | None = None):
        pid = uuid.uuid4().hex
        settings = (settings or new_project_settings(payload['form'])).model_dump()
        with self.connect() as db:
            db.execute('INSERT INTO projects VALUES (?,?,?,?,?,?,?)', (pid, name.strip() or 'Untitled Kobo project', datetime.now().isoformat(), json.dumps(payload, ensure_ascii=False), json.dumps(settings), form_raw, data_raw))
        return self.get(pid)

    def get(self, pid):
        with self.connect() as db:
            row = db.execute('SELECT name,payload,settings,updated FROM projects WHERE id=?', (pid,)).fetchone()
        if not row: raise KeyError('Project not found.')
        return {'id': pid, 'name': row[0], **json.loads(row[1]), 'settings': json.loads(row[2]), 'updated': row[3]}

    def save(self, pid, settings: KoboSettings):
        with self.connect() as db:
            cursor = db.execute('UPDATE projects SET settings=?,updated=? WHERE id=?', (settings.model_dump_json(), datetime.now().isoformat(), pid))
            if not cursor.rowcount: raise KeyError('Project not found.')

    def rename(self, pid, name):
        name = name.strip()
        if not name or len(name) > 200:
            raise ValueError('Enter a project name between 1 and 200 characters.')
        updated = datetime.now().isoformat()
        with self.connect() as db:
            cursor = db.execute('UPDATE projects SET name=?,updated=? WHERE id=?', (name, updated, pid))
            if not cursor.rowcount: raise KeyError('Project not found.')
        return {'id': pid, 'name': name, 'updated': updated}

    def sources(self, pid):
        with self.connect() as db:
            row = db.execute('SELECT form,data FROM projects WHERE id=?', (pid,)).fetchone()
        if not row: raise KeyError('Project not found.')
        return bytes(row[0]), bytes(row[1])

    def repair(self, pid, payload, form_raw, data_raw):
        # Keep the project identity, latest display settings until the reviewed sources are applied.
        with self.connect() as db:
            row = db.execute('SELECT settings,payload FROM projects WHERE id=?', (pid,)).fetchone()
            if not row: raise KeyError('Project not found.')
            settings = json.loads(row[0])
            known = {q['id'] for q in payload['form']['questions']}
            settings['hidden'] = list(set(settings.get('hidden', [])) | {q['id'] for q in payload['form']['questions'] if q['identifying'] and q['id'] not in {old['id'] for old in json.loads(row[1])['form']['questions']}})
            settings['hidden'] = [qid for qid in settings['hidden'] if qid in known]
            settings['filterFields'] = [qid for qid in settings.get('filterFields', []) if qid in known]
            settings['filters'] = [f for f in settings.get('filters', []) if f['field'] in known]
            settings['questions'] = {qid: {**cfg, 'breakdown': cfg.get('breakdown', '') if cfg.get('breakdown', '') in known else '', 'breakdowns': [field for field in cfg.get('breakdowns', []) if field in known]} for qid, cfg in settings.get('questions', {}).items() if qid in known}
            from .kobo_indicators import indicator_issues
            for definition in settings.get('indicators', []):
                item = IndicatorDefinition.model_validate({**definition, 'reviewIssues': []})
                definition['reviewIssues'] = indicator_issues(payload, item)
            from .kobo_saved_views import saved_view_issues
            for view in settings.get('savedViews', []):
                if view.get('questionKinds'):
                    added = known - view['questionKinds'].keys()
                    view['snapshot']['hidden'] = list(dict.fromkeys([*view['snapshot'].get('hidden', []), *sorted(added)]))
                    view['questionKinds'].update({q['id']: q['kind'] for q in payload['form']['questions'] if q['id'] in added})
                view['reviewIssues'] = saved_view_issues(payload, SavedAnalysisView.model_validate(view))
            if any(view['id'] == settings.get('activeViewId') and view['reviewIssues'] for view in settings.get('savedViews', [])): settings['activeViewId'] = ''
            db.execute('UPDATE projects SET payload=?,form=?,data=?,settings=?,updated=? WHERE id=?', (json.dumps(payload, ensure_ascii=False), form_raw, data_raw, json.dumps(settings), datetime.now().isoformat(), pid))
        return self.get(pid)

    def delete(self, pid):
        with self.connect() as db:
            db.execute('DELETE FROM projects WHERE id=?', (pid,))


def matches(value, rule: KoboFilter):
    if rule.kind == 'category':
        if not rule.values: return True
        return any(v in rule.values for v in (value if isinstance(value, list) else [str(value) if value is not None else '__blank__']))
    if missing(value): return False
    if rule.kind == 'date': value = str(value)[:10]
    try:
        low = float(rule.minimum) if rule.kind == 'number' and rule.minimum is not None else rule.minimum
        high = float(rule.maximum) if rule.kind == 'number' and rule.maximum is not None else rule.maximum
        return (low is None or value >= low) and (high is None or value <= high)
    except (ValueError, TypeError): return False


def filtered_records(project: dict, settings: KoboSettings):
    from .kobo_trends import validate_chart_settings
    validate_chart_settings({q['id']: q for q in project['form']['questions']}, settings)
    records = project['records']; lookup = {r['id']: r for r in records}
    questions = {q['id']: q for q in project['form']['questions']}
    for rule in settings.filters:
        if rule.field not in questions: raise ValueError('Filter references an unknown question.')
    def passes(record, rules):
        return all(matches(record['values'].get(f.field), f) for f in rules)
    main_rules = [f for f in settings.filters if not questions[f.field]['repeat']]
    parents = {r['id'] for r in records if not r['scope'] and passes(r, main_rules)}
    repeat_rules = [f for f in settings.filters if questions[f.field]['repeat']]
    for scope in {questions[f.field]['repeat'] for f in repeat_rules}:
        scope_rules = [f for f in repeat_rules if questions[f.field]['repeat'] == scope]
        matching_parents = set()
        for r in records:
            if r['scope'] != scope or not passes(r, scope_rules): continue
            parent = r
            while parent.get('parent') in lookup: parent = lookup[parent['parent']]
            matching_parents.add(parent['id'])
        parents &= matching_parents
    out = []
    for r in records:
        if not r['scope']:
            if r['id'] in parents: out.append(r)
            continue
        root = r
        while root.get('parent') in lookup: root = lookup[root['parent']]
        if root['id'] in parents and passes(r, [f for f in repeat_rules if questions[f.field]['repeat'] == r['scope']]): out.append(r)
    return out


def breakdown_fields(cfg, questions, target):
    ids = cfg.breakdowns or ([cfg.breakdown] if cfg.breakdown else [])
    if len(set(ids)) != len(ids): raise ValueError('Choose a different question for each breakdown.')
    fields = []
    for qid in ids:
        field = questions.get(qid)
        if not field: raise ValueError('Breakdown references an unknown question.')
        if field['kind'] not in {'category', 'text'} or (field['repeat'] and field['repeat'] != target['repeat']):
            raise ValueError('Breakdown must be a category/text question in this section or its parent submission.')
        fields.append(field)
    return fields


def breakdown_values(record, dimensions, lookup):
    values = []
    for dimension in dimensions:
        current = record
        while dimension['id'] not in current['values'] and current.get('parent') in lookup:
            current = lookup[current['parent']]
        value = current['values'].get(dimension['id'])
        values.append('__blank__' if missing(value) else str(value))
    return tuple(values)


CHART_PALETTES = {
    'Solid blue': ['#1683D8'],
    'Ocean': ['#1E3A5F', '#247BA0', '#329D9C', '#5E60CE', '#427AA1', '#6B8E9B'],
    'Botanical': ['#245C4A', '#4E7C59', '#849B55', '#B18A45', '#8F5D5D', '#5D7285'],
    'Jewel': ['#3155A6', '#7851A9', '#B54B73', '#C77B32', '#218577', '#506680'],
    'Contrast': ['#0072B2', '#D55E00', '#009E73', '#CC79A7', '#7A5A00', '#555555'],

    'Balanced': ['#1877c9','#2f9e68','#d97706','#7c5cc4','#d9485f','#168a91','#9a6b35','#536d88','#b85c9e','#6d8f38','#e06b3c','#3f73a8'],
    'Blue': ['#174A7E','#2166AC','#3787C0','#65A9CF','#9CCAE1','#C6E2EF'],
    'Teal': ['#145A58','#207D78','#339D94','#62B6A8','#95CFC0','#C4E5D8'],
    'Warm': ['#8C3B2A','#B45232','#D57641','#E6A15B','#EEC486','#F3DEB7'],
}


def palette_color(palette, option):
    colors = CHART_PALETTES[palette]
    value = 0
    for char in option: value = (value * 31 + ord(char)) & 0xffffffff
    return colors[0 if option == 'series' else value % len(colors)]


def question_categories(q: dict, answered: list[dict], cfg):
    denominator = len(answered)
    values = [r['values'][q['id']] for r in answered]
    stats, bins = {}, []
    if q['kind'] == 'number' and values:
        stats = {'count': denominator, 'mean': statistics.mean(values), 'median': statistics.median(values), 'min': min(values), 'max': max(values), 'standardDeviation': statistics.pstdev(values)}
        unique = sorted(set(values))
        if (cfg.chart == 'histogram' or cfg.numericMode == 'bins' or (cfg.numericMode == 'auto' and len(unique) > 12)) and min(values) != max(values):
            low, high = min(values), max(values); step = (high - low) / cfg.bins
            bins = [{'id': f'bin:{i}', 'label': f'{low + i * step:g} - {low + (i + 1) * step:g}' + (' (inclusive)' if i == cfg.bins - 1 else ' (upper excluded)'), 'min': low + i * step, 'max': low + (i + 1) * step} for i in range(cfg.bins)]
    def categories(value):
        if missing(value): return []
        if q['kind'] == 'multiple': return value
        if q['kind'] == 'number' and bins:
            i = min(int((value - bins[0]['min']) / (bins[-1]['max'] - bins[0]['min']) * len(bins)), len(bins) - 1)
            return [bins[i]['id']]
        if q['kind'] == 'date': return [str(value)[:{'day': 10, 'month': 7, 'year': 4}[cfg.dateUnit]]]
        return [str(value)]
    return stats, bins, categories


def question_result(q: dict, records: list[dict], settings: KoboSettings, questions: dict, lookup: dict):
    cfg = settings.questions.get(q['id'], QuestionSettings())
    cohort = [r for r in records if r['scope'] == q['repeat']]
    answered = [r for r in cohort if not missing(r['values'].get(q['id']))]
    denominator = len(answered)
    stats, bins, categories = question_categories(q, answered, cfg)
    counts = Counter(c for record in answered for c in categories(record['values'][q['id']]))
    answer_labels = {c['name']: label(c, settings.language) for c in q['choices']}
    answer_labels.update({b['id']: b['label'] for b in bins})
    universe = list(answer_labels) if q['choices'] or bins else sorted(counts, key=lambda v: float(v) if q['kind'] == 'number' else v)
    rows = [{'id': c, 'label': answer_labels.get(c, c), 'count': counts[c], 'percent': counts[c] / denominator * 100 if denominator else None} for c in universe]
    breakdowns = []
    dimensions = breakdown_fields(cfg, questions, q)
    if dimensions:
        dimension_labels = [{c['name']: label(c, settings.language) for c in dim['choices']} for dim in dimensions]
        cohorts = defaultdict(list)
        for record in answered: cohorts[breakdown_values(record, dimensions, lookup)].append(record)
        for values, members in sorted(cohorts.items()):
            captions = [dimension_labels[i].get(value, value) if value != '__blank__' else 'Missing breakdown' for i, value in enumerate(values)]
            bid = values[0] if len(values) == 1 else json.dumps(values, ensure_ascii=False)
            breakdowns.append({'id': bid, 'label': ' / '.join(captions), 'denominator': len(members), 'values': list(values), 'labels': captions})
        for row in rows:
            row['series'] = []
            for b in breakdowns:
                members = cohorts[tuple(b['values'])]
                n = sum(row['id'] in categories(r['values'][q['id']]) for r in members)
                row['series'].append({'id': b['id'], 'count': n, 'percent': n / b['denominator'] * 100 if b['denominator'] else None})
    apply_display_percentages(rows, breakdowns, denominator, q['kind'], settings.percentageDecimals)
    if settings.hideZero and cfg.chart != 'histogram': rows = [r for r in rows if r['count']]
    if cfg.answerSort != 'form' and cfg.chart != 'histogram':
        by_count = cfg.answerSort.startswith('count')
        rows.sort(key=lambda r: (r['count'], r['label']) if by_count else r['label'].casefold(), reverse=cfg.answerSort.endswith('desc'))
    if not settings.hideBlanks and len(cohort) > denominator: rows.append({'id': '__blank__', 'label': 'Blank / unanswered' if settings.language == 'en' else 'فارغ / دون إجابة', 'count': len(cohort) - denominator, 'percent': None, 'displayPercent': None})
    colors = {key: palette_color(cfg.palette, key) for key in [*[row['id'] for row in rows], *[column['id'] for column in breakdowns], 'series']} if cfg.palette else {}
    colors.update(cfg.colors)
    from .kobo_trends import answer_trend
    trend = answer_trend(q, cfg, questions, answered, lookup, categories, rows)
    return {'trend': trend, 'id': q['id'], 'name': q['name'], 'label': label(q, settings.language), 'group': q['group'], 'kind': q['kind'], 'repeat': q['repeat'], 'denominator': denominator, 'missing': len(cohort) - denominator, 'total': len(cohort), 'rows': rows, 'statistics': stats, 'bins': bins, 'breakdowns': breakdowns, 'breakdownDisplay': cfg.breakdownDisplay, 'breakdownFields': [{'id': dim['id'], 'label': label(dim, settings.language)} for dim in dimensions], 'percentageBasis': 'Answered respondents within each breakdown category' if breakdowns else 'Respondents who answered this question', 'percentageDecimals': settings.percentageDecimals, 'chart': cfg.chart or ('line' if q['kind'] == 'date' else 'column'), 'colors': colors, 'pointColors': cfg.pointColors, 'appearance': cfg.appearance.model_dump()}


def analyze(project: dict, settings: KoboSettings, question: str = ''):
    records = filtered_records(project, settings)
    questions = {q['id']: q for q in project['form']['questions']}
    lookup = {r['id']: r for r in project['records']}
    results = []
    for q in questions.values():
        if question and q['id'] != question: continue
        if q['id'] in settings.hidden or (q['metadata'] and not settings.metadata): continue
        result = question_result(q, records, settings, questions, lookup)
        if not settings.hideEmpty or result['denominator']: results.append(result)
    if settings.questionSort != 'form':
        by_count = settings.questionSort.startswith('answered')
        # Sorting only within groups preserves the original group hierarchy.
        groups = defaultdict(list)
        for result in results: groups[result['group']].append(result)
        for group in groups.values(): group.sort(key=lambda r: (r['denominator'], r['name']) if by_count else r['name'].casefold(), reverse=settings.questionSort.endswith('desc'))
        iterators = {g: iter(rs) for g, rs in groups.items()}
        results = [next(iterators[r['group']]) for r in results]
    filter_summary = []
    for rule in settings.filters:
        q = questions[rule.field]
        choices = {c['name']: label(c, settings.language) for c in q['choices']}
        value = ', '.join('Blank / unanswered' if v == '__blank__' else choices.get(v, v) for v in rule.values) if rule.kind == 'category' else f'{rule.minimum if rule.minimum is not None else "Any"} - {rule.maximum if rule.maximum is not None else "Any"}'
        filter_summary.append(f'{label(q, settings.language)}: {value}')
    return {'total': sum(not r['scope'] for r in project['records']), 'filtered': sum(not r['scope'] for r in records), 'questions': results, 'filterSummary': filter_summary, 'groups': [{**g, 'label': label(g, settings.language)} for g in project['form']['groups']], 'settings': settings.model_dump()}
