"""Export the original rows referenced by an indicator, grouped by source."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .excel_export import Font, PatternFill, append_frame
from .legal_platform import format_excel_dates, _safe_export


SOURCE_SHEETS = {
    "assessments": "Assessments",
    "legalservices": "Legal Services",
    "deportationrecords": "Deportation Records",
    "awareness": "Awareness",
    "beneficiaries": "Beneficiaries",
}


def build_source_record_workbook(frames):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for dataset, title in SOURCE_SHEETS.items():
        if dataset not in frames:
            continue
        frame = frames[dataset]
        frame = frame[[column for column in frame.columns if not str(column).startswith("__")]]
        sheet = workbook.create_sheet(title)
        append_frame(sheet, _safe_export(frame))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for index, column in enumerate(frame.columns, 1):
            width = max([len(str(column)), *[len(str(value)) for value in frame[column].head(250)]]) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(12, width))
    format_excel_dates(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
