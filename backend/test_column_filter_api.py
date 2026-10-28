from io import BytesIO

from openpyxl import load_workbook

from backend.test_kobo import sample
from backend.test_kobo_preview import preview_client
from backend.test_page_filter_audit import store


def test_kobo_filters_values_and_exports_use_the_complete_cohort(preview_client):
    client,store,pid=preview_client
    project=store.get(pid)
    query={'settings':project['settings'],'columns':['group/needs','group/age'],'pageSize':1,'columnFilters':[{'column':'group/age','type':'number','condition':{'operator':'gte','value':'20'}}]}
    response=client.post(f'/api/kobo/projects/{pid}/data',json=query)
    assert response.status_code==200,response.text
    result=response.json()
    assert result['matched']==1 and float(result['rows'][0]['values']['group/age'])==20
    values=client.post(f'/api/kobo/projects/{pid}/data/values',json={**query,'column':'group/age'}).json()
    assert values['total']==3 and values['type']=='number'
    assert '' in values['values'] and any(value and float(value)==0 for value in values['values'])
    export=client.post(f'/api/kobo/projects/{pid}/data/export',json=query)
    assert export.status_code==200,export.text
    rows=list(load_workbook(BytesIO(export.content)).active.values)
    assert len(rows)==2 and float(rows[1][-1])==20
    zero=next(record for record in project['records'] if record['values'].get('group/age')==0)
    selected=client.post(f'/api/kobo/projects/{pid}/data/export',json={**query,'selectedIds':[zero['id']]})
    rows=list(load_workbook(BytesIO(selected.content)).active.values)
    assert len(rows)==2 and float(rows[1][-1])==0
    assert store.get(pid)==project


def test_empty_selections_and_invalid_columns_are_explicit(preview_client):
    client,store,pid=preview_client
    query={'settings':store.get(pid)['settings'],'columns':['group/age']}
    response=client.post(f'/api/kobo/projects/{pid}/data',json={**query,'columnFilters':[{'column':'group/age','values':[]}]})
    assert response.status_code==200 and response.json()['matched']==0
    response=client.post(f'/api/kobo/projects/{pid}/data',json={**query,'columnFilters':[{'column':'unknown','values':[]}]})
    assert response.status_code==400


def test_quality_and_answer_record_filters_keep_summary_calculations(preview_client):
    client,store,pid=preview_client
    settings=store.get(pid)['settings']
    quality=client.post(f'/api/kobo/projects/{pid}/quality-review',json={'settings':settings}).json()
    check=quality['checks'][0]['id']
    filtered=client.post(f'/api/kobo/projects/{pid}/quality-review',json={'settings':settings,'check':check,'columnFilters':[{'column':'row','values':[]}]}).json()
    assert filtered['matched']==0 and filtered['rows']==[]
    assert filtered['checks']==quality['checks']
    query={'settings':settings,'question':'group/age','columnFilters':[{'column':'row','type':'number','condition':{'operator':'gte','value':'3'}}]}
    answers=client.post(f'/api/kobo/projects/{pid}/responses',json=query)
    assert answers.status_code==200,answers.text
    assert all(row['row']>=3 for row in answers.json()['rows'])
    values=client.post(f'/api/kobo/projects/{pid}/responses/values',json={**query,'column':'row'}).json()
    assert values['total']>=answers.json()['total']


def test_legal_header_facets_cases_and_dashboard_rows(store,monkeypatch):
    from backend import main
    monkeypatch.setattr(main,'require_legal_store',lambda:store)
    rule={'column':'Assessment ID','values':['1']}
    result=main.legal_explorer(main.LegalQuery(dataset='assessments',columnFilters=[rule]))
    assert result['total']==1 and result['rows'][0]['Assessment ID']=='1'
    facets=main.legal_record_values(main.RecordValuesQuery(dataset='assessments',column='Assessment ID',columnFilters=[rule]))
    assert facets['values']==['1','2','3'] and facets['type']=='text'
    baseline=main.legal_analytics_dashboard(main.LegalAnalyticsDashboardRequest(dataset='assessments'))
    filtered=main.legal_analytics_dashboard(main.LegalAnalyticsDashboardRequest(dataset='assessments',columnFilters=[rule]))
    assert filtered['matchedRows']==1 and filtered['kpis']==baseline['kpis'] and filtered['charts']==baseline['charts']
    detention_base=store.detention_cases()
    column=detention_base['columns'][0];selected=detention_base['rows'][0][column]
    detention=main.legal_detention(main.LegalQuery(columnFilters=[{'column':column,'values':[selected]}]))
    assert detention['total']==1 and detention['rows'][0][column]==selected
    cases=store.case('',view_mode='table',column_filters=[{'column':'beneficiaries::Case ID','values':['1']}])
    assert cases['totalCases']==1 and cases['cases'][0]['beneficiary']['Case ID']=='1'


def test_legal_date_facets_follow_conditions_but_ignore_their_own_rule(store):
    rules=[{'column':'Assessment ID','values':['1']},{'column':'Project','values':['UNHCR 2026 - AMAL CAMP']}]
    options={column['name']:column['values'] for column in store.explorer_filters('assessments',column_filters=rules)['columns']}
    assert options['Assessment ID']==['1']
    assert options['Project']==['UNHCR 2026 - AMAL CAMP']
