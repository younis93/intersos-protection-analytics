from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.filter_selection import EXCLUDE_PREFIX
from backend.legal_platform import LegalStore, REGISTERED_RULES
from backend.test_page_filter_audit import store


@pytest.fixture
def checkbox_store(store, monkeypatch):
    local=LegalStore(store.frames,"synthetic review checkbox filters",[])
    monkeypatch.setattr(local,"_name_match_flags",lambda *args,**kwargs:[])
    monkeypatch.setattr(local,"_contact_name_match_flags",lambda *args,**kwargs:[])
    monkeypatch.setattr(local,"_assessment_month_flags",lambda *args,**kwargs:([],"",[]))
    monkeypatch.setattr(local,"_service_month_flags",lambda *args,**kwargs:([],"",[]))
    fields={"beneficiaries":"identificationDate","assessments":"assessmentDate","legalservices":"serviceDate","awareness":"awarenessDate"}
    for dataset,date_field in fields.items():
        rule=next(rule for rule in REGISTERED_RULES[dataset] if rule not in {"Possible duplicate name","Possible duplicate contact and name"})
        local.flags[dataset]=[{"row":index+2,"rule":rule,"name":f"Synthetic person {index}","project":project,"location":location,"severity":"High",date_field:date}
                              for index,(project,location,date) in enumerate(zip(["A","B","C"],["One","Two","Three"],["01/01/2026","02/04/2026","03/07/2026"]))]
    return local


@pytest.mark.parametrize("dataset",["beneficiaries","assessments","legalservices","awareness"])
@pytest.mark.parametrize("field,values",[("project",["A","B","C"]),("location",["One","Two","Three"]),("date",["2026-01","2026-04","2026-07"])])
def test_review_checkbox_selection_and_excel_agree(checkbox_store,dataset,field,values):
    local=checkbox_store
    rule=local.flags[dataset][0]["rule"]
    assert local.review(dataset,rule=rule)["total"]==3
    for selected,expected in [(values,3),(values[:2],2),([values[0]],1),([EXCLUDE_PREFIX+values[0]],2),(["missing"],0),([],3)]:
        filters={field:selected}
        result=local.review(dataset,rule=rule,filters=filters)
        assert result["total"]==expected
        assert result["ruleCounts"][rule]==expected
        assert set(result["filterOptions"][field])==set(values),"A field must ignore its own selection when offering options"
        book=load_workbook(BytesIO(local.review_export(dataset,selected_rules=[rule],filters=filters)))
        # Each nonempty finding sheet has one header row; an empty export has a message sheet.
        if expected: assert sum(sheet.max_row-1 for sheet in book.worksheets)==expected
        else: assert result["rows"]==[]


@pytest.mark.parametrize("dataset",["beneficiaries","assessments","legalservices","awareness"])
def test_review_checkbox_facets_link_in_both_directions(checkbox_store,dataset):
    local=checkbox_store
    result=local.review(dataset,filters={"project":["A","B"]})
    assert set(result["filterOptions"]["location"])=={"One","Two"}
    result=local.review(dataset,filters={"location":["Two"]})
    assert result["filterOptions"]["project"]==["B"]
    result=local.review(dataset,filters={"project":["A","B"],"date":["2026-04"]})
    assert result["total"]==1
    assert result["rows"][0]["project"]=="B"
