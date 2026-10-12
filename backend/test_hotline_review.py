from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from backend.duplicate_exclusions import DuplicateExclusionRegistry
from backend.hotline import dashboard
from backend.hotline_review import RULES, COMBINED_NAME_RULE
from backend.legal_platform import LegalStore


def store_for(rows, exclusions=None):
    frame = pd.DataFrame(rows)
    return LegalStore.from_files({"legalhotlines": frame.to_csv(index=False).encode("utf-8-sig")}, "hotline review test", exclusions=exclusions or [])


def row(identifier, name="", **changes):
    return {"Hotline ID": identifier, "Name of the detained person اسم الشخص المعتقل": name,
            "Name of the caller اسم المتصل": "Different Caller", "Contact Number": "7501234567",
            "Contact Date تاريخ الاتصال": "2/1/2026", "Is the beneficiary detained هل المستفيد موقوف": "",
            "Has the beneficiary referred to the helpline? هل تم احالة المستفيد الى خط المساعدة": "",
            "Priority الاولوية": "Low", "Refer to Lawyer": "", **changes}


def review(store, **kwargs):
    return store.review("legalhotlines", allow_name_variations=True, **kwargs)


def test_all_source_field_checkbox_filters_and_linked_options():
    store=store_for([
        row("1", "Same Person", **{"Gender":"Male","Detaining Authority":"Police","Custom field":"A"}),
        row("2", "Same Person", **{"Gender":"Female","Detaining Authority":"Court","Custom field":"B"}),
        row("3", "Other Person", **{"Gender":"","Detaining Authority":"Police","Custom field":"C","Contact Number":"123"}),
    ])
    initial=review(store)
    assert set(initial["fieldFilterOptions"]) == set(store.frames["legalhotlines"].columns)-{"Name of the caller اسم المتصل","Name of the detained person اسم الشخص المعتقل"}
    assert initial["fieldFilterOptions"]["Gender"] == ["Female","Male","Not recorded"]
    matched=review(store,filters={"Gender":["Male","Female"],"Custom field":["B"]})
    assert {item["hotlineId"] for item in matched["rows"]} == {"id:2"}
    assert matched["ruleCounts"][COMBINED_NAME_RULE] == 1
    assert matched["fieldFilterOptions"]["Custom field"] == ["A","B"]
    assert matched["fieldFilterOptions"]["Gender"] == ["Female"]
    blanks=review(store,filters={"Gender":["Not recorded"]})
    assert {item["hotlineId"] for item in blanks["rows"]} == {"id:3"}
    assert review(store,filters={"Custom field":["missing"]})["total"] == 0
    assert review(store,filters={"Custom field":["\x00exclude:A","\x00exclude:B"]})["total"] == blanks["total"]
    assert review(store)["total"] == initial["total"]


def test_checkbox_filters_apply_to_excel_export():
    store=store_for([
        row("1", "Same Person", **{"Custom field":"A"}),
        row("2", "Same Person", **{"Custom field":"B"}),
    ])
    workbook=load_workbook(BytesIO(store.review_export("legalhotlines",selected_rules=[COMBINED_NAME_RULE],filters={"Custom field":["B"]})))
    sheet=workbook.active
    records=list(sheet.values)
    # One header and the selected source record, while duplication was detected
    # using both records before the display filter was applied.
    assert len(records) == 2
    assert "B" in records[1]
    assert "A" not in records[1]


CONTACT_DATE="Contact Date تاريخ الاتصال"


def test_month_options_group_dates_and_hide_only_requested_filter_fields():
    source=[row("1", "Same Person", **{CONTACT_DATE:"1/5/2026","Caller himself or someone else المتصل نفسه":"Someone else"}),
            row("2", "Same Person", **{CONTACT_DATE:"2026-01-25"}),
            row("3", **{CONTACT_DATE:"2026-02-10"}),
            row("4", **{CONTACT_DATE:"invalid"}),row("5", **{CONTACT_DATE:""})]
    store=store_for(source)
    result=review(store)
    assert result["fieldFilterOptions"][CONTACT_DATE] == ["2026-02","2026-01","Not recorded"]
    assert not any("Name of the caller" in field or "Name of the detained person" in field or "Caller himself" in field for field in result["fieldFilterOptions"])
    assert result["rows"][0]["name"] == "Same Person"
    assert result["rows"][0]["callerName"] == "Different Caller"
    assert result["rows"][0]["contactDate"] == "1/5/2026"
    assert "Caller himself or someone else المتصل نفسه" in result["rows"][0]["sourceFields"]


def cross_month_rows():
    return [row("1", "Same Person", **{CONTACT_DATE:"2026-01-10","Project":"A"}),
            row("2", "Same Person", **{CONTACT_DATE:"2025-12-10","Project":"A"}),
            row("3", "Same Person", **{CONTACT_DATE:"2026-02-10","Project":"B"}),
            row("4", "Unrelated person", **{CONTACT_DATE:"2025-12-15","Priority الاولوية":"High عالية"}),
            row("5", "January individual", **{CONTACT_DATE:"2026-01-20","Priority الاولوية":"High عالية"}),
            row("6", "Other Group", **{CONTACT_DATE:"2025-12-20"}),
            row("7", "Other Group", **{CONTACT_DATE:"2026-02-20"})]


def test_month_filter_expands_only_seeded_duplicate_groups_and_matches_export():
    store=store_for(cross_month_rows())
    filters={CONTACT_DATE:["2026-01"]}
    result=review(store,filters=filters)
    assert {flag["hotlineId"] for flag in result["rows"]} == {"id:1","id:2","id:3","id:5"}
    assert result["ruleCounts"][COMBINED_NAME_RULE] == 3
    assert result["ruleCounts"][RULES[5]] == 1
    assert result["fieldFilterOptions"][CONTACT_DATE] == ["2026-02","2026-01","2025-12"]
    selected=review(store,rule=COMBINED_NAME_RULE,filters=filters,page_size=1,page=2)
    assert selected["total"] == 3 and len(selected["rows"]) == 1
    for flag in result["rows"]:
        if flag["rule"]==COMBINED_NAME_RULE:
            assert flag["duplicateContext"]["members"] == ["id:1","id:2","id:3"]
    workbook=load_workbook(BytesIO(store.review_export("legalhotlines",selected_rules=[COMBINED_NAME_RULE],filters=filters)))
    exported=list(workbook.active.values)
    assert len(exported) == 4
    assert {record[exported[0].index("Contact date")].strftime("%Y-%m-%d") for record in exported[1:]} == {"2026-01-10","2025-12-10","2026-02-10"}
    assert review(store,filters={CONTACT_DATE:["2030-01"]})["total"] == 0
    multi=review(store,filters={CONTACT_DATE:["2026-01","2025-12"]})
    assert {flag["hotlineId"] for flag in multi["rows"]} == {"id:1","id:2","id:3","id:4","id:5","id:6","id:7"}


def test_cross_month_peers_respect_other_filters_search_and_exclusions():
    store=store_for(cross_month_rows())
    filters={CONTACT_DATE:["2026-01"],"Project":["A"]}
    result=review(store,filters=filters)
    assert {flag["hotlineId"] for flag in result["rows"]} == {"id:1","id:2"}
    assert {flag["hotlineId"] for flag in review(store,filters={CONTACT_DATE:["2026-01"]},search="2026-01-10")["rows"]} == {"id:1"}
    excluded=[{"dataset":"legalhotlines","rule":COMBINED_NAME_RULE,"identifierType":"hotlineId","identifierValue":"id:2"}]
    restricted=store_for(cross_month_rows(),excluded)
    assert {flag["hotlineId"] for flag in review(restricted,rule=COMBINED_NAME_RULE,filters={CONTACT_DATE:["2026-01"]})["rows"]} == {"id:1","id:3"}


def test_not_recorded_month_includes_missing_and_invalid_dates():
    store=store_for([row("1", **{CONTACT_DATE:"","Contact Number":"123"}),row("2", **{CONTACT_DATE:"invalid","Contact Number":"123"}),row("3", **{CONTACT_DATE:"2026-01-01","Contact Number":"123"})])
    result=review(store,filters={CONTACT_DATE:["Not recorded"]})
    assert {flag["hotlineId"] for flag in result["rows"]} == {"id:1","id:2"}


def test_exact_and_similar_name_tables_expand_across_months_without_changing_placement():
    store=store_for([
        row("1","Exact person",**{CONTACT_DATE:"2026-01-01","Name of the caller اسم المتصل":"Caller A"}),
        row("2","Exact person",**{CONTACT_DATE:"2025-12-01","Name of the caller اسم المتصل":"Caller B"}),
        row("3","Ahmed Mohammed Hassan",**{CONTACT_DATE:"2026-01-01","Name of the caller اسم المتصل":"Caller C"}),
        row("4","Ahmed Mohammad Hassan",**{CONTACT_DATE:"2026-02-01","Name of the caller اسم المتصل":"Caller D"}),
    ])
    result=review(store,filters={CONTACT_DATE:["2026-01"]})
    assert {flag["hotlineId"] for flag in result["rows"] if flag["rule"]==RULES[0]} == {"id:1","id:2"}
    assert {flag["hotlineId"] for flag in result["rows"] if flag["rule"]==RULES[1]} == {"id:3","id:4"}
    assert result["ruleCounts"][COMBINED_NAME_RULE] == 0


def test_review_rules_and_hotline_only_source_preservation():
    rows = [row("1", "Ahmed Mohammed Hassan"), row("2", "AHMED  MOHAMMED HASSAN"),
            row("3", "Ahmed Mohammad Hassan"),
            row("4", "Ali Ahmed", **{"Name of the caller اسم المتصل": "Ali"}),
            row("5", **{"Is the beneficiary detained هل المستفيد موقوف": "Yes نعم"}),
            row("6", **{"Is the beneficiary detained هل المستفيد موقوف": "No لا", "Has the beneficiary referred to the helpline? هل تم احالة المستفيد الى خط المساعدة": "No لا"}),
            row("7", **{"Priority الاولوية": "High عالية"}), row("8", **{"Contact Number": "123"})]
    store = store_for(rows)
    source = store.frames["legalhotlines"].copy(deep=True)
    result = review(store)
    assert result["ruleCounts"] == dict(zip(RULES, [0, 1, 1, 1, 1, 1, 1, 2]))
    assert not result["unavailableRules"]
    assert all(flag["caseId"] == "" for flag in result["rows"])
    assert {flag["severity"] for flag in result["rows"] if flag["rule"] in RULES[3:6]} == {"High"}
    assert len({flag["duplicateGroup"] for flag in review(store, rule=COMBINED_NAME_RULE)["rows"]}) == 1
    assert review(store, date="2026-02")["total"] == result["total"]
    assert review(store, date="2026-01")["total"] == 0
    pd.testing.assert_frame_equal(store.frames["legalhotlines"], source)
    assert dashboard(store.frames["legalhotlines"])["total"] == len(rows)


def test_caller_match_complete_words_and_not_across_records():
    store = store_for([row("1", "Ali Ahmed", **{"Name of the caller اسم المتصل": "Ali"}),
                       row("2", "Alison", **{"Name of the caller اسم المتصل": "Ali"}),
                       row("3", "Omar Hassan", **{"Name of the caller اسم المتصل": "Ahmed"}),
                       row("4", "", **{"Name of the caller اسم المتصل": ""}),
                       row("5", "محمد أحمد", **{"Name of the caller اسم المتصل": "أحمد"})])
    assert {flag["hotlineId"] for flag in review(store, rule=RULES[2])["rows"]} == {"id:1", "id:5"}


def test_combined_name_check_is_first_and_name_tables_are_mutually_exclusive():
    store = store_for([
        row("1", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Ahmed Mohammed", "Is the beneficiary detained هل المستفيد موقوف": "Yes"}),
        row("2", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "AHMED  MOHAMMED"}), row("3", "Ahmed Mohammad Hassan"),
        row("4", "Omar Khaled", **{"Name of the caller اسم المتصل": "Omar"}),
        row("5", "محمد أحمد", **{"Name of the caller اسم المتصل": "علي حسن"}), row("6", "محمد أحمد", **{"Name of the caller اسم المتصل": "على حسن"}),
        row("7", "Ahmed Mohammed Hassan"),
    ])
    result = review(store)
    assert result["rules"][0] == next(iter(result["ruleCounts"])) == COMBINED_NAME_RULE
    expected = {COMBINED_NAME_RULE: {"id:1", "id:2", "id:5", "id:6"}, RULES[0]: {"id:7"}, RULES[1]: {"id:3"}, RULES[2]: {"id:4"}}
    name_rows = []
    for rule, identifiers in expected.items():
        table = review(store, rule=rule)
        assert {flag["hotlineId"] for flag in table["rows"]} == identifiers
        assert table["total"] == result["ruleCounts"][rule] == len(identifiers)
        name_rows.extend(table["rows"])
    assert len(name_rows) == len({flag["hotlineId"] for flag in name_rows})
    combined = review(store, rule=COMBINED_NAME_RULE, page_size=1, page=2)
    assert combined["total"] == 4 and len(combined["rows"]) == 1
    assert all(flag["duplicateGroup"] and flag["duplicateSimilarity"] == 100 for flag in combined["rows"])
    # Unrelated findings remain visible for the same record.
    assert {flag["hotlineId"] for flag in review(store, rule=RULES[3])["rows"]} == {"id:1"}
    assert all(flag["hotlineId"] != "id:1" for flag in review(store, rule=RULES[0], search="id:1")["rows"])


def test_only_combined_exclusions_hide_combined_findings_without_moving_tables():
    rows = [row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan")]
    for excluded_rule in (COMBINED_NAME_RULE, RULES[0], RULES[2]):
        store = store_for(rows, [{"dataset": "legalhotlines", "rule": excluded_rule, "identifierType": "hotlineId", "identifierValue": "id:1"}])
        result = review(store)
        expected = {"id:2"} if excluded_rule == COMBINED_NAME_RULE else {"id:1", "id:2"}
        assert {flag["hotlineId"] for flag in result["rows"]} == expected
        assert all(flag["rule"] == COMBINED_NAME_RULE for flag in result["rows"])


def test_combined_export_contains_one_name_finding_per_record():
    store = store_for([row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan"), row("3", "Ahmed Mohammad Hassan")])
    book = load_workbook(BytesIO(store.review_export("legalhotlines", selected_rules=[COMBINED_NAME_RULE, *RULES[:3]], allow_name_variations=True)))
    records = []
    for sheet in book:
        headers = [cell.value for cell in sheet[1]]
        records.extend(value for cells in sheet.iter_rows(min_row=2) if isinstance(value := cells[headers.index("Record identifier")].value, str) and value.startswith("id:"))
    assert len(records) == 3 and set(records) == {"id:1", "id:2", "id:3"}


def test_combined_requires_both_names_to_match_the_same_other_record():
    caller_column = "Name of the caller اسم المتصل"
    store = store_for([
        row("1", "Ahmed Hassan", **{caller_column: "Omar Khaled"}),
        row("2", "Ahmed Hassan", **{caller_column: "Ali Salman"}),
        row("3", "Mohammed Saleh", **{caller_column: "Omar Khaled"}),
        row("4", "Unique Person", **{caller_column: "Unique Person"}),
        row("5", "", **{caller_column: "Omar Khaled"}),
        row("6", "", **{caller_column: "Omar Khaled"}),
        row("7", "Another Person", **{caller_column: ""}),
        row("8", "Another Person", **{caller_column: ""}),
    ])
    assert review(store, rule=COMBINED_NAME_RULE)["total"] == 0
    assert {flag["hotlineId"] for flag in review(store, rule=RULES[2])["rows"]} == {"id:4"}


def test_combined_group_references_only_peers_matching_both_columns_and_ignores_fuzzy_settings():
    rows = [row(str(i), "Ahmed Hassan", Projects="A" if i == 1 else "B") for i in range(1, 4)]
    rows.append(row("4", "Ahmed Hassan", **{"Name of the caller اسم المتصل": "Other Caller"}))
    store = store_for(rows)
    for compare_chars, variations in ((10, True), (30, False)):
        result = store.review("legalhotlines", rule=COMBINED_NAME_RULE,
                              name_compare_chars=compare_chars, allow_name_variations=variations)
        assert result["total"] == 3
        for flag in result["rows"]:
            assert flag["duplicateContext"]["members"] == ["id:1", "id:2", "id:3"]
            assert set(flag["duplicateMatchRows"]) == {2, 3, 4} - {flag["row"]}
            assert flag["duplicateGroup"].startswith("hotline-name-caller:")
            assert "Both normalized" in flag["detail"]
    filtered = review(store, rule=COMBINED_NAME_RULE, project="A")
    assert filtered["total"] == 1
    assert filtered["rows"][0]["duplicateContext"]["members"] == ["id:1", "id:2", "id:3"]


def test_saved_old_combined_context_is_rechecked_under_new_pair_criteria(tmp_path):
    rows = [row("1", "Ahmed Hassan"), row("2", "Ahmed Hassan")]
    registry = DuplicateExclusionRegistry(tmp_path / "old-combined.json")
    old_context = {"version": 1, "key": "hotline-exact:ahmedhassan",
                   "members": ["id:1", "id:2"], "settings": {"callerName": "ahmed"}}
    registry.exclude_record("legalhotlines", COMBINED_NAME_RULE, "hotlineId", "id:1", duplicate_context=old_context)
    store = store_for(rows, registry.exclusion_rows())
    entries, changed = registry.reconcile_duplicate_contexts(store.duplicate_contexts_for_exclusions(registry.exclusion_rows()))
    assert changed and entries[0]["pendingRecheck"]
    store.set_review_exclusions(registry.exclusion_rows())
    assert any(flag["hotlineId"] == "id:1" and flag["duplicateRecheck"] for flag in review(store, rule=COMBINED_NAME_RULE)["rows"])


def test_caller_match_takes_priority_over_similar_name_and_missing_caller_is_unavailable():
    store = store_for([row("1", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Ahmed"}), row("2", "Ahmed Mohammad Hassan")])
    assert {flag["hotlineId"] for flag in review(store, rule=RULES[2])["rows"]} == {"id:1"}
    assert {flag["hotlineId"] for flag in review(store, rule=RULES[1])["rows"]} == {"id:2"}
    store = store_for([{"Hotline ID": "1", "Name of the detained person": "Ahmed Hassan"}, {"Hotline ID": "2", "Name of the detained person": "Ahmed Hassan"}])
    result = review(store)
    assert COMBINED_NAME_RULE in result["unavailableRules"]
    assert result["ruleCounts"][COMBINED_NAME_RULE] == 0 and result["ruleCounts"][RULES[0]] == 2


def test_combined_duplicate_exclusions_recheck_after_group_changes(tmp_path):
    rows = [row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan")]
    for rule in (COMBINED_NAME_RULE,):
        store = store_for(rows)
        target = review(store, rule=COMBINED_NAME_RULE)["rows"][0]
        registry = DuplicateExclusionRegistry(tmp_path / "combined.json")
        registry.exclude_record("legalhotlines", rule, "hotlineId", "id:1", duplicate_context=target["duplicateContext"])
        unchanged = store_for(rows, registry.exclusion_rows())
        assert review(unchanged, rule=COMBINED_NAME_RULE)["total"] == 1
        assert registry.reconcile_duplicate_contexts(unchanged.duplicate_contexts_for_exclusions(registry.exclusion_rows()))[1] is False
        changed = store_for(rows + [row("3", "Ahmed Mohammed Hassan")], registry.exclusion_rows())
        entries, did_change = registry.reconcile_duplicate_contexts(changed.duplicate_contexts_for_exclusions(registry.exclusion_rows()))
        assert did_change and entries[0]["pendingRecheck"]
        changed.set_review_exclusions(registry.exclusion_rows())
        assert review(changed, rule=COMBINED_NAME_RULE)["total"] == 3


def test_phone_validation_allows_blank_and_only_formatting():
    phones = ["", "7501234567", "750 123-4567", "(750)1234567", "+9647501234567", "07501234567", "750123456", "75012345678", "abc7501234567", "750/1234567"]
    store = store_for([row(str(i), **{"Contact Number": phone}) for i, phone in enumerate(phones)])
    assert review(store, rule=RULES[6])["total"] == 6


def test_missing_columns_do_not_generate_false_findings():
    store = store_for([{"Hotline ID": "1", "Priority": "High"}])
    result = review(store)
    assert result["total"] == 0
    assert set(result["unavailableRules"]) == set(RULES)
    assert "Refer to Lawyer" in result["unavailableRules"][RULES[5]]


def test_review_includes_remaining_source_fields_without_repeating_existing_columns():
    store = store_for([
        row("1", "Ahmed Hassan", **{"Gender": "Male", "Detaining Authority": "Authority A", "Notes": "Check source", "Priority الاولوية": "High"}),
        row("2", "Ahmed Hassan", **{"Gender": "", "Detaining Authority": "Authority B", "Notes": ""}),
    ])
    result = review(store, rule=COMBINED_NAME_RULE, page_size=1)
    flag = result["rows"][0]
    assert flag["priority"] == "High"
    assert flag["sourceFields"] == {"Hotline ID": "1", "Gender": "Male", "Detaining Authority": "Authority A", "Notes": "Check source"}
    second = review(store, rule=COMBINED_NAME_RULE, page_size=1, page=2)["rows"][0]
    assert list(second["sourceFields"]) == list(flag["sourceFields"])
    assert second["sourceFields"]["Gender"] == ""
    assert review(store, rule=COMBINED_NAME_RULE, search="Check source")["total"] == 1


def test_arabic_headers_and_blank_answers():
    store = store_for([{"اسم الشخص المعتقل": "محمد علي", "اسم المتصل": "محمد", "هل المستفيد موقوف": "نعم", "هل تم احالة المستفيد الى خط المساعدة": "لا", "الاولوية": "عالية", "احالة الى محامي": "", "رقم الهاتف": "7501234567"},
                       {"اسم الشخص المعتقل": "", "اسم المتصل": "", "هل المستفيد موقوف": "", "هل تم احالة المستفيد الى خط المساعدة": "لا", "الاولوية": "", "احالة الى محامي": "", "رقم الهاتف": ""}])
    counts = review(store)["ruleCounts"]
    assert counts[RULES[2]] == counts[RULES[3]] == counts[RULES[5]] == 1
    assert counts[RULES[4]] == 0


def test_stable_identity_exclusion_restore_and_recheck(tmp_path):
    rows = [row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Other Caller"})]
    store = store_for(rows)
    target = review(store, rule=RULES[0])["rows"][0]
    registry = DuplicateExclusionRegistry(tmp_path / "exclusions.json")
    registry.exclude_record("legalhotlines", RULES[0], "hotlineId", target["hotlineId"], duplicate_context=target["duplicateContext"])
    store.set_review_exclusions(registry.exclusion_rows())
    assert review(store, rule=RULES[0])["total"] == 1
    assert dashboard(store.frames["legalhotlines"])["total"] == 2
    reordered = store_for(list(reversed(rows)), registry.exclusion_rows())
    assert review(reordered, rule=RULES[0])["total"] == 1
    assert registry.reconcile_duplicate_contexts(reordered.duplicate_contexts_for_exclusions(registry.exclusion_rows()))[1] is False
    changed = store_for(rows + [row("3", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Third Caller"})], registry.exclusion_rows())
    entries, did_change = registry.reconcile_duplicate_contexts(changed.duplicate_contexts_for_exclusions(registry.exclusion_rows()))
    assert did_change and entries[0]["pendingRecheck"]
    changed.set_review_exclusions(registry.exclusion_rows())
    assert review(changed, rule=RULES[0])["total"] == 3
    registry.restore(target["hotlineId"], RULES[0], "legalhotlines", "hotlineId")
    changed.set_review_exclusions(registry.exclusion_rows())
    assert review(changed, rule=RULES[0])["total"] == 3


def test_fallback_ids_stable_after_reorder_including_identical_rows():
    rows = [row("", "Alpha Person"), row("", "Beta Person"), row("", "Alpha Person")]
    first = store_for(rows); second = store_for(list(reversed(rows)))
    from backend.hotline_review import prepare
    ids_a = prepare(first.frames["legalhotlines"])[1]
    ids_b = prepare(second.frames["legalhotlines"])[1]
    assert len(set(ids_a.values())) == 3
    assert set(ids_a.values()) == set(ids_b.values())


def test_export_filtering_and_pagination():
    store = store_for([row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Other Caller"}), row("3", **{"Priority الاولوية": "High"})])
    result = review(store, rule=RULES[0], page_size=1, page=2)
    assert result["total"] == 2 and len(result["rows"]) == 1
    workbook = load_workbook(BytesIO(store.review_export("legalhotlines", selected_rules=[RULES[5]], allow_name_variations=True)))
    headers = [cell.value for cell in workbook.active[1]]
    assert "Record identifier" in headers and "Caller name" in headers and "Lawyer referral" in headers
    assert "Case ID" not in headers and "Assessment" not in headers and "Service" not in headers
    values = [cell.value for sheet in workbook for cells in sheet.iter_rows() for cell in cells]
    assert RULES[5] in values and RULES[0] not in values


def test_review_api_exports_and_exclusion_import(monkeypatch, tmp_path):
    import asyncio
    from fastapi import UploadFile
    from backend import main
    store = store_for([row("1", "Ahmed Mohammed Hassan"), row("2", "Ahmed Mohammed Hassan", **{"Name of the caller اسم المتصل": "Other Caller"})])
    registry = DuplicateExclusionRegistry(tmp_path / "api-exclusions.json")
    monkeypatch.setattr(main, "legal_store", store)
    monkeypatch.setattr(main, "duplicate_exclusions", registry)
    response = main.legal_review(main.LegalQuery(dataset="legalhotlines", rule=RULES[0], allowNameVariations=True))
    assert response["total"] == 2
    export = main.legal_review_export("legalhotlines", rules=RULES[0])
    book = load_workbook(BytesIO(export.body))
    assert "Record identifier" in [cell.value for cell in book.active[1]]
    response = asyncio.run(main.import_duplicate_exclusions(UploadFile(filename="ids.csv", file=BytesIO(b"Record identifier\nid:1\n")), "legalhotlines", "hotlineId", RULES[0]))
    assert response["imported"] == 1
    assert main.legal_review(main.LegalQuery(dataset="legalhotlines", rule=RULES[0]))["total"] == 1
