"""Reusable workbook styles and row writers. No exported data is cached."""
from functools import lru_cache
from math import isinf
from datetime import date, timedelta

import pandas as pd
from openpyxl.styles import Alignment as _Alignment, Border as _Border, Font as _Font
from openpyxl.styles import PatternFill as _PatternFill, Side as _Side


# These shared styles are immutable after construction and assignment to cells.
Alignment = lru_cache(maxsize=256)(_Alignment)
Border = lru_cache(maxsize=256)(_Border)
Font = lru_cache(maxsize=256)(_Font)
PatternFill = lru_cache(maxsize=256)(_PatternFill)
Side = lru_cache(maxsize=64)(_Side)


def excel_value(value):
    """Match pandas' Excel scalar handling without its per-cell writer overhead."""
    if isinstance(value, str):
        return value
    if not pd.api.types.is_scalar(value):
        return str(value)
    if pd.isna(value):
        return ""
    if pd.api.types.is_bool(value):
        return bool(value)
    if pd.api.types.is_integer(value):
        return int(value)
    if pd.api.types.is_float(value):
        return ("inf" if value > 0 else "-inf") if isinf(value) else float(value)
    if isinstance(value, date):
        return value
    if isinstance(value, timedelta):
        return value.total_seconds() / 86400
    return str(value)


def append_frame(sheet, frame):
    """Write a flat frame with the existing pandas-style header appearance."""
    if not len(frame.columns):
        return
    sheet.append([excel_value(value) for value in frame.columns])
    thin = Side(style="thin")
    header_border = Border(top=thin, right=thin, bottom=thin, left=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="top")
    for row_number,row in enumerate(frame.itertuples(index=False, name=None),2):
        sheet.append([excel_value(value) for value in row])
        for column,value in enumerate(row,1):
            if isinstance(value,timedelta):
                sheet.cell(row_number,column).number_format="0"
