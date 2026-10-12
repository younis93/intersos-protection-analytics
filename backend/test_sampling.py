import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from .sampling import LEADING, export, pipeline, table


def record(identifier="1", phone="7712345678", gender="Male ذكر", completed="01/10/2026", **changes):
    row = {names[0]: "" for _, *names in LEADING}
    row.update({"Service ID": identifier, "Beneficiary ID": identifier, "Contact Number": phone,
                "Gender": gender, "Date Service Completed": completed, "Service Status": "Completed اکتملت",
                "Type of Service Provided": "Legal Representation - تمثيل", "Projects": "P1", "Project Location": "L1"})
    row.update(changes)
    return row


def ids(result):
    return [row["Service ID"] for row in result["rows"]]


@pytest.mark.parametrize("phone,valid", [
    ("7712345678", True), ("7812345678", True), ("7512345678", True), ("٧٧١٢٣٤٥٦٧٨", True),
    ("۷۸۱۲۳۴۵۶۷۸", True), ("77 123-45(678)", True), ("\u200e7712345678", True),
    ("", False), (None, False), ("77123456", False), ("771234567", False), ("77123456789", False),
    ("07712345678", False), ("+9647712345678", False), ("7412345678", False), ("7912345678", False),
])
def test_exact_primary_phone(phone, valid):
    source = pd.DataFrame([record(phone=phone, **{"Secondary Contact Number": "7712345678"})])
    result = table(source)
    assert result["total"] == int(valid)
    if valid:
        assert result["rows"][0]["Contact Number"] == phone


@pytest.mark.parametrize("status,service,valid", [
    ("Completed اکتملت", "Legal Representation - تمثيل", True),
    ("Completed", "Legal Representation", True), ("اكتملت", "تمثيل", True),
    ("Uncompleted لم تكتمل", "Legal Representation", False),
    ("In-Process قيد المعالجة", "Legal Representation", False),
    ("Completed", "Legal Assistance - مساعدة", False),
    ("Completed", "Legal Counselling - استشارة", False), ("", "Legal Representation", False),
])
def test_service_eligibility(status, service, valid):
    assert table(pd.DataFrame([record(**{"Service Status": status, "Type of Service Provided": service})]))["total"] == int(valid)


def conflicts():
    return pd.DataFrame([
        record("old", "7711111118", completed="01/01/2026", **{"Beneficiary ID": "A"}),
        record("latest", "7822222228", completed="03/01/2026", **{"Beneficiary ID": "A"}),
        record("phone-conflict", "7822222228", completed="02/01/2026", **{"Beneficiary ID": "B"}),
        record("retained-after-skip", "7533333338", completed="01/01/2026", **{"Beneficiary ID": "B"}),
        record("empty-one", "7544444448", completed="invalid", **{"Beneficiary ID": ""}),
        record("empty-two", "7555555558", completed="", **{"Beneficiary ID": None}),
        record("tie", "7566666668", completed="03/01/2026", **{"Beneficiary ID": "A"}),
    ])


def test_duplicate_flags_independent_and_reversible():
    source = conflicts()
    original = source.copy(deep=True)
    assert ids(table(source, removeDuplicateBeneficiaryIds=True)) == ["latest", "phone-conflict", "empty-one", "empty-two"]
    assert ids(table(source, removeDuplicateContacts=True)) == ["latest", "tie", "old", "retained-after-skip", "empty-one", "empty-two"]
    assert ids(table(source, removeDuplicateBeneficiaryIds=True, removeDuplicateContacts=True)) == ["latest", "retained-after-skip", "empty-one", "empty-two"]
    assert table(source)["total"] == 7
    pd.testing.assert_frame_equal(source, original)


def test_normalized_contact_duplicates_keep_source_values():
    source = pd.DataFrame([record("older", phone="٧٧١٢٣٤٥٦٧٨", completed="01/01/2026"),
                           record("newer", phone="77 123-45(678)", completed="02/01/2026")])
    assert table(source)["total"] == 2
    result = table(source, removeDuplicateContacts=True)
    assert ids(result) == ["newer"]
    assert result["rows"][0]["Contact Number"] == "77 123-45(678)"


def test_search_and_filters_before_dedup_month_and_self_excluding_facets():
    source = conflicts()
    source.loc[0, "Projects"] = "P2"
    assert ids(table(source, filters={"Projects": ["P2"]}, removeDuplicateBeneficiaryIds=True)) == ["old"]
    assert ids(table(source, search="old", removeDuplicateBeneficiaryIds=True)) == ["old"]
    result = table(source, filters={"Projects": ["P2"], "Date Service Completed": ["2026-01"]})
    assert ids(result) == ["old"]
    assert result["filterOptions"]["Projects"] == ["P1", "P2"]
    assert table(source, filters={"Projects": ["unavailable"]})["total"] == 0
    assert table(source, search="no results")["total"] == 0


def test_global_gender_order_pagination_and_export_parity():
    source = pd.DataFrame([record(str(index), gender="Male" if index % 2 else "Female", completed=f"{index % 28 + 1:02}/01/2026") for index in range(205)])
    source.loc[0, "Gender"] = "Other"
    source.loc[1, "Gender"] = ""
    source.loc[2, "Name"] = "=HYPERLINK(\"test\")"
    source.loc[3, "Name"] = "اسم عربي"
    whole = table(source, pageSize=10000)
    assert [row["Gender"] for row in whole["rows"][:200]] == ["Male", "Female"] * 100
    assert [row["Gender"] for row in whole["rows"][-2:]] == ["", "Other"]
    pages = [table(source, page=page) for page in (1, 2, 3)]
    assert sum((ids(page) for page in pages), []) == ids(whole)
    assert table(source, page=999)["page"] == 3
    workbook = load_workbook(io.BytesIO(export(source)))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers == whole["columns"]
    assert headers[:12] == [names[0] for _, *names in LEADING]
    service = headers.index("Service ID")
    assert [str(row[service]) for row in sheet.iter_rows(min_row=2, values_only=True)] == ids(whole)
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    assert sheet.freeze_panes == "A2"


def test_unequal_genders_and_invalid_dates_last_with_stable_ties():
    source = pd.DataFrame([record("m1", completed="invalid"), record("f1", gender="Female"),
                           record("m2"), record("m3"), record("other", gender="Other"),
                           record("arabic", gender="أنثى", completed="02/10/2026")])
    assert ids(table(source)) == ["m2", "arabic", "m3", "f1", "m1", "other"]


def test_iso_and_local_dates_share_correct_months_and_priority():
    source = pd.DataFrame([record("iso", completed="2026-10-12", **{"Beneficiary ID": "same"}),
                           record("local", completed="11/10/2026", **{"Beneficiary ID": "same"}),
                           record("new", completed="2026-11-01", **{"Beneficiary ID": "same"})])
    assert ids(table(source, removeDuplicateBeneficiaryIds=True)) == ["new"]
    result = table(source, filters={"Date Service Completed": ["2026-10"]}, removeDuplicateBeneficiaryIds=True)
    assert ids(result) == ["iso"]
    assert result["filterOptions"]["Date Service Completed"] == ["2026-10", "2026-11"]


def test_bilingual_columns_preserved_and_lookalikes_not_selected():
    source = pd.DataFrame([record()]).rename(columns={
        "Projects": "Projects - المشروع", "Name": "Name / الأسم", "Contact Number": "Contact Number / رقم الاتصال",
        "Type of Document": "Type of Document نوع الوثيقة", "Gender": "Gender النوع الاجتماعي"})
    source.insert(0, "Nationality Stateless", "No")
    source.insert(0, "Age Gender Group", "Adult")
    result = table(source)
    assert result["columns"][:4] == ["Projects - المشروع", "Project Location", "Name / الأسم", "Gender النوع الاجتماعي"]
    assert len(result["columns"]) == len(set(result["columns"])) == len(source.columns)
    secondary_only = source.drop(columns=["Contact Number / رقم الاتصال"])
    assert not table(secondary_only)["available"]
    assert "primary Contact Number" in table(secondary_only)["reason"]


def test_partial_import_unavailable_and_empty_eligible_frame():
    assert not table(None)["available"]
    assert not table(pd.DataFrame({"Service ID": ["1"]}))["available"]
    with pytest.raises(ValueError, match="not been imported"):
        export(None)
    result = table(pd.DataFrame([record(phone="")]))
    assert result["available"] and result["total"] == 0 and result["columns"]
    assert load_workbook(io.BytesIO(export(pd.DataFrame([record(phone="")])))).active.max_row == 1


def test_excel_duplicate_pipeline_matches_table():
    source = conflicts()
    query = dict(removeDuplicateBeneficiaryIds=True, removeDuplicateContacts=True, filters={"Projects": ["P1"]})
    result = table(source, **query)
    sheet = load_workbook(io.BytesIO(export(source, **query))).active
    service = result["columns"].index("Service ID")
    assert [row[service] for row in sheet.iter_rows(min_row=2, values_only=True)] == ids(result)


def test_read_only_endpoints_and_validation(monkeypatch):
    from types import SimpleNamespace
    from fastapi.testclient import TestClient
    monkeypatch.setenv("INTERSOS_DEFER_LEGAL_LOAD", "1")
    from . import main
    source = conflicts()
    original = source.copy(deep=True)
    monkeypatch.setattr(main, "require_legal_store", lambda: SimpleNamespace(frames={"legalservices": source}))
    with TestClient(main.app) as client:
        result = client.post("/api/legal/sampling", json={"removeDuplicateContacts": True})
        assert result.status_code == 200
        assert result.json()["total"] == 6
        assert client.post("/api/legal/sampling", json={"page": 0}).status_code == 400
        assert client.post("/api/legal/sampling", json={"pageSize": 10001}).status_code == 400
        download = client.post("/api/legal/sampling/export", json={"removeDuplicateContacts": True})
        assert download.status_code == 200
        assert "sampling.xlsx" in download.headers["content-disposition"]
        assert load_workbook(io.BytesIO(download.content)).active.max_row == 7
    pd.testing.assert_frame_equal(source, original)
