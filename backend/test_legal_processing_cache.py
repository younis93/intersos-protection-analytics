import io
import json
import os
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from backend import legal_processing_cache as cache
from backend.import_progress import ImportCancelled, ImportOperation
from backend.legal_platform import LegalStore
from backend.test_legal_platform import required_payload, csv


def comparable_metadata(store):
    return {key: value for key, value in store.metadata().items() if key not in {"source", "revision"}}


def test_cache_skips_expensive_work_and_restores_independent_store():
    payload = required_payload()
    original = LegalStore.from_files(payload, "old folder")
    assert original.import_timings["cache_miss"] == 1
    operation = ImportOperation()
    progress = []
    with patch("pandas.read_csv", side_effect=AssertionError("CSV parsing repeated")), \
         patch.object(LegalStore, "_build_flags", side_effect=AssertionError("Validation repeated")), \
         patch.object(LegalStore, "review", side_effect=AssertionError("Review preparation repeated")), \
         patch("backend.linked_csv_enrichment.enrich_linked_csvs", side_effect=AssertionError("Enrichment repeated")):
        restored = LegalStore.from_files(payload, "new folder", progress.append, operation=operation)
    assert restored.import_timings["cache_hit"] == 1
    assert restored.import_timings["cache_restore"] > 0
    assert restored.revision != original.revision and restored.data_revision != original.data_revision
    assert restored.source == "new folder" and restored.metadata()["source"] == "new folder"
    assert restored.metadata()["revision"] == restored.revision
    assert comparable_metadata(restored) == comparable_metadata(original)
    assert restored.flags == original.flags
    assert restored.source_files == payload
    assert restored._cache_lock is not original._cache_lock
    assert restored._processing_lock is not original._processing_lock
    assert restored._import_operation is None
    assert progress == sorted(progress)
    assert not any(task.startswith(("Parsing:", "Cleaning:", "Validating:")) for task in operation.tasks)
    assert operation.tasks["Restoring saved processing"] == 1
    for name in payload:
        pd.testing.assert_frame_equal(restored.frames[name], original.frames[name])
        assert restored.review(name) == original.review(name)
        # Compare actual workbook data, ignoring ZIP timestamps and styles.
        before = load_workbook(io.BytesIO(original.export(name)), data_only=True)
        after = load_workbook(io.BytesIO(restored.export(name)), data_only=True)
        assert [list(sheet.values) for sheet in before] == [list(sheet.values) for sheet in after]
    restored.frames["beneficiaries"].iloc[0, 0] = "modified"
    assert original.frames["beneficiaries"].iloc[0, 0] == "B1"


@pytest.mark.parametrize("change", ["contents", "added", "removed", "exclusions", "date", "version", "app_version", "pandas", "python"])
def test_cache_invalidates_when_inputs_change(change, monkeypatch):
    payload = {**required_payload(), "awareness": csv(**{"Awareness ID": ["W1"], "Name": ["Test"]})}
    LegalStore.from_files(payload, "first")
    exclusions = []
    if change == "contents": payload["beneficiaries"] += b"\n"
    if change == "added": payload["legalfees"] = csv(**{"Fee ID": ["F1"]})
    if change == "removed": payload.pop("awareness")
    if change == "exclusions": exclusions = [("Invalid age", "B2")]
    if change == "date":
        tomorrow = date.today() + timedelta(days=1)
        class Tomorrow(date):
            @classmethod
            def today(cls): return tomorrow
        monkeypatch.setattr(cache, "date", Tomorrow)
    if change == "version": monkeypatch.setattr(cache, "CACHE_VERSION", cache.CACHE_VERSION + 1)
    if change == "app_version": monkeypatch.setattr(cache, "APP_VERSION", "new-build")
    if change == "pandas": monkeypatch.setattr(cache.pd, "__version__", "incompatible")
    if change == "python": monkeypatch.setattr(cache, "sys", SimpleNamespace(version_info=(99, 0, 0)))
    result = LegalStore.from_files(payload, "second", exclusions=exclusions)
    assert result.import_timings["cache_miss"] == 1
    assert result.import_timings["validation"] >= 0


def test_reordered_dataset_selection_uses_cache_and_current_sheet_order():
    payload = required_payload()
    first = LegalStore.from_files(payload, "first")
    reordered = dict(reversed(list(payload.items())))
    second = LegalStore.from_files(reordered, "second")
    assert second.import_timings["cache_hit"] == 1
    assert list(second.frames) == list(reordered)
    assert [sheet["id"] for sheet in second.metadata()["sheets"]] == list(reordered)
    assert second.flags == first.flags


def test_current_reconciliation_runs_on_hit_and_rebuilds_changed_reviews():
    payload = required_payload()
    LegalStore.from_files(payload, "first")
    observed = []
    def reconcile(store):
        observed.append(store)
        store.set_review_exclusions([("Invalid contact number", "B1")])
        store._commit_import_exclusions = lambda: None
    restored = LegalStore.from_files(payload, "second", reconcile_exclusions=reconcile)
    assert restored.import_timings["cache_hit"] == 1
    assert observed == [restored]
    assert restored._review_cache
    assert not any(row["rule"] == "Invalid contact number" and row.get("caseId") == "B1"
                   for row in restored.review("beneficiaries")["rows"])


def test_snapshot_whitelists_state_and_keeps_only_default_reviews():
    store = LegalStore.from_files(required_payload(), "first")
    store.review("beneficiaries", name_compare_chars=20)
    store._commit_import_exclusions = lambda: None
    store._lawyer_cache[()] = {"not": "cached"}
    key = cache.cache_key(store.source_files, [])
    cache.save(key, store)
    state = cache.load(key)
    assert "_commit_import_exclusions" not in state and "_lawyer_cache" not in state
    assert all(key[1:] == ("", 15, False) for key in state["_review_cache"])


@pytest.mark.parametrize("damage", ["signature", "truncated", "key", "missing"])
def test_damaged_cache_is_never_deserialized_and_falls_back(damage):
    payload = required_payload()
    LegalStore.from_files(payload, "first")
    snapshot = cache.cache_directory() / "snapshot.bin"
    if damage == "missing": snapshot.unlink()
    elif damage == "key": (cache.cache_directory() / "authentication.key").write_bytes(b"wrong key")
    elif damage == "truncated": snapshot.write_bytes(snapshot.read_bytes()[:20])
    else:
        raw = bytearray(snapshot.read_bytes())
        raw[-1] ^= 1
        snapshot.write_bytes(raw)
    with patch.object(cache.pickle, "loads", side_effect=AssertionError("Unauthenticated deserialization")):
        result = LegalStore.from_files(payload, "second")
    assert result.import_timings["cache_miss"] == 1
    assert result.metadata()["ready"]
    assert LegalStore.from_files(payload, "third").import_timings["cache_hit"] == 1


def test_authenticated_but_incompatible_state_falls_back():
    payload = required_payload()
    LegalStore.from_files(payload, "first")
    with patch.object(cache.pickle, "loads", return_value={"old_schema": True}):
        assert LegalStore.from_files(payload, "second").import_timings["cache_miss"] == 1
    with patch.object(cache.pickle, "loads", side_effect=ValueError("incompatible pickle")):
        assert LegalStore.from_files(payload, "third").import_timings["cache_miss"] == 1


def test_oversize_and_failed_atomic_replace_preserve_last_snapshot(monkeypatch):
    first = LegalStore.from_files(required_payload(), "first")
    snapshot = cache.cache_directory() / "snapshot.bin"
    previous = snapshot.read_bytes()
    with monkeypatch.context() as context:
        context.setattr(cache, "MAX_SNAPSHOT_BYTES", 20)
        cache.save(cache.cache_key(first.source_files, []), first)
    assert snapshot.read_bytes() == previous
    with patch.object(Path, "replace", side_effect=OSError("interrupted write")):
        cache.save(cache.cache_key(first.source_files, []), first)
    assert snapshot.read_bytes() == previous
    assert not list(cache.cache_directory().glob("*.tmp"))
    # Oversized existing snapshots are rejected before deserialization.
    with monkeypatch.context() as context:
        context.setattr(cache, "MAX_SNAPSHOT_BYTES", 20)
        with patch.object(cache.pickle, "loads", side_effect=AssertionError("Oversized read")):
            assert cache.load(cache.cache_key(first.source_files, [])) is None


def test_unwritable_storage_does_not_fail_import(monkeypatch, tmp_path):
    blocked = tmp_path / "blocked"
    blocked.write_text("file instead of directory")
    monkeypatch.setattr(cache, "cache_directory", lambda: blocked)
    result = LegalStore.from_files(required_payload(), "source")
    assert result.import_timings["cache_miss"] == 1 and result.metadata()["ready"]


def test_failed_and_cancelled_publications_preserve_previous_cache(monkeypatch):
    from backend import main
    payload = required_payload()
    previous = LegalStore.from_files(payload, "previous")
    monkeypatch.setattr(main, "legal_store", previous)
    monkeypatch.setattr(main, "prepare_import_exclusions", lambda store: None)
    snapshot = cache.cache_directory() / "snapshot.bin"
    saved = snapshot.read_bytes()
    changed = {**payload, "beneficiaries": payload["beneficiaries"] + b"\n"}
    for failure in ("cancel", "remember"):
        operation = ImportOperation()
        candidate = LegalStore.from_files(changed, "new", operation=operation)
        assert snapshot.read_bytes() == saved
        if failure == "cancel": operation.cancel()
        def cannot_remember(): raise OSError("remembered selection write failed")
        with pytest.raises(ImportCancelled if failure == "cancel" else OSError):
            main.publish_legal_candidate(candidate, operation, before_publish=cannot_remember)
        operation.fail("failed")
        assert main.legal_store is previous and snapshot.read_bytes() == saved
    operation = ImportOperation()
    candidate = LegalStore.from_files(changed, "new", operation=operation)
    main.publish_legal_candidate(candidate, operation)
    assert main.legal_store is candidate and snapshot.read_bytes() != saved
    assert not operation.completion_callbacks


def test_cancel_during_cached_restore_preserves_previous_cache(monkeypatch):
    payload = required_payload()
    LegalStore.from_files(payload, "first")
    snapshot = cache.cache_directory() / "snapshot.bin"
    previous = snapshot.read_bytes()
    operation = ImportOperation()
    start = operation.start
    def cancel(name):
        if name == "Restoring saved processing": operation.cancel()
        start(name)
    monkeypatch.setattr(operation, "start", cancel)
    with pytest.raises(ImportCancelled):
        LegalStore.from_files(payload, "cancelled", operation=operation)
    assert snapshot.read_bytes() == previous


def test_snapshot_reopens_in_fresh_process(tmp_path):
    payload = required_payload()
    original = LegalStore.from_files(payload, "parent")
    folder = tmp_path / "source"
    folder.mkdir()
    for name, raw in payload.items(): (folder / (name + ".csv")).write_bytes(raw)
    script = '''
import json, sys
from pathlib import Path
from unittest.mock import patch
from backend import legal_processing_cache as cache
from backend.legal_platform import LegalStore
cache.cache_directory = lambda: Path(sys.argv[1])
with patch("pandas.read_csv", side_effect=AssertionError("repeated parsing")), patch.object(LegalStore, "_build_flags", side_effect=AssertionError("repeated validation")):
    store = LegalStore.from_folder(Path(sys.argv[2]))
print(json.dumps({"hit": store.import_timings["cache_hit"], "revision": store.revision, "sheets": store.metadata()["sheets"]}))
'''
    result = subprocess.run([sys.executable, "-c", script, str(cache.cache_directory()), str(folder)],
        cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, check=True, timeout=60)
    restored = json.loads(result.stdout)
    assert restored["hit"] == 1 and restored["revision"] != original.revision
    assert sorted(restored["sheets"], key=lambda sheet: sheet["id"]) == sorted(original.metadata()["sheets"], key=lambda sheet: sheet["id"])


def test_folder_cache_survives_move_and_keeps_version_selection(tmp_path):
    first = tmp_path / "first"
    second = tmp_path / "second"
    first.mkdir()
    second.mkdir()
    payload = required_payload()
    for name, raw in payload.items():
        (first / (name + ".csv")).write_bytes(raw)
        (second / (name + ".csv")).write_bytes(raw)
    original = LegalStore.from_folder(first)
    restored = LegalStore.from_folder(second)
    assert original.import_timings["cache_miss"] == restored.import_timings["cache_hit"] == 1
    assert restored.source == str(second.resolve())
    assert restored.metadata()["source"] == str(second.resolve())
    # A selected newer dataset version must invalidate the prior selection.
    (second / "beneficiaries (2).csv").write_bytes(payload["beneficiaries"] + b"\n")
    changed = LegalStore.from_folder(second)
    assert changed.source_files["beneficiaries"] == payload["beneficiaries"] + b"\n"
    assert changed.import_timings["cache_miss"] == 1


def test_same_size_and_timestamp_still_detect_changed_contents(tmp_path):
    for name, raw in required_payload().items(): (tmp_path / (name + ".csv")).write_bytes(raw)
    LegalStore.from_folder(tmp_path)
    source = tmp_path / "beneficiaries.csv"
    previous_stat = source.stat()
    changed = source.read_bytes().replace(b"B1", b"C1")
    source.write_bytes(changed)
    os.utime(source, ns=(previous_stat.st_atime_ns, previous_stat.st_mtime_ns))
    assert source.stat().st_size == previous_stat.st_size
    result = LegalStore.from_folder(tmp_path)
    assert result.import_timings["cache_miss"] == 1
    assert result.source_files["beneficiaries"] == changed


def test_restore_stage_is_active_during_deserialization():
    payload = required_payload()
    LegalStore.from_files(payload, "first")
    operation = ImportOperation()
    deserialize = cache.pickle.loads
    def observed(raw):
        assert operation.current == "Restoring saved processing"
        assert "Parsing:beneficiaries" not in operation.tasks
        return deserialize(raw)
    with patch.object(cache.pickle, "loads", observed):
        assert LegalStore.from_files(payload, "second", operation=operation).import_timings["cache_hit"] == 1


@pytest.mark.parametrize("damage", ["identity", "magic"])
def test_incompatible_header_does_not_read_snapshot_body(damage):
    from unittest.mock import Mock
    LegalStore.from_files(required_payload(), "first")
    snapshot = cache.cache_directory() / "snapshot.bin"
    key = cache.cache_key(required_payload(), [])
    if damage == "identity":
        key = "0" * 64
    else:
        raw = snapshot.read_bytes()
        snapshot.write_bytes(b"!" + raw[1:])
    with snapshot.open("rb") as source:
        observed = Mock(wraps=source)
        observed.__enter__ = Mock(return_value=observed)
        observed.__exit__ = Mock(return_value=False)
        with patch.object(Path, "open", return_value=observed), patch.object(cache.pickle, "loads") as deserialize:
            assert cache.load(key) is None
        assert [call.args for call in observed.read.call_args_list] == [(len(cache.MAGIC) + 96,)]
        deserialize.assert_not_called()
