from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from backend.analytics import DataStore
from backend.legal_platform import EXCEL_DATE_FORMAT, EXPLORER_DATE_FORMAT, LegalStore


def test_legal_explorer_download_preserves_dates_and_uses_english_month_names():
    frame = pd.DataFrame({"Case ID": ["001", "002", "003", "004"],
                          "Date of Identification": ["01/10/2026", "2026-10-01", "invalid", ""]})
    store = LegalStore({"beneficiaries": frame}, "synthetic", [])
    sheet = load_workbook(BytesIO(store.explorer_export("beneficiaries")))["Filtered data"]
    for row in (2, 3):
        assert sheet.cell(row, 2).value == datetime(2026, 10, 1)
        assert sheet.cell(row, 2).number_format == EXPLORER_DATE_FORMAT
    assert sheet.cell(4, 2).value == "invalid"
    assert sheet.cell(5, 2).value is None
    assert sheet.cell(2, 1).value == "001"
    assert frame.iloc[0, 1] == "01/10/2026"


def test_workbook_explorer_download_uses_the_same_date_format():
    # Build the minimal raw-sheet store without importing private source workbooks.
    store = object.__new__(DataStore)
    store.raw_sheets = {"synthetic": pd.DataFrame({"Date": ["2026-10-01", pd.Timestamp("2026-10-02")]})}
    sheet = load_workbook(BytesIO(store.explorer_export("synthetic", "", [], None, "asc", [], "xlsx")))["Filtered data"]
    assert sheet.cell(2, 1).value == datetime(2026, 10, 1)
    assert sheet.cell(3, 1).value == datetime(2026, 10, 2)
    assert sheet.cell(2, 1).number_format == EXPLORER_DATE_FORMAT
    assert sheet.cell(3, 1).number_format == EXPLORER_DATE_FORMAT
    assert sheet.column_dimensions["A"].width >= 22


def test_raw_review_and_case_downloads_use_the_shared_date_format():
    from backend.test_legal_platform import required_payload
    store = LegalStore.from_files(required_payload(), "synthetic")
    downloads = [store.export("beneficiaries"), store.review_export("beneficiaries"), store.case_export("", {})]
    for download in downloads:
        workbook = load_workbook(BytesIO(download))
        dates = [cell for sheet in workbook for row in sheet for cell in row if cell.is_date]
        assert dates, "The fixture must exercise exported dates"
        assert all(cell.number_format == EXCEL_DATE_FORMAT for cell in dates)


def test_table_and_exclusion_downloads_format_dates_without_changing_numbers(monkeypatch):
    from types import SimpleNamespace
    monkeypatch.setenv("INTERSOS_DEFER_LEGAL_LOAD", "1")
    from backend import main
    response = main.table_workbook(main.TableWorkbookRequest(columns=["Date", "Count"],
        rows=[{"Date": "2026-10-01", "Count": 42}, {"Date": "01/10/2026", "Count": 12}]))
    sheet = load_workbook(BytesIO(response.body)).active
    assert sheet["A2"].value == sheet["A3"].value == datetime(2026, 10, 1)
    assert sheet["A2"].number_format == EXCEL_DATE_FORMAT
    assert sheet["B2"].value == 42
    monkeypatch.setattr(main, "duplicate_exclusions", SimpleNamespace(entries=lambda: [
        {"identifierValue": "001", "excludedAt": "2026-10-01T12:30:00+03:00"}]))
    sheet = load_workbook(BytesIO(main.export_duplicate_exclusions().body)).active
    assert sheet["G2"].value == datetime(2026, 10, 1)
    assert sheet["G2"].number_format == EXCEL_DATE_FORMAT
    assert sheet["D2"].value == "001"
