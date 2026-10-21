import queue
import threading
from copy import deepcopy

from openpyxl import load_workbook

from backend.kobo import KoboSettings, analyze
from backend.kobo_excel import ExportJobs
from backend.kobo_excel_prepare import worker_count, prepare_workbook
from backend.kobo_excel_runtime import ExportProgress, timeout_message
from backend.kobo_excel_source import response_layout, query_config
from backend.test_kobo import sample


def test_resource_budget_reserves_memory_and_limits_workers():
    gib = 1024**3
    assert worker_count(100, 32*gib, 16) == 1
    assert worker_count(1_000_000, 32*gib, 16) == 4
    assert worker_count(1_000_000, 10*gib, 16) == 2
    assert worker_count(1_000_000, 4*gib, 16) == 1
    assert worker_count(1_000_000, 32*gib, 2) == 1


def test_healthy_twenty_minute_export_is_not_cancelled_and_stalls_name_stage():
    assert timeout_message(1300, 0, 1290, 'queries') is None
    assert 'queries' in timeout_message(1300, 0, 600, 'queries')
    assert 'ten minutes' in timeout_message(1300, 0, 600, 'queries')
    assert 'sixty-minute' in timeout_message(3601, 0, 3600, 'saving')


def test_stage_telemetry_retains_legacy_fields(monkeypatch):
    import backend.kobo_excel_runtime as runtime
    ticks = iter([10, 12, 12, 15, 15])
    monkeypatch.setattr(runtime.time, 'monotonic', lambda: next(ticks))
    out = queue.Queue(); progress = ExportProgress(out)
    progress.total = 10; progress.begin('preparation', 5, 'Preparing tables')
    progress.completed = 2; progress.begin('pivots', 20, 'Creating pivots')
    out.get(); value = out.get()
    assert value['stageSeconds']['analysis'] == 2
    assert value['stageSeconds']['preparation'] == 3
    assert value['progress'] == 20 and value['message'] == 'Creating pivots'
    assert value['totalQuestions'] == 10 and value['completedQuestions'] == 2


def test_status_exposes_elapsed_and_last_progress_without_resetting_watchdog(monkeypatch):
    import backend.kobo_excel as native
    jobs = ExportJobs()
    jobs.jobs['test'] = {'id': 'test', 'state': 'running', 'progress': 50, 'message': 'Building',
                         'stage': 'pivots', 'started': 100, 'lastProgress': 150, 'completed': None}
    monkeypatch.setattr(native.time, 'monotonic', lambda: 180)
    result = jobs.status('test')
    assert result['elapsedSeconds'] == 80 and result['lastProgressSecondsAgo'] == 30
    assert jobs.jobs['test']['lastProgress'] == 150
    jobs.jobs.clear()


def test_file_preparation_keeps_literals_keys_and_editable_calculated_columns(sample, tmp_path):
    project = deepcopy(sample[0])
    field = project['form']['questions'][0]
    project['records'][0]['values'][field['id']] = '=literal text'
    settings = KoboSettings(hidden=['group/respondent_name'], hideEmpty=False)
    layouts = response_layout(project, settings)
    config = query_config(project, analyze(project, settings), settings, [], layouts)
    target = tmp_path/'prepared.xlsx'
    prepare_workbook(target, layouts, config, settings, queue.Queue())
    workbook = load_workbook(target)
    data = workbook['Data']; table = data.tables['KoboResponses000']
    assert data['D2'].value.lstrip("'") == '=literal text' and data['D2'].data_type == 's'
    assert data.max_row == 4
    assert 'Name of respondent' not in [cell.value for cell in data[1]]
    helper = next(col for col in table.tableColumns if col.name.endswith('answer'))
    assert '#This Row' in helper.calculatedColumnFormula.attr_text
    assert 'Technical' not in workbook.defined_names['KoboRaw000'].attr_text
    assert workbook['_Kobo choices'].sheet_state == 'hidden'
    workbook.close()
