import io
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from backend.legal_platform import LegalStore, clean_id, _find
from backend.linked_csv_enrichment import enrich_linked_csvs, PROFILE_FIELDS, ASSESSMENT_FIELDS, field_for
from backend.linked_csv_layouts import LEGACY_LAYOUTS
from backend.test_legal_platform import csv


NAME = "Name / الأسم"
DOB = "DoB / تأريخ الولاده"
PHONE = "Contact Number / رقم الاتصال"
ASSESSMENT_DATE = "Date of Assessment تاريخ التقييم"
CLOSURE_DATE = "Date of Assessment Closure تاريخ إغلاق التقييم"
SERVICE_DATE = "Date of Service Provision /  تاريخ بدء الخدمة"
STATUS = "Assessment Status حالة التقييم"


def payload(layout):
    data = {
        "beneficiaries": {"Case ID": ["001"], "Name (Filter Color Red)": ["سارة علي"],
                          DOB: ["01/02/1990"], PHONE: ["07501234567"], "Age": [36],
                          "Gender النوع الاجتماعي": ["Female"], "Community Type": ["Refugee"],
                          "Nationality الجنسية": ["Syrian"]},
        "assessments": {"Beneficiary ID": ["001"], "Assessment ID": ["A1"],
                        ASSESSMENT_DATE: ["02/03/2026"], CLOSURE_DATE: ["04/03/2026"], STATUS: ["Closed"],
                        "Is the beneficiary detained هل المستفيد موقوف": ["Yes"],
                        "Is it an immigration related charge? هل هو معتقل على اساس قانون الاقامة ؟": ["Yes"],
                        "Type of Released نوع الافراج": ["Released"],
                        "Date of the released or deported تاريخ الافراج او الترحيل": ["03/03/2026"]},
        "legalservices": {"Beneficiary ID": ["001"], "Assessment ID": ["A1"], "Service ID": ["S1"],
                          SERVICE_DATE: ["03/03/2026"], "Total Cost of Service": [100]},
        "legalfees": {"Beneficiary ID": ["001"], "Assessment ID": ["A1"], "Legal Service ID": ["S1"],
                      "Fee ID": ["F1"], "Amount Spent (IQD) / المبلغ المنفق": [100]},
        "followupslogbooks": {"Beneficiary ID": ["001"], "Assessment ID": ["A1"], "Service ID": ["S1"],
                              "Follow-ups & Logbook ID": ["L1"], "Follow up objective الهدف من التابعة": ["Check"]},
    }
    if layout in {"old", "mixed"}:
        for dataset in LEGACY_LAYOUTS:
            if layout == "mixed" and dataset in {"assessments", "legalfees"}:
                continue
            data[dataset].update({NAME: ["Outdated name"], DOB: ["01/01/1980"], PHONE: ["00000000000"]})
        data["legalservices"].update({STATUS: ["Open"], ASSESSMENT_DATE: ["01/01/2025"]})
        data["followupslogbooks"].update({ASSESSMENT_DATE: ["01/01/2025"], SERVICE_DATE: ["01/01/2025"]})
    return {name: csv(**columns) for name, columns in data.items()}


@pytest.mark.parametrize("layout", ["old", "reduced", "mixed"])
def test_profiles_and_parent_fields_reach_views_filters_and_excel(layout):
    store = LegalStore.from_files(payload(layout), "test")
    for dataset in LEGACY_LAYOUTS:
        frame = store.frames[dataset]
        assert len(frame) == 1
        assert frame.iloc[0][NAME] == "سارة علي"
        assert frame.iloc[0][PHONE] == "07501234567"
        assert frame.iloc[0][DOB] == pd.Timestamp("1990-02-01")
        assert store.explorer(dataset, search="سارة", filters={NAME: ["سارة علي"]})["total"] == 1
        for workbook_bytes in (store.export(dataset), store.explorer_export(dataset, search="سارة")):
            sheet = load_workbook(io.BytesIO(workbook_bytes)).active
            headers = [cell.value for cell in sheet[1]]
            row = dict(zip(headers, [cell.value for cell in sheet[2]]))
            assert row[NAME] == "سارة علي"
            assert row[PHONE] == "07501234567"
            assert isinstance(row[DOB], datetime)
            assert sheet.cell(2, headers.index(DOB) + 1).is_date
    services = store.explorer("legalservices")["rows"][0]
    assert services[STATUS] == "Closed"
    assert services[ASSESSMENT_DATE] == "02/03/2026"
    assert services[CLOSURE_DATE] == "04/03/2026"
    for field in ASSESSMENT_FIELDS:
        column = next(c for c in store.frames["assessments"] if field_for(c, ASSESSMENT_FIELDS) == field)
        assert store.frames["legalservices"].iloc[0][column] == store.frames["assessments"].iloc[0][column]
    followup = store.explorer("followupslogbooks")["rows"][0]
    assert followup[ASSESSMENT_DATE] == "02/03/2026"
    assert followup[SERVICE_DATE] == "03/03/2026"
    assert store.frames["legalfees"].iloc[0]["Amount Spent (IQD) / المبلغ المنفق"] == "100"


def enrich(frames):
    warnings = []
    dates = {name: [] for name in frames}
    enrich_linked_csvs(frames, dates, warnings, clean_id=clean_id, checkpoint=lambda: None)
    return warnings


def test_blank_authoritative_values_and_unmatched_links_do_not_multiply_rows():
    frames = {
        "beneficiaries": pd.DataFrame({"Case ID": ["001", "001", ""],
                                       "Name (Filter Color Red)": [None, "Duplicate", "Blank ID person"]}),
        "assessments": pd.DataFrame({"Assessment ID": ["A1", "A2", "A3"], "Beneficiary ID": ["001", "missing", ""],
                                     NAME: ["Old", "Unmatched", "No ID"], "Nationality Stateless": ["Yes"] * 3}),
    }
    warnings = enrich(frames)
    assert len(frames["assessments"]) == 3
    assert pd.isna(frames["assessments"].iloc[0][NAME])
    assert frames["assessments"][NAME].iloc[1:].tolist() == ["Unmatched", "No ID"]
    assert frames["assessments"]["Nationality Stateless"].tolist() == ["Yes"] * 3
    assert any("duplicate" in warning for warning in warnings)
    assert any("2 rows" in warning for warning in warnings)


def test_aliases_missing_source_fields_and_existing_column_order():
    frames = {
        "beneficiaries": pd.DataFrame({"Case ID": ["001"], "Name (Filter Color Red)": ["Current"],
                                       "Governorate  / المحافظه": ["Ninewa"]}),
        "assessments": pd.DataFrame({"Custom": ["Keep"], "Assessment ID": ["A1"], "Beneficiary ID": [" 001 "],
                                     "Name": ["Old"], "DoB": ["Legacy DOB"], "Governorate / المحافظه": ["Old"],
                                     "Nationality Stateless": ["Yes"]}),
    }
    original = list(frames["assessments"])
    enrich(frames)
    result = frames["assessments"]
    assert result.iloc[0]["Name"] == "Current"
    assert NAME not in result
    assert result.iloc[0]["DoB"] == "Legacy DOB"
    assert result.iloc[0]["Governorate / المحافظه"] == "Ninewa"
    assert [column for column in result if column in original] == original
    assert result.iloc[0]["Custom"] == "Keep"
    assert field_for("Name of the reporting person", PROFILE_FIELDS) is None
    assert field_for("Date of Assessment Closure", ("Date of Assessment",)) is None
    assert _find(["Age Gender Group", "Gender النوع الاجتماعي"], "Gender") == "Gender النوع الاجتماعي"
    assert _find([CLOSURE_DATE, ASSESSMENT_DATE], "Date of Assessment") == ASSESSMENT_DATE


def test_refresh_rebuilds_linked_fields_without_mutating_csv_bytes():
    data = payload("reduced")
    snapshot = data.copy()
    first = LegalStore.from_files(data, "first")
    assert data == snapshot
    data["beneficiaries"] = csv(**{"Case ID": ["001"], "Name (Filter Color Red)": ["Updated"]})
    second = LegalStore.from_files(data, "second")
    assert first.explorer("legalservices")["rows"][0][NAME] == "سارة علي"
    assert second.explorer("legalservices")["rows"][0][NAME] == "Updated"
    assert first.revision != second.revision


def test_empty_datasets_restore_columns():
    frames = {
        "beneficiaries": pd.DataFrame(columns=["Case ID", "Name (Filter Color Red)"]),
        "assessments": pd.DataFrame(columns=["Assessment ID", "Beneficiary ID"]),
        "legalservices": pd.DataFrame(columns=["Service ID", "Assessment ID", "Beneficiary ID"]),
    }
    assert enrich(frames) == []
    assert all(frame.empty for frame in frames.values())
    assert NAME in frames["legalservices"]
    assert STATUS in frames["legalservices"]


def test_blank_assessment_values_clear_legacy_service_fields():
    data = payload("old")
    assessment = pd.read_csv(io.BytesIO(data["assessments"]), dtype=object)
    assessment[STATUS] = ""
    assessment[ASSESSMENT_DATE] = ""
    data["assessments"] = assessment.to_csv(index=False).encode("utf-8")
    store = LegalStore.from_files(data, "blank assessment")
    row = store.explorer("legalservices")["rows"][0]
    assert row[STATUS] == ""
    assert row[ASSESSMENT_DATE] == ""
    sheet = load_workbook(io.BytesIO(store.export("legalservices"))).active
    columns = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, columns.index(ASSESSMENT_DATE) + 1).value is None
