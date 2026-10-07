from __future__ import annotations

import csv
import hashlib
import html
import io
import itertools
import json
import os
import re
import threading
from collections import defaultdict
from datetime import date
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from typing import Any, Callable, Literal

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from .issue_message_details import arabic_detail
from .issue_affected_fields import affected_fields
from .issue_field_values import IssueFieldValues
from .issue_duplicate_matches import IssueDuplicateMatches
from .lawyer_contact_defaults import DEFAULT_LAWYER_CONTACTS
from .file_security import validate_xlsx_archive
from .legal_platform import REVIEW_FINDING_ARABIC, _find

DATASETS = {"beneficiaries": "Beneficiaries Review", "assessments": "Assessments Review", "legalservices": "Legal Services Review", "legalhotlines": "Hotline Review", "awareness": "Awareness Review"}
UNASSIGNED = "Unassigned"


def normalize_name(value: str) -> str:
    return " ".join(value.split()).casefold()


def valid_email(value: str) -> bool:
    return bool(re.fullmatch(r"[^\s@<>;,\r\n]+@[^\s@<>;,\r\n]+\.[^\s@<>;,\r\n]+", value)) and len(value) <= 254


def whatsapp_phone(value: str) -> str:
    if not re.fullmatch(r"\+?[0-9\s()-]+", value.strip()):
        raise ValueError("Enter a valid WhatsApp number. Iraq (+964) is added automatically for local numbers.")
    digits = re.sub(r"\D", "", value)
    if digits.startswith("00"):
        digits = digits[2:]
    if not value.strip().startswith(("+", "00")) and re.fullmatch(r"0?7[0-9]{9}", digits):
        digits = "964" + digits.removeprefix("0")
    if not re.fullmatch(r"[1-9][0-9]{7,14}", digits):
        raise ValueError("Enter a valid WhatsApp number. Iraq (+964) is added automatically for local numbers.")
    return digits


class Contacts:
    def __init__(self, path: Path | None = None):
        self.defaults = DEFAULT_LAWYER_CONTACTS if path is None else {}
        self.path = path or Path(os.getenv("LOCALAPPDATA", str(Path.home()))) / "INTERSOS Legal Platform" / "lawyer-contacts.json"
        self.lock = threading.RLock()

    def read(self) -> dict[str, str]:
        with self.lock:
            try:
                data = json.loads(self.path.read_text(encoding="utf-8"))
                return {**self.defaults, **(data if isinstance(data, dict) else {})}
            except FileNotFoundError:
                return dict(self.defaults)

    def for_lawyers(self, lawyers: list[str]) -> dict[str, str]:
        contacts = self.read()
        normalized: dict[str, list[str]] = defaultdict(list)
        for name in contacts:
            normalized[normalize_name(name)].append(name)
        result = {}
        for lawyer in lawyers:
            if lawyer in contacts:
                result[lawyer] = contacts[lawyer]
            else:
                matches = normalized.get(normalize_name(lawyer), [])
                if len(matches) == 1:
                    result[lawyer] = contacts[matches[0]]
        return result

    def save(self, contacts: dict[str, str]):
        if any(email and not valid_email(email) for email in contacts.values()):
            raise ValueError("Enter a valid email address for each contact.")
        with self.lock:
            data = self.read()
            data.update({name: email.strip() for name, email in contacts.items()})
            self.path.parent.mkdir(parents=True, exist_ok=True)
            temporary = self.path.with_suffix(".tmp")
            temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            temporary.replace(self.path)

    def read_whatsapp(self) -> dict[str, str]:
        return Contacts(self.path.with_suffix('.whatsapp.json')).read()

    def save_whatsapp(self, numbers: dict[str, str]):
        normalized = {name: whatsapp_phone(value) if value.strip() else '' for name, value in numbers.items()}
        with self.lock:
            path = self.path.with_suffix('.whatsapp.json')
            data = self.read_whatsapp()
            data.update(normalized)
            path.parent.mkdir(parents=True, exist_ok=True)
            temporary = path.with_suffix('.tmp')
            temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
            temporary.replace(path)


def collect_findings(store) -> tuple[list[str], list[dict[str, Any]]]:
    field_values = IssueFieldValues(store.frames)
    names = set()
    for dataset, frame in store.frames.items():
        column = _find(list(frame.columns), "Lawyers", "Lawyer")
        if dataset == "legalhotlines":
            column = next((column for column in frame.columns if normalize_name(str(column)) in {"lawyer", "lawyers"}), None)
        if column:
            names.update(str(value).strip() for value in frame[column].fillna("") if str(value).strip())
    canonical: dict[str, list[str]] = defaultdict(list)
    for name in sorted(names):
        canonical[normalize_name(name)].append(name)
    rows = []
    duplicate_matches = IssueDuplicateMatches(store, canonical, normalize_name)
    duplicate_rules={"Possible duplicate name","Possible duplicate contact and name"}
    core_available = all(name in store.frames for name in ("beneficiaries", "assessments", "legalservices"))
    for dataset in DATASETS:
        if dataset not in store.frames:
            continue
        if dataset in {"beneficiaries", "assessments", "legalservices"} and not core_available:
            continue
        # Match review-page defaults; query each rule to preserve exact duplicate semantics.
        summary = store.review(dataset, page_size=1, allow_name_variations=dataset == "legalhotlines")
        duplicate_rows={}
        if dataset=="beneficiaries" and hasattr(store,"_name_match_flags"):
            excluded=store._excluded_case_ids("Possible duplicate name")
            combined={}
            for item in store._name_match_flags(exact_only=True,excluded_case_ids=excluded)+store._name_match_flags(allow_variations=True,excluded_case_ids=excluded):
                key=item["row"]
                if key not in combined: combined[key]=dict(item)
                peers={peer["row"]:peer for peer in combined[key].get("duplicateMatches",[])}
                peers.update({peer["row"]:peer for peer in item.get("duplicateMatches",[])})
                combined[key]["duplicateMatches"]=sorted(peers.values(),key=lambda peer:peer["row"])
            duplicate_rows["Possible duplicate name"]=list(combined.values())
            duplicate_rows["Possible duplicate contact and name"]=store._contact_name_match_flags(excluded_case_ids=store._excluded_case_ids("Possible duplicate contact and name"))
            for rule,items in duplicate_rows.items():
                eligible=[item for item in items if not store._is_excluded(item)]
                eligible_rows={item["row"] for item in eligible}
                for item in eligible:
                    item["duplicateMatches"]=[peer for peer in item.get("duplicateMatches",[]) if peer["row"] in eligible_rows]
                duplicate_rows[rule]=[item for item in eligible if item["duplicateMatches"]]
        for rule in dict.fromkeys([*summary["rules"],*duplicate_rows]):
            page = 1
            while True:
                if rule in duplicate_rows:
                    result={"rows":duplicate_rows[rule],"total":len(duplicate_rows[rule])}
                else:
                    result = store.review(dataset, rule=rule, page=page, page_size=5000,
                                          allow_name_variations=dataset == "legalhotlines",
                                          exact_matches_only=dataset != "legalhotlines")
                for flag in result["rows"]:
                    if hasattr(store,"_is_excluded") and store._is_excluded(flag): continue
                    raw_name = str(flag.get("lawyer", "")).strip()
                    matches = canonical.get(normalize_name(raw_name), [])
                    if dataset != "legalhotlines" and raw_name and not matches:
                        names.add(raw_name)
                        matches = [raw_name]
                    lawyer = matches[0] if len(matches) == 1 else UNASSIGNED
                    safe = {key: flag.get(key, "") for key in ("dataset", "rule", "severity", "row", "recordId", "caseId", "assessmentId", "serviceId", "hotlineId", "awarenessId", "detail", "action", "project", "location", "assessmentStatus", "requestForClosedStatus", "linkedServiceCount", "linkedServiceStatuses")}
                    safe.update(lawyer=lawyer, ruleArabic=REVIEW_FINDING_ARABIC.get(rule, rule), reviewPage=DATASETS[dataset])
                    if rule in duplicate_rules:
                        safe["duplicateMatches"]=[{**peer,"lawyer":(canonical.get(normalize_name(peer.get("lawyer", "")),[]) or [peer.get("lawyer") or UNASSIGNED])[0]} for peer in flag.get("duplicateMatches",[])]
                        safe["detail"]=f"Possible duplicate with {len(safe['duplicateMatches'])} matching record(s). Verify the case references and responsible lawyers below."
                        safe["matchingCases"]="; ".join(f"Case {peer['caseId'] or 'row '+str(peer['row'])} - {peer['lawyer']} ({peer['matchType']})" for peer in safe["duplicateMatches"])
                        safe["action"]="Please verify these possible duplicates with the responsible lawyer(s). If confirmed, notify the IM Officer to arrange deletion of the redundant case and update the physical file and PR record with the correct Case ID."
                    if rule=="Open assessment with all services closed":
                        requested=str(flag.get("requestForClosedStatus", "")).strip().casefold()=="yes"
                        safe["rule"]=rule+(" (Closure requested by lawyer)" if requested else " (Closure not requested by lawyer)")
                        safe["ruleArabic"] += " (طلب المحامي الإغلاق)" if requested else " (لم يطلب المحامي الإغلاق)"
                    if flag.get('duplicateMatchRows') or flag.get('duplicateMatches'):
                        safe['duplicateMatches'] = duplicate_matches.matches(flag)
                        safe['matchingCases'] = '\n'.join(peer['description'] for peer in safe['duplicateMatches'])
                    safe["detailArabic"]=arabic_detail(str(safe.get("detail", "")))
                    safe["affectedFields"]=affected_fields(dataset, safe["rule"], str(safe.get("detail", "")))
                    safe["affectedFieldValues"]=field_values.values(flag, safe["affectedFields"])
                    safe["id"] = hashlib.sha256(json.dumps(safe, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()
                    rows.append(safe)
                if page * 5000 >= result["total"]:
                    break
                page += 1
    return sorted(names), rows


class FindingsQuery(BaseModel):
    page: int = Field(1, ge=1)
    pageSize: int = Field(500, ge=1, le=5000)
    lawyer: str = ""
    filters: dict[str, str] = Field(default_factory=dict)


class ContactUpdate(BaseModel):
    contacts: dict[str, str]
    whatsappNumbers: dict[str, str] = Field(default_factory=dict)


class DraftRequest(BaseModel):
    language: Literal["en", "ar", "bilingual"] = "bilingual"
    revision: str
    lawyer: str
    ids: list[str]
    assignments: dict[str, str] = Field(default_factory=dict)
    subject: str = "Legal data review - action required"
    introduction: str = "Please review the findings below, correct the relevant records, and confirm completion. Thank you for your cooperation."
    introductionArabic: str = "يرجى مراجعة الملاحظات أدناه وتصحيح السجلات ذات الصلة وتأكيد إتمام التصحيحات. شكراً لتعاونكم."
    deadline: str = ""
    signature: str = ""


def render_draft(request: DraftRequest, email: str, rows: list[dict]) -> tuple[str, str, bytes]:
    if not valid_email(email) or any(char in request.subject for char in "\r\n"):
        raise ValueError("Invalid recipient or subject.")
    if request.deadline:
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", request.deadline): raise ValueError()
            date.fromisoformat(request.deadline)
        except ValueError as exc: raise ValueError("Invalid correction deadline.") from exc
    esc = lambda value: html.escape(str(value if value not in (None, "") else "Not available"))
    # Retain the request language field for older clients. Only finding details are translated.
    def label(en, ar):
        return en
    text=[]; blocks=[]
    def paragraph(value, rtl=False):
        text.append(value)
        direction=' dir="rtl"' if rtl else ''
        blocks.append(f'<p{direction} style="white-space:pre-wrap;line-height:1.7">{esc(value)}</p>')
    paragraph(f"Dear {request.lawyer},")
    if request.introduction: paragraph(request.introduction)
    if request.deadline: paragraph(f"{label('Correction deadline','الموعد النهائي للتصحيح')}: {request.deadline}")
    groups=defaultdict(list)
    for row in rows: groups[(row['reviewPage'],row['rule'])].append(row)
    number=0
    for (page,rule),findings in groups.items():
        title=label(rule,findings[0]['ruleArabic'])
        paragraph(f"{page}: {title}")
        identifiers=[('Case ID','رقم الحالة','caseId'),('Assessment ID','رقم التقييم','assessmentId'),('Service ID','رقم الخدمة','serviceId'),('Hotline ID','رقم الخط الساخن','hotlineId'),('Awareness ID','رقم التوعية','awarenessId')]
        columns=[(label(en,ar),key) for en,ar,key in identifiers if any(item.get(key) for item in findings)]
        if not columns or any(not any(item.get(key) for _,_,key in identifiers) for item in findings): columns.append((label('Record','السجل'),'recordId'))
        identifier_labels = {heading.casefold() for heading, _ in columns}
        if 'case id' in identifier_labels:
            identifier_labels.add('beneficiary id')
        fields = list(dict.fromkeys(field for item in findings for field in item.get('affectedFields', []) if field.casefold() not in identifier_labels))
        columns.extend((field, ('field', field)) for field in fields)
        if any(item.get('matchingCases') for item in findings): columns.append((label('Matching records and lawyers','السجلات والمحامون المطابقون'),'matchingCases'))
        if rule.startswith('Open assessment with all services closed'): columns.append((label('Closure context','تفاصيل الإغلاق'),'closureContext'))
        columns.append(('Finding detail','detail'))
        columns.append(('Finding detail (Arabic)','arabicDetail'))
        direction=''
        blocks.append(f'<table{direction} cellpadding="0" cellspacing="0" border="0" style="width:100%;min-width:780px;border-collapse:collapse;font-size:13px;margin:12px 0 24px;border:1px solid #d4e0eb"><thead><tr>')
        headers=[('#','number'),*columns]
        for heading,_ in headers: blocks.append(f'<th scope="col" style="padding:11px 10px;background:#eaf1f8;color:#183c60;border:1px solid #d4e0eb;text-align:start;font-weight:700">{esc(heading)}</th>')
        blocks.append('</tr></thead><tbody>')
        for index,row in enumerate(findings):
            number+=1
            translated=row.get('detailArabic') or arabic_detail(str(row.get('detail','')))
            values={**row,'number':number,'recordId':row.get('recordId') or f"Row {row.get('row','')}", 'arabicDetail':translated or str(row.get('detail') or 'Not available')}
            if rule.startswith('Open assessment with all services closed'):
                values['closureContext']='\n'.join(f"{label(en,ar)}: {row.get(key) if row.get(key) not in (None,'') else 'Not available'}" for en,ar,key in [('Assessment status','حالة التقييم','assessmentStatus'),('Closure request','طلب الإغلاق','requestForClosedStatus'),('Linked services','الخدمات المرتبطة','linkedServiceCount'),('Service statuses','حالات الخدمات','linkedServiceStatuses')])
            background='#ffffff' if index%2==0 else '#f5f8fc'
            blocks.append(f'<tr style="background:{background}">')
            text.append(f"{number}. {title}")
            for heading,key in headers:
                value=row.get('affectedFieldValues', {}).get(key[1]) if isinstance(key, tuple) else values.get(key)
                text.append(f"{heading}: {value if value not in (None,'') else 'Not available'}")
                cell_direction=' dir="ltr"' if key in {item[2] for item in identifiers} else ' dir="auto"'
                display_value=re.sub(r',\s*', ',\n',str(value)) if key in {item[2] for item in identifiers} and value is not None else value
                blocks.append(f'<td{cell_direction} style="padding:11px 10px;vertical-align:top;line-height:1.65;border:1px solid #d4e0eb;white-space:pre-wrap;overflow-wrap:anywhere">{esc(display_value)}</td>')
            blocks.append('</tr>')
        blocks.append('</tbody></table>')
    paragraph('Please confirm once the corrections are complete. If any finding needs clarification, let us know. Thank you for your cooperation.')
    body='<html><head><meta charset="utf-8"></head><body style="font-family:Arial,sans-serif;color:#172b42;max-width:1200px;margin:auto;padding:24px;background:#f7f9fc">'+''.join(blocks)+'</body></html>'
    plain='\n\n'.join(text)
    message = EmailMessage(policy=SMTP)
    message["To"] = email
    message["Subject"] = request.subject
    message["X-Unsent"] = "1"
    message.set_content(plain)
    message.add_alternative(body, subtype="html")
    return body, plain, message.as_bytes()


def make_router(get_store: Callable, contacts: Contacts | None = None) -> APIRouter:
    router = APIRouter(prefix="/api/legal/send-issues")
    registry = contacts or Contacts()
    cache: dict[str, Any] = {}
    lock = threading.RLock()

    def snapshot():
        store = get_store()
        with lock:
            if cache.get("revision") != store.revision:
                revision = store.revision
                lawyers, rows = collect_findings(store)
                if revision != store.revision:
                    raise HTTPException(409, "Review data changed. Refresh Send Issues and select findings again.")
                cache.update(revision=revision, lawyers=lawyers, rows=rows)
            return dict(cache)

    @router.post("/findings")
    def findings(query: FindingsQuery):
        data = snapshot()
        rows = [row for row in data["rows"] if (not query.lawyer or row["lawyer"] == query.lawyer) and all(not value or row.get(key) == value for key, value in query.filters.items() if key in {"dataset", "rule", "severity", "project", "location"})]
        start = (query.page - 1) * query.pageSize
        return {"revision": data["revision"], "lawyers": data["lawyers"], "rows": rows[start:start + query.pageSize], "total": len(rows), "contacts": registry.for_lawyers(data["lawyers"]), "whatsappNumbers": registry.read_whatsapp()}

    @router.get("/contacts")
    def read_contacts():
        lawyers = snapshot()["lawyers"]
        return {"contacts": registry.for_lawyers(lawyers), "whatsappNumbers": registry.read_whatsapp(), "lawyers": lawyers}

    @router.put("/contacts")
    def save_contacts(request: ContactUpdate):
        known = set(snapshot()["lawyers"])
        if any(name not in known for name in set(request.contacts) | set(request.whatsappNumbers)):
            raise HTTPException(400, "Choose lawyer names from the current data.")
        try:
            numbers = {name: whatsapp_phone(value) if value.strip() else '' for name, value in request.whatsappNumbers.items()}
            registry.save({name: email.strip() for name, email in request.contacts.items()})
            registry.save_whatsapp(numbers)
        except ValueError as exc: raise HTTPException(400, str(exc)) from exc
        return {"contacts": registry.for_lawyers(snapshot()["lawyers"]), "whatsappNumbers": registry.read_whatsapp()}

    @router.post("/contacts/preview")
    async def preview_contacts(file: UploadFile = File(...), nameColumn: str = Form(""), emailColumn: str = Form("")):
        try:
            raw = await file.read(5 * 1024 * 1024 + 1)
            if len(raw) > 5 * 1024 * 1024: raise ValueError("Contact files must be 5 MB or smaller.")
            suffix = Path(file.filename or "").suffix.lower()
            if suffix == ".xlsx":
                validate_xlsx_archive(raw)
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                try: values = list(itertools.islice(workbook.active.iter_rows(values_only=True), 10002))
                finally: workbook.close()
            elif suffix == ".csv":
                source = raw.decode("utf-8-sig")
                try: dialect = csv.Sniffer().sniff(source[:4096], delimiters=",;\t")
                except csv.Error: dialect = csv.excel
                values = list(itertools.islice(csv.reader(io.StringIO(source), dialect), 10002))
            else: raise ValueError("Upload an .xlsx or UTF-8 .csv contact list.")
            if not values: raise ValueError("The contact file is empty.")
            columns = [str(value or "").strip() for value in values[0]]
            if len(set(columns)) != len(columns): raise ValueError("Contact column names must be unique.")
            if len(values) > 10001: raise ValueError("Contact files may contain at most 10,000 rows.")
            if not nameColumn or not emailColumn: return {"columns": columns, "rows": []}
            if nameColumn == emailColumn: raise ValueError("Choose different name and email columns.")
            ni, ei = columns.index(nameColumn), columns.index(emailColumn)
            names = snapshot()["lawyers"]
            rows = []
            for index, value in enumerate(values[1:], 2):
                name = str(value[ni] or "").strip() if ni < len(value) else ""
                email = str(value[ei] or "").strip() if ei < len(value) else ""
                if not name and not email: continue
                matches = [candidate for candidate in names if normalize_name(candidate) == normalize_name(name)]
                rows.append({"row": index, "name": name, "email": email, "lawyer": matches[0] if len(matches) == 1 else "", "status": "Invalid email" if not valid_email(email) else "Matched" if len(matches) == 1 else "Ambiguous name" if matches else "Unmatched name"})
            return {"columns": columns, "rows": rows}
        except (ValueError, UnicodeError) as exc: raise HTTPException(400, str(exc)) from exc
        finally: await file.close()

    def prepare(request: DraftRequest):
        data = snapshot()
        if request.revision != data["revision"]:
            raise HTTPException(409, "Review data changed. Refresh Send Issues and select findings again.")
        if request.lawyer not in data["lawyers"]: raise HTTPException(400, "Choose a lawyer from the current data.")
        index = {row["id"]: row for row in data["rows"]}
        if not request.ids or len(set(request.ids)) != len(request.ids): raise HTTPException(400, "Select distinct findings for this lawyer.")
        rows = []
        for identifier in request.ids:
            row = index.get(identifier)
            if not row or (request.assignments.get(identifier, "") if row["lawyer"] == UNASSIGNED else row["lawyer"]) != request.lawyer:
                raise HTTPException(400, "Every selected finding must belong to this lawyer.")
            rows.append(row)
        recipient = registry.for_lawyers(data["lawyers"]).get(request.lawyer, "")
        try: body, plain, eml = render_draft(request, recipient, rows)
        except ValueError as exc: raise HTTPException(400, str(exc)) from exc
        if get_store().revision != request.revision: raise HTTPException(409, "Review data changed. Refresh Send Issues.")
        return recipient, body, plain, eml

    @router.post("/preview")
    def preview(request: DraftRequest):
        recipient, body, plain, _ = prepare(request)
        return {"recipient": recipient, "html": body, "text": plain, "subject": request.subject}

    @router.post("/draft")
    def draft(request: DraftRequest):
        _, _, _, eml = prepare(request)
        return Response(eml, media_type="message/rfc822", headers={"Content-Disposition": 'attachment; filename="legal-review-draft.eml"'})

    return router
