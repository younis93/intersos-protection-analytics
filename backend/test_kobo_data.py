import copy

from backend.kobo import KoboSettings
from backend.kobo_data import data_page
from backend.test_kobo import sample
from backend.test_kobo_preview import preview_client


def test_data_grid_labels_zero_missing_search_and_sort_without_saving(preview_client):
    client, store, pid = preview_client
    before = store.get(pid)
    body = {'settings': {**before['settings'], 'language': 'en'}, 'columns': ['group/needs', 'group/age'], 'pageSize': 1, 'sortColumn': 'group/age', 'sortDirection': 'desc'}
    result = client.post(f'/api/kobo/projects/{pid}/data', json=body)
    assert result.status_code == 200
    assert result.json()['total'] == 3
    assert result.json()['rows'][0]['values']['group/age'] == '20.0'
    second = client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'page': 2}).json()
    assert float(second['rows'][0]['values']['group/age']) == 0
    third = client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'page': 3}).json()
    assert third['rows'][0]['values']['group/age'] == ''
    arabic = client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'settings': {**body['settings'], 'language': 'ar'}, 'search': 'المياه النظيفة'}).json()
    assert arabic['matched'] == 1
    assert 'المياه النظيفة' in arabic['rows'][0]['values']['group/needs']
    assert store.get(pid) == before


def test_data_grid_filters_and_validation(preview_client):
    client, store, pid = preview_client
    body = {'settings': {**store.get(pid)['settings'], 'filters': [{'field': 'group/gender', 'kind': 'category', 'values': ['f']}]}, 'columns': ['group/gender']}
    result = client.post(f'/api/kobo/projects/{pid}/data', json=body).json()
    assert result['total'] == result['matched'] == 2
    assert client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'columns': ['missing']}).status_code == 400
    assert client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'sortColumn': 'group/age'}).status_code == 400
    assert client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'scope': 'missing'}).status_code == 400
    assert client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'pageSize': 1000}).status_code == 422


def test_repeat_rows_stay_separate_and_keep_parent_filters(sample):
    project = copy.deepcopy(sample[0])
    child = {**project['form']['questions'][2], 'id': 'visits/count', 'repeat': 'visits'}
    project['form']['questions'].append(child)
    parent = project['records'][0]
    project['records'].append({'id': 'child', 'parent': parent['id'], 'scope': 'visits', 'row': 8, 'values': {'visits/count': 2}})
    settings = KoboSettings(filters=[{'field': 'group/gender', 'kind': 'category', 'values': ['f']}])
    assert data_page(project, settings, '', ['group/age'])['total'] == 2
    result = data_page(project, settings, 'visits', ['visits/count'])
    assert result['total'] == 1
    assert result['rows'][0]['values'] == {'visits/count': '2'}


def read_export(response):
    from io import BytesIO
    from openpyxl import load_workbook
    assert response.status_code == 200, response.text if response.status_code != 200 else ''
    return load_workbook(BytesIO(response.content))


def test_data_export_all_rows_selected_order_and_empty(preview_client):
    client, store, pid = preview_client
    before = store.get(pid)
    body = {'settings': {**before['settings'], 'language': 'en'}, 'columns': ['group/age', 'group/needs'], 'sortColumn': 'group/age', 'sortDirection': 'desc', 'pageSize': 1, 'page': 2}
    grid = client.post(f'/api/kobo/projects/{pid}/data', json={**body, 'page': 1, 'pageSize': 100}).json()
    export = lambda payload: read_export(client.post(f'/api/kobo/projects/{pid}/data/export', json=payload))
    workbook = export(body)
    rows = list(workbook.active.values)
    assert len(rows) == 4  # Ignores pagination for a complete filtered export.
    assert rows[1][1] == grid['rows'][0]['values']['group/age']
    assert rows[2][1] == grid['rows'][1]['values']['group/age']
    selected = export({**body, 'selectedIds': [grid['rows'][1]['id'], grid['rows'][0]['id']]})
    assert list(selected.active.values)[1:] == rows[1:3]
    assert len(list(export({**body, 'selectedIds': []}).active.values)) == 1
    assert len(list(export({**body, 'search': 'no such response'}).active.values)) == 1
    assert store.get(pid) == before


def test_data_export_filters_language_and_validation(preview_client):
    client, store, pid = preview_client
    body = {'settings': {**store.get(pid)['settings'], 'language': 'ar', 'filters': [{'field': 'group/gender', 'kind': 'category', 'values': ['f']}]}, 'columns': ['group/needs', 'group/age'], 'search': 'المياه النظيفة'}
    response = client.post(f'/api/kobo/projects/{pid}/data/export', json=body)
    workbook = read_export(response)
    assert workbook.active.sheet_view.rightToLeft
    assert len(list(workbook.active.values)) == 2
    assert 'المياه النظيفة' in list(workbook.active.values)[1][1]
    for changes in [{'columns': ['missing']}, {'scope': 'missing'}, {'sortColumn': 'group/gender'}]:
        assert client.post(f'/api/kobo/projects/{pid}/data/export', json={**body, **changes}).status_code == 400


def test_data_export_many_rows_duplicate_labels_and_literal_text(sample):
    from io import BytesIO
    from openpyxl import load_workbook
    from backend.kobo_data import data_workbook
    project = copy.deepcopy(sample[0])
    questions = project['form']['questions'][:2]
    questions[0]['labels'] = questions[1]['labels'] = {'en': 'Source row'}
    project['form']['questions'] = questions
    columns = [q['id'] for q in questions]
    project['records'] = [{'id': str(i), 'scope': '', 'parent': '', 'row': i + 2, 'values': {columns[0]: '=1+1', columns[1]: 'text'}} for i in range(125)]
    workbook = load_workbook(BytesIO(data_workbook(project, KoboSettings(language='en'), '', columns)))
    rows = list(workbook.active.values)
    assert len(rows) == 126
    assert len(set(rows[0])) == 3
    assert all(q['name'] in rows[0][i + 1] for i, q in enumerate(questions))
    assert workbook.active['B2'].value == '=1+1'
    assert workbook.active['B2'].data_type == 's'


def test_data_export_repeat_selection_and_parent_filters(sample):
    from io import BytesIO
    from openpyxl import load_workbook
    from backend.kobo_data import data_workbook
    project = copy.deepcopy(sample[0])
    project['form']['questions'].append({**project['form']['questions'][2], 'id': 'visits/count', 'repeat': 'visits'})
    parent = project['records'][0]
    project['records'].append({'id': 'child', 'parent': parent['id'], 'scope': 'visits', 'row': 8, 'values': {'visits/count': 0}})
    settings = KoboSettings(filters=[{'field': 'group/gender', 'kind': 'category', 'values': ['f']}])
    workbook = load_workbook(BytesIO(data_workbook(project, settings, 'visits', ['visits/count'], selected_ids=['child', parent['id']])))
    assert list(workbook.active.values)[1:] == [(8, '0')]
