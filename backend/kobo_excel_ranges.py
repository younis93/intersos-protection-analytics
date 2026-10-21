"""Use one COM dispatch per worksheet range, rather than chained Cells objects."""
from openpyxl.utils import get_column_letter


def excel_cell(sheet, row, column):
    return sheet.Range(f'{get_column_letter(column)}{row}')


def excel_range(sheet, first_row, first_column, last_row, last_column):
    return sheet.Range(f'{get_column_letter(first_column)}{first_row}:{get_column_letter(last_column)}{last_row}')
