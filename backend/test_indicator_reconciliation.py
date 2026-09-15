from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook

from backend.indicator_reconciliation import IndicatorMasterWorkbook, build_reconciliation_workbook, reconcile


def workbook_bytes(include_reporting_tool=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporting Tool" if include_reporting_tool else "Other"
    if include_reporting_tool:
        sheet["D1"] = "2026-01-01"
        sheet["A18"] = "# of persons receiving legal counselling for civil documentation"
        sheet["B18"] = "AMAL Camp"
        sheet["C18"] = "AMAL Camp"
        sheet["AD18"] = 3
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_master_parser_uses_reporting_tool_and_known_dimensions():
    master = IndicatorMasterWorkbook(workbook_bytes(), "master.xlsx")
    key = ("civil-counselling", "2026-01", "UNHCR 2026 - AMAL CAMP", "AMAL Camp", "idp", "male", "00-04")
    assert master.values[key] == 3
    assert master.metadata()["sheet"] == "Reporting Tool"


def test_master_parser_allows_selecting_a_sheet():
    master = IndicatorMasterWorkbook(workbook_bytes(False), "master.xlsx", sheet_name="Other")
    assert master.metadata()["sheet"] == "Other"
    assert master.metadata()["availableSheets"] == ["Other"]


def test_reconciliation_statuses_and_export():
    master = SimpleNamespace(
        months=["2026-01"], warnings=[],
        metadata=lambda: {"ready": True, "filename": "master.xlsx", "sheet": "Reporting Tool", "months": ["2026-01"], "loadedAt": "", "remembered": False, "warnings": []},
        values={
            ("civil-counselling", "2026-01", "UNHCR 2026 - AMAL CAMP", "AMAL Camp", "idp", "male", "00-04"): 3,
            ("civil-counselling", "2026-01", "UNHCR 2026 - AMAL CAMP", "AMAL Camp", "idp", "male", "05-11"): 0,
            ("civil-counselling", "2026-01", "UNHCR 2026 - AMAL CAMP", "AMAL Camp", "idp", "male", "12-17"): None,
            ("civil-counselling", "2026-01", "UNHCR 2026 - AMAL CAMP", "AMAL Camp", "idp", "male", "18-59"): 5,
        },
    )
    values = [3, 2, 4] + [0] * 10
    report = {"filterOptions": {"months": ["2026-01"]}, "groups": [{"indicators": [{"id": "civil-counselling", "title": "Civil counselling", "children": [], "sections": [{"id": "idp", "rows": [{"project": "UNHCR 2026 - AMAL CAMP", "location": "AMAL Camp", "values": values}]}]}]}]}
    request = SimpleNamespace(fromDate="", toDate="", projects=[], projectLocations=[], years=[], quarters=[], months=["2026-01"], communityTypes=[])
    with patch("backend.indicator_reconciliation.build_indicator_report", return_value=report):
        result = reconcile(master, {}, request)
    assert result["summary"] == {"matched": 1, "different": 0, "workbookOnly": 1, "platformOnly": 1, "missingWorkbook": 1, "warnings": 0}
    assert len(result["rows"]) == 4

    exported = build_reconciliation_workbook(result)
    assert len(exported) > 1000
    exported_book = load_workbook(BytesIO(exported), read_only=True)
    assert exported_book["Comparison"]["I1"].value == "Excel value"
