"""Read-only, paged Kobo response grids. Repeat scopes remain separate tables."""
from .kobo import KoboFilter, filtered_records, label, missing, question_result, question_categories, QuestionSettings, breakdown_fields, breakdown_values


def pivot_records(project, settings, scope, records, selection):
    questions = {q['id']: q for q in project['form']['questions']}
    question = questions.get(selection['question'])
    if question is None:
        raise ValueError('Question not found.')
    if question['repeat'] != scope:
        raise ValueError('The selected question belongs to a different response table.')
    lookup = {record['id']: record for record in project['records']}
    snapshot = question_result(question, records, settings, questions, lookup)
    answer, breakdown = selection.get('answer'), selection.get('breakdown')
    if answer is not None and not any(row['id'] == answer for row in snapshot['rows']):
        raise ValueError('The selected answer is no longer available in this pivot.')
    column = next((column for column in snapshot['breakdowns'] if column['id'] == breakdown), None)
    if breakdown is not None and column is None:
        raise ValueError('The selected breakdown is no longer available in this pivot.')
    if answer == '__blank__' and breakdown is not None:
        raise ValueError('Blank answers do not have breakdown values.')
    answered = [record for record in records if not missing(record['values'].get(question['id']))]
    cfg = settings.questions.get(question['id'], QuestionSettings())
    _, _, categories = question_categories(question, answered, cfg)
    dimensions = breakdown_fields(cfg, questions, question)
    selected = []
    for record in records:
        value = record['values'].get(question['id'])
        matches = missing(value) if answer == '__blank__' else not missing(value) and (answer is None or answer in categories(value))
        if matches and (column is None or list(breakdown_values(record, dimensions, lookup)) == column['values']):
            selected.append(record)
    return selected


def data_rows(project, settings, scope, columns, search='', sort_column='', direction='asc', drilldown=None, column_filters=None):
    questions = {q['id']: q for q in project['form']['questions'] if q['repeat'] == scope}
    if scope and not questions:
        raise ValueError('Response table not found.')
    if any(column not in questions for column in columns):
        raise ValueError('A selected column is not in this response table.')
    if sort_column and sort_column != '__row__' and sort_column not in columns:
        raise ValueError('Select a visible column to sort.')
    columns = list(dict.fromkeys(columns))
    captions = {qid: {c['name']: label(c, settings.language) for c in questions[qid]['choices']} for qid in columns}

    def display(qid, value):
        if value is None: return ''
        if isinstance(value, list): return ', '.join(display(qid, item) for item in value)
        return captions[qid].get(str(value), str(value))

    records = [record for record in filtered_records(project, settings) if record['scope'] == scope]
    if drilldown is not None:
        if drilldown.get('baseFilters') is not None:
            # Keep the clicked pivot cohort and numeric bin boundaries stable.
            base_settings = settings.model_copy(update={'filters': [KoboFilter.model_validate(rule) for rule in drilldown['baseFilters']]})
            base_records = [record for record in filtered_records(project, base_settings) if record['scope'] == scope]
            selected_ids = {record['id'] for record in pivot_records(project, base_settings, scope, base_records, drilldown)}
            records = [record for record in records if record['id'] in selected_ids]
        else:
            records = pivot_records(project, settings, scope, records, drilldown)
    query = search.strip().casefold()
    rows = []
    for record in records:
        values = {qid: display(qid, record['values'].get(qid)) for qid in columns}
        if query and not any(query in text.casefold() for text in [str(record['row']), *values.values(), *(str(record['values'].get(qid, '')) for qid in columns)]):
            continue
        rows.append((record, {'id': record['id'], 'row': record['row'], 'values': values}))
    if sort_column:
        if sort_column == '__row__':
            rows.sort(key=lambda item: item[0]['row'], reverse=direction == 'desc')
        else:
            filled = [item for item in rows if item[1]['values'][sort_column] != '']
            blank = [item for item in rows if item[1]['values'][sort_column] == '']
            def sort_key(item):
                value = item[0]['values'].get(sort_column)
                if questions[sort_column]['kind'] == 'number':
                    try: return (0, float(value))
                    except (ValueError, TypeError): pass
                return (1, item[1]['values'][sort_column].casefold())
            filled.sort(key=sort_key, reverse=direction == 'desc')
            rows = filled + blank
    if column_filters:
        import pandas as pd
        from .column_filters import apply_column_filters
        frame = pd.DataFrame([{'__row__':item[1]['row'], **item[1]['values']} for item in rows], columns=['__row__', *columns])
        indexes = apply_column_filters(frame, column_filters).index
        rows = [rows[index] for index in indexes]
    return {'total': len(records), 'matched': len(rows), 'rows': [item[1] for item in rows]}


def data_page(project, settings, scope, columns, search='', sort_column='', direction='asc', page=1, page_size=50, drilldown=None, column_filters=None):
    result = data_rows(project, settings, scope, columns, search, sort_column, direction, drilldown, column_filters)
    start = (page - 1) * page_size
    return {**result, 'rows': result['rows'][start:start + page_size]}


def data_workbook(project, settings, scope, columns, search='', sort_column='', direction='asc', selected_ids=None, drilldown=None, column_filters=None):
    from collections import Counter
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    columns = list(dict.fromkeys(columns))
    result = data_rows(project, settings, scope, columns, search, sort_column, direction, drilldown, None if selected_ids is not None else column_filters)
    selected = None if selected_ids is None else set(selected_ids)
    questions = {q['id']: q for q in project['form']['questions']}
    labels = [label(questions[qid], settings.language) for qid in columns]
    counts = Counter(['Source row', *labels])
    headers, used = ['Source row'], {'Source row'}
    for qid, caption in zip(columns, labels):
        title = f"{caption} [{questions[qid]['name']}]" if counts[caption] > 1 else caption
        if title in used: title = f'{title} [{qid}]'
        while title in used: title += ' [column]'
        used.add(title); headers.append(title)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet('Response data')
    sheet.freeze_panes = 'B2'
    sheet.sheet_view.rightToLeft = settings.language == 'ar'
    from openpyxl.utils import get_column_letter
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 14 if index == 1 else 32

    def cells(values, header=False):
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            # Export response text literally, including strings beginning with '='.
            if isinstance(value, str): cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if header:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='245B83')
            yield cell

    sheet.append(list(cells(headers, True)))
    for row in result['rows']:
        if selected is None or row['id'] in selected:
            sheet.append(list(cells([row['row'], *(row['values'][qid] for qid in columns)])))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
