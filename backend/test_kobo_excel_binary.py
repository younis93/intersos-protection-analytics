from copy import deepcopy
from queue import Queue

from openpyxl import load_workbook
import pytest

from backend.kobo import analyze, KoboSettings, QuestionSettings
from backend.kobo_excel_direct import direct_helpers
from backend.kobo_excel_prepare import helper_values, prepare_workbook
from backend.kobo_excel_source import response_layout, query_config
from backend.kobo_excel_slicers import slicer_specs, pivot_value
from backend.kobo_excel import native_breakdown_columns
from backend.test_kobo import sample


def setup(project, **kwargs):
    settings = KoboSettings(hidden=['group/respondent_name'], hideZero=False, hideEmpty=False, **kwargs)
    snapshot = analyze(project, settings)
    layouts = response_layout(project, settings)
    config = query_config(project, snapshot, settings, slicer_specs(project, settings), layouts)
    config['excelMode'] = 'binary'
    return settings, snapshot, layouts, config


@pytest.mark.parametrize('hide_blanks', [True, False])
def test_binary_cached_counts_and_bases_match_analysis_without_expansion(sample, hide_blanks):
    project = deepcopy(sample[0])
    project['records'][0]['values']['group/needs'].append('water')
    settings, snapshot, layouts, config = setup(project, hideBlanks=hide_blanks, filterFields=['group/gender'])
    helpers, sources = direct_helpers(config, layouts)
    cached = helper_values(layouts, config, sources, helpers)
    values = [dict(zip(helpers[''], row)) for row in cached[0]]
    source = sources['group/needs']
    normalized = deepcopy(project)
    for record in normalized['records']:
        value = record['values'].get('group/needs')
        if value: record['values']['group/needs'] = list(dict.fromkeys(value))
    result = next(q for q in analyze(normalized, settings)['questions'] if q['id'] == 'group/needs')
    assert source['binary'] and source['direct'] and source['table'] == 'KoboResponses000'
    assert len(values) == len(project['records'])
    assert sum(row[source['fields']['Answered']] for row in values) == result['denominator'] == 2
    for option in source['options']:
        assert sum(row[option['field']] for row in values) == next(row['count'] for row in result['rows'] if row['id'] == option['id'])
        assert all(row[option['field']] in (0, 1) for row in values)
        formula = helpers[''][option['field']]
        if option['id'] != '__blank__': assert 'FIND(" '+option['id']+' "' in formula


def test_binary_formulas_are_exact_case_sensitive_and_headers_unique(sample):
    _, _, layouts, config = setup(sample[0])
    q = next(q for q in config['questions'] if q['kind'] == 'multiple')
    q['label'] = 'A'*300
    q['answers'] = [{'id': code, 'label': 'Same'*100} for code in ('a', 'aa', 'A', '*', '?', '"')]
    helpers, sources = direct_helpers(config, layouts)
    options = sources[q['id']]['options']
    assert len({o['field'].casefold() for o in options}) == len(options)
    assert len({o['caption'].casefold() for o in options}) == len(options)
    assert all(len(o['field']) <= 255 for o in options)
    assert all('FIND(' in helpers[''][o['field']] and 'SEARCH(' not in helpers[''][o['field']] for o in options)
    assert '" "" "' in helpers[''][options[-1]['field']]


def test_binary_preparation_has_formulas_cached_values_and_no_reference_tables(sample, tmp_path):
    settings, _, layouts, config = setup(sample[0], hideBlanks=False)
    path = tmp_path/'binary.xlsx'
    prepared = prepare_workbook(path, layouts, config, settings, Queue())
    assert prepared['multi'] == {}
    values = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    assert values.sheetnames == ['Data', '_Kobo choices']
    source = prepared['sources']['group/needs']
    headers = list(next(values['Data'].values))
    for option in source['options']:
        col = headers.index(option['field'])+1
        assert values['Data'].cell(2, col).value in (0, 1)
        assert formulas['Data'].cell(2, col).value.startswith('=')
        assert formulas['Data'].tables['KoboResponses000'].tableColumns[col-1].calculatedColumnFormula is not None
    values.close(); formulas.close()


def test_binary_getpivotdata_separates_option_counts_and_respondent_denominators(sample):
    _, snapshot, layouts, config = setup(sample[0])
    _, sources = direct_helpers(config, layouts)
    result = next(q for q in snapshot['questions'] if q['id'] == 'group/needs')
    result.update(_source=sources[result['id']], _countPivot="'Analysis'!$A$20")
    count = pivot_value("'Totals'!$A$1", 'Responses', 'water', result=result)
    base = pivot_value("'Totals'!$A$1", 'Answered total', 'water', result=result)
    assert '"Clean water"' in count and "'Analysis'!$A$20" in count
    assert '"Answered total"' in base and "'Totals'!$A$1" in base
    assert 'Clean water' not in base and 'AnswerCode' not in count


def test_repeat_binary_helpers_resolve_parent_breakdowns_and_trend_dates(sample):
    project = deepcopy(sample[0])
    project['form']['groups'].append({'id':'repeat','parent':'group','repeat':True,'labels':{'en':'Visits'}})
    q = deepcopy(next(q for q in project['form']['questions'] if q['kind'] == 'multiple'))
    q.update(id='repeat/needs', repeat='repeat', group='repeat')
    project['form']['questions'].append(q)
    project['records'].append({'id':'visit','parent':project['records'][0]['id'],'scope':'repeat','row':2,'values':{q['id']:['water','water']}})
    settings, snapshot, layouts, config = setup(project, questions={q['id']:QuestionSettings(breakdowns=['group/gender'])})
    question = next(item for item in config['questions'] if item['id'] == q['id'])
    question.update(trendDateField='group/day', trendUnit='month')
    helpers, sources = direct_helpers(config, layouts)
    cached = helper_values(layouts, config, sources, helpers)
    row = dict(zip(helpers['repeat'], cached[1][0])); source = sources[q['id']]
    assert row[source['fields']['Breakdown1']] == 'Female'
    assert row[source['fields']['TrendPeriod']] == '2026-01'
    assert 'INDEX(KoboResponses000[' in helpers['repeat'][source['fields']['Breakdown1']]
    assert sum(row[option['field']] for option in source['options']) == 1


def test_single_option_binary_pivot_has_a_real_answer_row(sample):
    project = deepcopy(sample[0])
    q = next(q for q in project['form']['questions'] if q['kind'] == 'multiple')
    q['choices'] = q['choices'][:1]
    for record in project['records']: record['values'][q['id']] = ['water']
    _, _, layouts, config = setup(project)
    helpers, sources = direct_helpers(config, layouts)
    source = sources[q['id']]
    assert source['single'] and source['options'][0]['metric'] == '#'
    cached = helper_values(layouts, config, sources, helpers)
    row = dict(zip(helpers[''], cached[0][0]))
    assert row[source['fields']['Answer']] == source['options'][0]['caption']
    formula = pivot_value("'Totals'!$A$1", 'Responses', 'water', result={'_source':source, '_countPivot':"'Report'!$A$1"})
    assert '"#"' in formula and '"Clean water"' in formula and source['fields']['Answer'] in formula


def test_binary_keeps_breakdown_columns_when_all_answers_are_missing(sample):
    project = deepcopy(sample[0])
    for record in project['records']: record['values']['group/needs'] = None
    settings, snapshot, layouts, config = setup(project, hideBlanks=False, questions={'group/needs':QuestionSettings(breakdowns=['group/gender'])})
    _, sources = direct_helpers(config, layouts)
    result = next(q for q in snapshot['questions'] if q['id'] == 'group/needs')
    assert result['breakdowns'] == []
    result['_source'] = sources[result['id']]
    columns = native_breakdown_columns(project, result, settings)
    assert {column['label'] for column in columns} == {'Female', 'Male'}
    assert all(column['denominator'] == 0 for column in columns)
