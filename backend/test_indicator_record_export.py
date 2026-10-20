from datetime import datetime
from io import BytesIO

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import main
from backend.indicator_reporting import build_indicator_report
from backend.legal_platform import LegalStore
from backend.indicator_record_export import SOURCE_SHEETS


@pytest.fixture
def store(monkeypatch):
    common = {"Projects": "UNHCR 2026 - Gov", "Project Location": "Anbar أنبار",
              "Community Type": "Syrian Refugee", "Gender": "Male", "UNHCR Age Group": "(18-39)",
              "Beneficiary ID": "B1", "Assessment ID": "A1", "Name": "أحمد علي", "Extra field": "full information"}
    counsel = {**common, "Date of Assessment": "15/01/2026", "Type of Legal Service Needed": "Legal Counselling",
               "Type of Service Provided": "Legal Counselling", "Is the beneficiary detained": "No"}
    representation = {**counsel, "Assessment ID": "A2", "Type of Legal Service Needed": "Legal Representation",
                      "Type of Service Provided": "Legal Representation"}
    service = {**common, "Service ID": "S1", "Date of Service Provision": "15/06/2025", "Date Service Completed": "15/01/2026",
               "Type of Service Provided": "Legal Representation", "Is the beneficiary detained": "No", "Service Status": "Completed اكتملت"}
    frames = {
        "assessments": pd.DataFrame([counsel, counsel, representation, representation,
                                     {**counsel, "Assessment ID": "A3", "Beneficiary ID": "", "Date of Assessment": "15/02/2026"}]),
        "legalservices": pd.DataFrame([service, service, {**service, "Date Service Completed": "15/02/2026"}]),
        "deportationrecords": pd.DataFrame([{**common, "Date of Deportation Knowledge": "15/01/2026", "PN ID": "PN1"},
                                            {**common, "Date of Deportation Knowledge": "15/02/2026", "PN ID": "PN2", "Beneficiary ID": ""}]),
        "awareness": pd.DataFrame([{**common, "Date of Session": "15/01/2026", "Awareness ID": "W1"}]),
        "beneficiaries": pd.DataFrame([{**common, "Case ID": "B1"}]),
    }
    source = LegalStore(frames, "Synthetic records", [])
    monkeypatch.setattr(main, "require_legal_store", lambda: source)
    return source


def export(client, store, records, revision=None):
    return client.post("/api/legal/indicators/records-export", json={"revision": revision or store.data_revision, "records": records})


def all_items(report):
    for group in report["groups"]:
        for item in group["indicators"]:
            yield item
            yield from item["children"]


def test_references_follow_counts_filters_deduplication_and_missing_ids(store):
    report = build_indicator_report(store.frames, include_source_records=True, months=["2026-01"])
    items = {item["id"]: item for item in all_items(report)}
    for item in items.values():
        for section in item["sections"]:
            for row in section["rows"]:
                assert [len(cell) for cell in row["sourceRecords"]] == row["values"]
            assert [len(cell) for cell in section["totalSourceRecords"]] == section["totals"]
    representation = items["legal-representation"]["sections"][0]["totalSourceRecords"][12]
    assert representation == [{"dataset": "assessments", "rowIndex": 2}, {"dataset": "legalservices", "rowIndex": 0}]
    counselling = items["legal-counselling"]["sections"][0]["totalSourceRecords"][12]
    assert counselling == [{"dataset": "assessments", "rowIndex": 0}, {"dataset": "assessments", "rowIndex": 1}]
    february = build_indicator_report(store.frames, include_source_records=True, months=["2026-02"])
    counselling = next(item for item in all_items(february) if item["id"] == "legal-counselling")["sections"][0]
    assert counselling["totalBeneficiaryIds"][12] == []
    assert counselling["totalSourceRecords"][12] == [{"dataset": "assessments", "rowIndex": 4}]
    empty = build_indicator_report(store.frames, include_source_records=True, locations=["AMAL Camp"])
    assert all(not cell for item in all_items(empty) for section in item["sections"] for cell in section["totalSourceRecords"])


def test_full_source_columns_and_distinct_physical_rows(store):
    client = TestClient(main.app)
    records = [{"dataset": name, "rowIndex": 0} for name in store.frames]
    records += [{"dataset": "assessments", "rowIndex": 0}, {"dataset": "assessments", "rowIndex": 1}]
    response = export(client, store, records)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == list(SOURCE_SHEETS.values())
    for name, frame in store.frames.items():
        sheet = workbook[SOURCE_SHEETS[name]]
        assert [cell.value for cell in sheet[1]] == list(frame.columns)
        assert sheet.max_row == (3 if name == "assessments" else 2)
        assert sheet.cell(2, list(frame.columns).index("Name") + 1).value == "أحمد علي"
        assert sheet.cell(2, list(frame.columns).index("Extra field") + 1).value == "full information"
        assert sheet.freeze_panes == "A2"
    date_column = list(store.frames["assessments"].columns).index("Date of Assessment") + 1
    assert workbook["Assessments"].cell(2, date_column).value == datetime(2026, 1, 15)


def test_export_uses_contributing_rows_from_mixed_indicator(store):
    client = TestClient(main.app)
    report = client.post("/api/legal/indicators", json={"months": ["2026-01"]}).json()
    assert report["revision"] == store.data_revision
    item = next(item for item in all_items(report) if item["id"] == "legal-representation")
    refs = item["sections"][0]["totalSourceRecords"][12]
    response = export(client, store, refs)
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Assessments", "Legal Services"]
    assert workbook["Assessments"].max_row == workbook["Legal Services"].max_row == 2
    assert workbook["Assessments"].cell(2, 7).value == "A2"
    # Neither same-ID services excluded by monthly deduplication nor February rows leak in.
    assert len(refs) == 2
    monthly = client.post("/api/legal/indicators/monthly", json={"months": ["2026-01", "2026-02"]}).json()
    assert all(entry["report"]["revision"] == store.data_revision for entry in monthly["reports"])


@pytest.mark.parametrize("records", [[], [{"dataset": "unknown", "rowIndex": 0}],
                                    [{"dataset": "assessments", "rowIndex": -1}],
                                    [{"dataset": "assessments", "rowIndex": 0.5}],
                                    [{"dataset": "assessments", "rowIndex": True}]])
def test_invalid_reference_validation(store, records):
    assert export(TestClient(main.app), store, records).status_code == 422


def test_stale_revision_and_unavailable_records_fail_without_partial_export(store):
    client = TestClient(main.app)
    records = [{"dataset": "assessments", "rowIndex": 0}]
    assert export(client, store, records, revision="old-import").status_code == 409
    records.append({"dataset": "assessments", "rowIndex": 999})
    assert export(client, store, records).status_code == 400
    del store.frames["awareness"]
    assert export(client, store, [{"dataset": "awareness", "rowIndex": 0}]).status_code == 400


def test_export_preserves_blanks_and_excludes_internal_fields(store):
    store.frames["assessments"]["__rowKey"] = "internal"
    store.frames["assessments"]["Empty field"] = None
    response = export(TestClient(main.app), store, [{"dataset": "assessments", "rowIndex": 4}])
    sheet = load_workbook(BytesIO(response.content))["Assessments"]
    headers = [cell.value for cell in sheet[1]]
    assert "__rowKey" not in headers
    assert sheet.cell(2, headers.index("Beneficiary ID") + 1).value is None
    assert sheet.cell(2, headers.index("Empty field") + 1).value is None


def test_reimport_rejects_references_from_previous_store(store, monkeypatch):
    original_revision = store.data_revision
    replacement = LegalStore({name: frame.copy() for name, frame in store.frames.items()}, "Reimported", [])
    monkeypatch.setattr(main, "require_legal_store", lambda: replacement)
    response = export(TestClient(main.app), replacement, [{"dataset": "assessments", "rowIndex": 0}], original_revision)
    assert response.status_code == 409
    assert "Refresh Indicator Reporting" in response.json()["detail"]


def test_export_does_not_truncate_to_generic_table_download_limit(store):
    store.frames["legalservices"] = pd.DataFrame({"Service ID": [f"S{index}" for index in range(10001)]})
    records = [{"dataset": "legalservices", "rowIndex": index} for index in range(10001)]
    response = export(TestClient(main.app), store, records)
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), read_only=True)["Legal Services"]
    assert sheet.max_row == 10002
