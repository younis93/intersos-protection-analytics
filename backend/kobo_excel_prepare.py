"""Prepare editable response sheets outside Excel's serial COM object model."""
import ctypes
import multiprocessing
import os
from zipfile import ZipFile, ZIP_STORED
from xml.etree import ElementTree as ET
from datetime import datetime
from concurrent.futures import ProcessPoolExecutor

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from .kobo import label
from .kobo_excel_direct import direct_helpers, choice_lookup_rows, structured_column


def available_memory():
    if os.name != 'nt': return 0
    class Memory(ctypes.Structure):
        _fields_ = [('length', ctypes.c_ulong), ('load', ctypes.c_ulong),
                    *[(name, ctypes.c_ulonglong) for name in ('total', 'available', 'page', 'free_page', 'virtual', 'free_virtual', 'extended')]]
    memory = Memory(); memory.length = ctypes.sizeof(memory)
    return memory.available if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(memory)) else 0


def worker_count(cells, available=None, cpus=None):
    """Leave 8 GiB free and budget 1 GiB per preparation process."""
    available = available_memory() if available is None else available
    cpus = (os.cpu_count() or 1) if cpus is None else cpus
    if cells < 250_000: return 1
    return max(1, min(4, max(1, cpus//2), max(0, (available-8*1024**3)//1024**3)))


def _worker_ready(messages):
    messages.put({'preparationPid': os.getpid()})


def _response_block(args):
    from .kobo_excel_source import response_values
    layout, records = args
    return [response_values(layout, record, escape=False) for record in records]


def file_formula(formula, layout):
    # OOXML uses fully qualified #This Row references, rather than the UI's @ shorthand.
    for header in layout['headers']:
        escaped = structured_column(header)
        formula = formula.replace(f'[@[{escaped}]]', f'{layout["table"]}[[#This Row],[{escaped}]]')
    return formula.replace('KoboChoices,', 'KoboChoices[],')


def helper_values(layouts, config, sources, helpers):
    """Seed formula caches so native caches can read them before final calculation."""
    from .kobo import missing
    from .kobo_excel_slicers import scalar_key
    records = {record['id']: record for layout in layouts for record in layout['rows']}
    choices = dict(choice_lookup_rows(config))

    def raw(record, qid):
        field = config['fields'][qid]; seen = set()
        while record['scope'] != field['scope']:
            if record['id'] in seen: raise ValueError('Circular parent response IDs.')
            seen.add(record['id']); record = records.get(record.get('parent'))
            if record is None: return None
        return record['values'].get(qid)

    def caption(record, qid, blank, namespace='field:'):
        value = raw(record, qid)
        if missing(value): return blank
        field = config['fields'][qid]
        key = scalar_key(value, field['kind'])
        return choices.get(namespace+qid+'|'+key, key)

    cached = []
    for layout in layouts:
        rows = []
        own = [q for q in config['questions'] if q['scope'] == layout['scope'] and (q['kind'] != 'multiple' or config.get('excelMode') == 'binary')]
        for record in layout['rows']:
            values = {'[Calc] Response': 1}
            for spec in config['filters']:
                values[spec['header']] = caption(record, spec['id'], 'Blank / unanswered', 'filter:')
            for question in own:
                value = raw(record, question['id']); answered = not missing(value)
                fields = sources[question['id']]['fields']
                answer = caption(record, question['id'], question['blankLabel'])
                if answered and question['bins']:
                    answer = next((bucket['label'] for i, bucket in enumerate(question['bins'])
                        if value >= bucket['min'] and (value < bucket['max'] or i == len(question['bins'])-1 and value <= bucket['max'])), None)
                    if answer is None: raise ValueError('Number outside exported bin limits.')
                elif answered and question['kind'] == 'date':
                    length = 4 if question['dateUnit'] == 'year' else 7 if question['dateUnit'] == 'month' else 10
                    key = str(value)[:length]
                    answer = choices.get('period:'+question['id']+'|'+key, key)
                if question.get('trendDateField'):
                    raw_date = raw(record, question['trendDateField'])
                    try: period = datetime.fromisoformat(str(raw_date)[:10]).isoformat()[:{'day': 10, 'month': 7, 'year': 4}[question['trendUnit']]]
                    except (ValueError, TypeError): period = ''
                    values[fields['TrendPeriod']] = period
                if sources[question['id']].get('binary'):
                    if sources[question['id']].get('single'): values[fields['Answer']] = sources[question['id']]['options'][0]['caption']
                    selected = set(value or [])
                    for option in sources[question['id']]['options']:
                        values[option['field']] = int(option['id'] in selected) if option['id'] != '__blank__' else int(not answered)
                else: values[fields['Answer']] = answer
                values[fields['Answered']] = int(answered)
                values[fields['IncludeInPivot']] = 'Included' if answered or question['blank'] else 'Excluded'
                for level, qid in enumerate(question['dimensions'], 1):
                    values[fields[f'Breakdown{level}']] = caption(record, qid, 'Missing breakdown')
            rows.append([values[header] for header in helpers[layout['scope']]])
        cached.append(rows)
    return cached


def seed_formula_caches(path, layouts, cached):
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', ns)
    temporary = path.with_suffix('.cached.xlsx')
    mapping = {f'xl/worksheets/sheet{index+1}.xml': (layout, rows) for index, (layout, rows) in enumerate(zip(layouts, cached))}
    try:
        with ZipFile(path) as source, ZipFile(temporary, 'w', ZIP_STORED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename in mapping:
                    layout, rows = mapping[item.filename]
                    root = ET.fromstring(data)
                    for row in root.findall(f'{{{ns}}}sheetData/{{{ns}}}row'):
                        index = int(row.get('r'))-2
                        if index < 0 or index >= len(rows): continue
                        formulas = [cell for cell in row if cell.find(f'{{{ns}}}f') is not None]
                        for cell, value in zip(formulas, rows[index]):
                            cell.set('t', 'str' if isinstance(value, str) else 'n')
                            node = cell.find(f'{{{ns}}}v')
                            if node is None: node = ET.SubElement(cell, f'{{{ns}}}v')
                            node.text = str(value)
                    data = ET.tostring(root, encoding='utf-8')
                target.writestr(item, data)
        temporary.replace(path)
    finally: temporary.unlink(missing_ok=True)


def materialize_multi(args):
    """Initial query results, matching the workbook-local M refresh transform."""
    config, records = args
    from .kobo import missing
    from .kobo_excel_slicers import scalar_key
    question = config['questions'][0]
    lookup = {record['id']: record for record in records}
    if len(lookup) != len(records): raise ValueError('Response IDs must be unique across all Data tables.')
    def owner(record, scope):
        visited = set()
        while record['scope'] != scope:
            if record['id'] in visited: raise ValueError('Circular parent response IDs.')
            visited.add(record['id']); record = lookup.get(record.get('parent'))
            if record is None: raise ValueError('Repeat entries must have a valid Parent ID.')
        return record
    def raw(record, qid):
        field = config['fields'][qid]
        return owner(record, field['scope'])['values'].get(qid)
    def caption(record, qid):
        value = raw(record, qid)
        return 'Missing breakdown' if missing(value) else config['fields'][qid]['captions'].get(scalar_key(value), scalar_key(value))
    def filters(record):
        return [f['captions'].get(scalar_key(raw(record, f['id']), config['fields'][f['id']]['kind']),
                    'Blank / unanswered' if missing(raw(record, f['id'])) else caption(record, f['id'])) for f in config['filters']]
    options = {answer['id']: answer['label'] for answer in question['answers']}
    if question['blank']: options['__blank__'] = question['blankLabel']
    prefix = [question['id'], question['label'], question['group'], question['groupLabel']]
    cohort = [record for record in records if record['scope'] == question['scope']]
    rows = []
    for record in cohort:
        value = raw(record, question['id'])
        selected = list(dict.fromkeys(value or []))
        answered = bool(selected)
        dimensions = [caption(record, qid) for qid in question['dimensions']]
        dimensions += ['']*(3-len(dimensions))
        details = [record['id'], record.get('parent') or '', record['row']]
        for field in config['details']:
            value = raw(record, field['id']); kind = config['fields'][field['id']]['kind']
            if isinstance(value, list): value = ' '.join(str(part) for part in value)
            if kind == 'date' and value: value = datetime.fromisoformat(str(value)[:10])
            details.append(value)
        trend = []
        if question.get('trendDateField'):
            try: period = datetime.fromisoformat(str(raw(record, question['trendDateField']))[:10]).isoformat()[:{'day': 10, 'month': 7, 'year': 4}[question['trendUnit']]]
            except (ValueError, TypeError): period = ''
            trend = [period]
        suffix = filters(record)+trend+details
        rows.append([*prefix, record['id'], '__base__', 'Respondent base', 0, int(answered), *dimensions, 'Denominator', *suffix])
        for code in selected if answered else (['__blank__'] if question['blank'] else []):
            rows.append([*prefix, record['id'], code, options.get(code, code), 1, 0, *dimensions, 'Included', *suffix])
    seed_filters = filters(cohort[0]) if cohort else [None]*len(config['filters'])
    for code, caption_value in options.items():
        rows.append([*prefix, None, code, caption_value, 0, 0, '', '', '', 'Category seed',
                     *seed_filters, *([None] if question.get('trendDateField') else []), *([None]*(3+len(config['details'])))])
    if len(rows) > 1_048_575: raise ValueError('Pivot support data exceeds the Excel row limit.')
    return rows


def prepare_workbook(path, layouts, config, settings, messages):
    from .kobo_excel_source import multi_config
    helpers, sources = direct_helpers(config, layouts)
    cells = sum(max(1, len(layout['rows']))*(len(layout['headers'])+len(helpers[layout['scope']])) for layout in layouts)
    records = [record for layout in layouts for record in layout['rows']]
    multis = []
    for index, question in enumerate(config['questions'], 1):
        if question['kind'] != 'multiple' or config.get('excelMode') == 'binary': continue
        own = multi_config(config, question)
        multis.append({'id': question['id'], 'config': own, 'seed': f'KoboSeed{index:03d}', 'sheet': f'_Kobo seed {index:03d}'})
        cells += sum(record['scope'] == question['scope'] for record in records)*3*len(own['headers'])
    workers = worker_count(cells)
    messages.put({'preparationWorkers': workers, 'message': f'Preparing editable tables using {workers} worker' + ('s' if workers != 1 else '')})
    pool = ProcessPoolExecutor(max_workers=workers, mp_context=multiprocessing.get_context('spawn'), initializer=_worker_ready, initargs=(messages,)) if workers > 1 else None
    book = Workbook(); book.remove(book.active)
    font = Font(name='Arial', size=10)
    try:
        for index, layout in enumerate(layouts):
            columns = helpers[layout['scope']]
            headers = [*layout['headers'], *columns]
            if len(headers) > 16384: raise ValueError('Response fields and helper columns exceed Excel\u2019s column limit.')
            sheet = book.create_sheet(layout['sheet']); sheet.append(headers)
            small = {key: layout[key] for key in ('fields',)}
            blocks = [(small, layout['rows'][start:start+512]) for start in range(0, len(layout['rows']), 512)]
            prepared = pool.map(_response_block, blocks) if pool else map(_response_block, blocks)
            formulas = [file_formula(formula, layout) for formula in columns.values()]
            for block in prepared:
                for row in block: sheet.append([*row, *formulas])
            # Excel requires an insertion row when opening a newly written empty
            # table. Remove that blank ListRow through Excel immediately on open.
            table = Table(displayName=layout['table'], ref=f'A1:{get_column_letter(len(headers))}{max(2, sheet.max_row)}')
            table._initialise_columns()
            for column, header in zip(table.tableColumns, headers): column.name = header
            for column, formula in zip(table.tableColumns[len(layout['headers']):], formulas):
                column.calculatedColumnFormula = TableFormula(attr_text=formula.lstrip('='))
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
            last = structured_column(layout['headers'][-1])
            book.defined_names.add(DefinedName(f'KoboRaw{index:03d}', attr_text=f'{layout["table"]}[[#All],[Response ID]:[{last}]]'))
            for row in sheet:
                for cell in row:
                    cell.font = font
                    if cell.row == 1:
                        cell.alignment = Alignment(wrap_text=True); cell.data_type = 's'
                    elif cell.column <= len(layout['headers']):
                        field = layout['fields'][cell.column-4] if cell.column >= 4 else None
                        cell.number_format = ('0.##########' if field and field['kind'] == 'number' else 'yyyy-mm-dd' if field and field['kind'] == 'date' else '@')
                        # Preserve literal source strings even when they start with '='.
                        if isinstance(cell.value, str): cell.data_type = 's'
            for col in range(1, len(headers)+1): sheet.column_dimensions[get_column_letter(col)].width = 28
            sheet.column_dimensions.group(get_column_letter(len(layout['headers'])+1), get_column_letter(len(headers)), hidden=True)
            sheet.sheet_properties.outlinePr.summaryRight = False
            sheet.row_dimensions[1].height = 42; sheet.freeze_panes = 'A2'; sheet.sheet_view.showGridLines = False
        lookup = book.create_sheet('_Kobo choices')
        for row in [('Key', 'Label'), *choice_lookup_rows(config)]:
            lookup.append(list(row))
            for cell in lookup[lookup.max_row]: cell.data_type = 's'
        lookup.add_table(Table(displayName='KoboChoices', ref=f'A1:B{lookup.max_row}'))
        legend = [('Question / field', 'Code', 'Label')]
        for layout in layouts:
            for question in layout['fields']:
                legend.extend((label(question, settings.language), choice['name'], label(choice, settings.language)) for choice in question['choices'])
        for row_index, row in enumerate(legend, 1):
            for column, value in enumerate(row, 5):
                cell = lookup.cell(row_index, column, value); cell.data_type = 's'
        lookup.sheet_state = 'hidden'
        materialized = pool.map(materialize_multi, [(item['config'], records) for item in multis]) if pool else map(materialize_multi, [(item['config'], records) for item in multis])
        for item, rows in zip(multis, materialized):
            # Keep staging data out of Excel until its own query is loaded.
            # Excel.CurrentWorkbook otherwise repeatedly enumerates every seed.
            item['rows'] = rows
        book.calculation.calcMode = 'manual'
        book.calculation.fullCalcOnLoad = False
        book.calculation.forceFullCalc = False
        book.calculation.calcOnSave = False
        # This is an ephemeral input file. Avoid compressing the large repeated
        # response fields twice; Excel compresses the final deliverable on save.
        with ZipFile(path, 'w', ZIP_STORED, allowZip64=True) as archive:
            ExcelWriter(book, archive).save()
        seed_formula_caches(path, layouts, helper_values(layouts, config, sources, helpers))
    finally:
        book.close()
        if pool: pool.shutdown(wait=True, cancel_futures=True)
    return {'sources': sources, 'multi': {item['id']: item for item in multis}}
