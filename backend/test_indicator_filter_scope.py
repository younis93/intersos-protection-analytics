from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.filter_selection import EXCLUDE_PREFIX
from backend.indicator_reporting import build_indicator_report, build_indicator_workbook, build_monthly_reports
from scripts.benchmark_indicator_performance import synthetic_frames


@pytest.fixture(scope="module")
def source():
    frames = synthetic_frames(120)
    return frames, build_indicator_report(frames)["filterOptions"]


def populations(report):
    return {group["id"] for group in report["groups"] if group["id"] in {"refugee", "idp"}}


def test_select_all_projects_keeps_both_populations_and_export_sheets(source):
    frames, options = source
    report = build_indicator_report(frames, projects=options["projects"])
    assert populations(report) == {"refugee", "idp"}
    assert report["groups"] == build_indicator_report(frames)["groups"]
    workbook = load_workbook(BytesIO(build_indicator_workbook(report)), read_only=True)
    assert any("refugee" in name.lower() for name in workbook.sheetnames)
    assert any("idp" in name.lower() for name in workbook.sheetnames)


def test_refugee_projects_and_amal_only_are_independent(source):
    frames, options = source
    refugee_projects = [project for project in options["projects"] if project != "UNHCR 2026 - AMAL CAMP"]
    assert populations(build_indicator_report(frames, projects=refugee_projects)) == {"refugee"}
    assert populations(build_indicator_report(frames, projects=["UNHCR 2026 - AMAL CAMP"])) == {"idp"}
    mixed = [refugee_projects[0], "UNHCR 2026 - AMAL CAMP"]
    assert populations(build_indicator_report(frames, projects=mixed)) == {"refugee", "idp"}
    assert populations(build_indicator_report(frames, projects=list(reversed(mixed)))) == {"refugee", "idp"}


def test_select_all_communities_keeps_both_populations(source):
    frames, options = source
    assert populations(build_indicator_report(frames, projects=options["projects"], community_types=options["communityTypes"])) == {"refugee", "idp"}
    assert populations(build_indicator_report(frames, projects=options["projects"], community_types=["Syrian Refugee", "Non-Syrian Refugee"])) == {"refugee"}
    assert populations(build_indicator_report(frames, projects=options["projects"], community_types=["IDP"])) == {"idp"}


def test_project_and_population_exclusions_do_not_hide_remaining_groups(source):
    frames, options = source
    assert populations(build_indicator_report(frames, projects=[EXCLUDE_PREFIX + "UNHCR 2026 - AMAL CAMP"])) == {"refugee"}
    assert populations(build_indicator_report(frames, projects=[EXCLUDE_PREFIX + options["projects"][-1]])) == {"refugee", "idp"}
    assert populations(build_indicator_report(frames, community_types=[EXCLUDE_PREFIX + "IDP"])) == {"refugee"}
    assert populations(build_indicator_report(frames, community_types=[EXCLUDE_PREFIX + "Syrian Refugee", EXCLUDE_PREFIX + "Non-Syrian Refugee"])) == {"idp"}
    assert populations(build_indicator_report(frames, community_types=[EXCLUDE_PREFIX + value for value in options["communityTypes"]])) == set()


def test_amal_location_matches_amal_project_in_reporting_analysis_and_export(source):
    frames, _ = source
    by_location = build_indicator_report(frames, locations=["AMAL Camp"])
    by_project = build_indicator_report(frames, projects=["UNHCR 2026 - AMAL CAMP"])
    assert populations(by_location) == {"idp"}
    assert by_location["groups"] == by_project["groups"]
    workbook = load_workbook(BytesIO(build_indicator_workbook(by_location)), read_only=True)
    assert not any("refugee" in name.lower() for name in workbook.sheetnames)
    monthly = build_monthly_reports(frames, locations=["AMAL Camp"])
    assert monthly["reports"]
    assert all(populations(entry["report"]) == {"idp"} for entry in monthly["reports"])


def test_location_scope_supports_refugee_mixed_excluded_and_conflicting_projects(source):
    frames, options = source
    refugee_location = next(location for project, locations in options["locationsByProject"].items()
                            if project != "UNHCR 2026 - AMAL CAMP" for location in locations)
    assert populations(build_indicator_report(frames, locations=[refugee_location])) == {"refugee"}
    assert populations(build_indicator_report(frames, locations=["AMAL Camp", refugee_location])) == {"refugee", "idp"}
    assert populations(build_indicator_report(frames, locations=[EXCLUDE_PREFIX + "AMAL Camp"])) == {"refugee"}
    refugee_projects = [project for project in options["projects"] if project != "UNHCR 2026 - AMAL CAMP"]
    assert populations(build_indicator_report(frames, projects=refugee_projects, locations=["AMAL Camp"])) == set()


@pytest.mark.parametrize("field,option", [("projects", "projects"), ("locations", "locations"),
                                       ("community_types", "communityTypes"), ("months", "months"),
                                       ("quarters", "quarters"), ("years", "years")])
def test_each_indicator_filter_keeps_mixed_populations_on_select_all(source, field, option):
    frames, options = source
    filters = {"projects": options["projects"], field: options[option]}
    report = build_indicator_report(frames, **filters)
    assert populations(report) == {"refugee", "idp"}
    monthly = build_monthly_reports(frames, **filters)
    assert monthly["reports"]
    assert all(populations(entry["report"]) == {"refugee", "idp"} for entry in monthly["reports"])
