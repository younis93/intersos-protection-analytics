from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .file_security import safe_spreadsheet_value, validate_xlsx_archive
from .indicator_reporting import AGE_GROUPS, build_indicator_report


MONTH_START_COLUMNS = tuple(4 + (40 * index) for index in range(12))
POPULATION_OFFSETS = {"syrian-refugee": 0, "non-syrian-refugee": 13, "idp": 26}
SEX_AGE = tuple((sex, age) for sex in ("male", "female") for age in AGE_GROUPS)

# Workbook row ranges are intentionally explicit. This makes template changes
# visible and testable instead of silently comparing the wrong cells.
INDICATOR_BLOCKS: dict[str, tuple[int, int, str, bool]] = {
    "civil-counselling": (18, 18, "idp", False),
    "secured-civil-documentation": (22, 22, "idp", False),
    "uid-secured": (23, 23, "idp", False),
    "civil-representation": (24, 24, "idp", False),
    "legal-awareness-participants": (25, 25, "idp", False),
    "detention-immigration": (29, 51, "refugee", False),
    "deported": (57, 79, "refugee", False),
    "released-immigration": (85, 107, "refugee", False),
    "06-1-1-legal-assistance": (113, 135, "refugee", False),
    "legal-counselling": (141, 163, "refugee", False),
    "legal-counselling-detainee": (169, 191, "refugee", False),
    "legal-counselling-other": (197, 219, "refugee", False),
    "legal-representation": (225, 247, "refugee", False),
    "legal-representation-detainee": (253, 275, "refugee", False),
    "legal-representation-other": (281, 303, "refugee", False),
    "individuals-reached": (309, 314, "all", True),
}

EXPECTED_WORKBOOK_TITLES = {
    18: "# of persons receiving legal counselling for civil documentation",
    22: "# of secured civil documentation",
    23: "# of persons who received UIDs",
    24: "# of persons receiving legal representation for civil documentation",
    25: "# of girls, boys, women, and men participating in legal awareness sessions",
    29: "# of persons identified in detention (with immigration related charges)",
    57: "# of persons deported from detention (with immigration related charges)",
    85: "# of persons successfully released from detention based on (with immigration related charges)",
    113: "06.1.1 Number of people who received legal assistance",
    141: "# of persons provided with legal counselling (Detainee + Other)",
    169: "# of persons provided with legal counselling (Detainee)",
    197: "# of persons provided with legal counselling (Other)",
    225: "# of persons provided with legal representation (Detainee + Other)",
    253: "# of persons provided with legal representation (Detainee)",
    281: "# of persons provided with legal representation (Other)",
    309: "# of individual beneficiaries reached",
}

PROJECT_ALIASES = {
    "amal camp": "UNHCR 2026 - AMAL CAMP",
    "erbil": "UNHCR 2026 - Erbil",
    "al sulaymaniyah": "UNHCR 2026 - SULI",
    "sulaymaniyah": "UNHCR 2026 - SULI",
    "mosul kirkuk": "UNHCR 2026 - Mosul & Kirkuk",
    "baghdad": "UNHCR 2026 - Baghdad",
    "governorate": "UNHCR 2026 - Gov",
}

LOCATION_ALIASES = {
    "kawergosk camp": "Kawrgawsk Camp",
    "rapareen pshdar rania": "Pshdar Urban (Refugees) + Rania",
    "arbat refugee camp": "Arbat Camp (Refugees)",
    "ninewa": "Ninewa نينوى",
    "kirkuk": "Kirkuk كركوك",
    "baghdad": "Baghdad بغداد",
    "al anbar": "Anbar أنبار",
    "al muthanna": "Al-Muthanna المثنى",
    "al qadissiya": "Al-Qadisiyyah القادسية",
    "babil": "Babil بابل",
    "al basrah": "Basra بصرة",
    "thi qar": "Dhi Qar ذي قار",
    "diyala": "Diyala ديالى",
    "kerbala": "Karbala كربلاء",
    "maysan": "Maysan ميسان",
    "al najaf": "Najaf نجف",
    "salah al din": "Salah Al-Din صلاح الدين",
    "wassit": "Wassit واسط",
    "amal camp": "AMAL Camp",
    "internal": "AMAL Camp",
    "urban": "Urban",
    "basirma camp": "Basirma Camp",
    "darashakran camp": "Darashakran Camp",
    "qushtapa camp": "Qushtapa Camp",
    "sulaymaniyah urban": "Sulaymaniyah Urban",
}


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _month(value: Any, fallback_index: int) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    parsed = datetime.fromisoformat(text) if re.match(r"^\d{4}-\d{2}-\d{2}", text) else None
    return parsed.strftime("%Y-%m") if parsed else f"2026-{fallback_index + 1:02d}"


def _number(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else round(number)


def _canonical_pair(project: Any, location: Any, aggregate: bool = False) -> tuple[str, str] | None:
    canonical_project = PROJECT_ALIASES.get(_key(project))
    if not canonical_project:
        return None
    if aggregate:
        return canonical_project, "All project locations"
    canonical_location = LOCATION_ALIASES.get(_key(location))
    return (canonical_project, canonical_location) if canonical_location else None


class IndicatorMasterWorkbook:
    def __init__(self, raw: bytes, filename: str, source_path: str = "", sheet_name: str = "") -> None:
        validate_xlsx_archive(raw)
        self.raw = raw
        self.filename = Path(filename).name
        self.source_path = source_path
        workbook = load_workbook(BytesIO(raw), read_only=True)
        self.available_sheets = list(workbook.sheetnames)
        workbook.close()
        if not self.available_sheets:
            raise ValueError("The workbook does not contain any worksheets.")
        self.sheet_name = sheet_name or ("Reporting Tool" if "Reporting Tool" in self.available_sheets else self.available_sheets[0])
        if self.sheet_name not in self.available_sheets:
            raise ValueError(f'The worksheet "{self.sheet_name}" is not available in this workbook.')
        self.loaded_at = datetime.now().isoformat(timespec="seconds")
        self.values: dict[tuple[str, str, str, str, str, str, str], int | None] = {}
        self.warnings: list[str] = []
        self.months: list[str] = []
        self._parse()

    @classmethod
    def from_path(cls, path: Path, sheet_name: str = "") -> "IndicatorMasterWorkbook":
        resolved = path.resolve()
        if not resolved.is_file() or resolved.suffix.lower() != ".xlsx":
            raise ValueError("The selected master workbook is unavailable or is not an .xlsx file.")
        return cls(resolved.read_bytes(), resolved.name, str(resolved), sheet_name)

    def metadata(self) -> dict[str, Any]:
        return {
            "ready": True,
            "filename": self.filename,
            "sheet": self.sheet_name,
            "availableSheets": self.available_sheets,
            "months": self.months,
            "loadedAt": self.loaded_at,
            "remembered": bool(self.source_path),
            "warnings": self.warnings,
        }

    def _parse(self) -> None:
        cached = load_workbook(BytesIO(self.raw), data_only=True, read_only=False)
        sheet = cached[self.sheet_name]
        self.months = [_month(sheet.cell(1, col).value, index) for index, col in enumerate(MONTH_START_COLUMNS)]
        for row, expected in EXPECTED_WORKBOOK_TITLES.items():
            actual = str(sheet.cell(row, 1).value or "").strip()
            if _key(actual) != _key(expected):
                self.warnings.append(f'Expected indicator row {row} to contain "{expected}"; found "{actual or "blank"}".')
        unmapped_locations: set[str] = set()
        for indicator_id, (start, end, population_group, aggregate) in INDICATOR_BLOCKS.items():
            populations = tuple(POPULATION_OFFSETS) if population_group == "all" else (("syrian-refugee", "non-syrian-refugee") if population_group == "refugee" else ("idp",))
            for row in range(start, end + 1):
                pair = _canonical_pair(sheet.cell(row, 2).value, sheet.cell(row, 3).value, aggregate)
                if not pair:
                    unmapped_locations.add(f"{sheet.cell(row, 2).value or 'blank'} / {sheet.cell(row, 3).value or 'blank'}")
                    continue
                project, location = pair
                for month_index, month_col in enumerate(MONTH_START_COLUMNS):
                    month = self.months[month_index]
                    for population in populations:
                        base_col = month_col + POPULATION_OFFSETS[population]
                        for offset, (sex, age) in enumerate(SEX_AGE):
                            value_cell = sheet.cell(row, base_col + offset)
                            value = _number(value_cell.value)
                            self.values[(indicator_id, month, project, location, population, sex, age)] = value
        for label in sorted(unmapped_locations):
            self.warnings.append(f"Workbook project/location is not mapped: {label}.")


def _walk_items(report: dict[str, Any]):
    for group in report.get("groups", []):
        for item in group.get("indicators", []):
            yield item
            yield from item.get("children", [])


def _platform_values(report: dict[str, Any], month: str) -> tuple[dict[tuple[str, str, str, str, str, str, str], int], dict[str, str]]:
    values: dict[tuple[str, str, str, str, str, str, str], int] = {}
    titles: dict[str, str] = {}
    for item in _walk_items(report):
        indicator_id = item["id"]
        if indicator_id not in INDICATOR_BLOCKS:
            continue
        titles[indicator_id] = item["title"]
        aggregate = INDICATOR_BLOCKS[indicator_id][3]
        for section in item.get("sections", []):
            population = section["id"]
            if aggregate:
                project_totals: dict[str, list[int]] = {}
                for row in section.get("rows", []):
                    target = project_totals.setdefault(row["project"], [0] * 12)
                    target[:] = [left + int(right or 0) for left, right in zip(target, row["values"][:12])]
                source_rows = ((project, "All project locations", row_values) for project, row_values in project_totals.items())
            else:
                source_rows = ((row["project"], row["location"], row["values"][:12]) for row in section.get("rows", []))
            for project, location, row_values in source_rows:
                for index, (sex, age) in enumerate(SEX_AGE):
                    values[(indicator_id, month, project, location, population, sex, age)] = int(row_values[index] or 0)
    return values, titles


def reconcile(master: IndicatorMasterWorkbook, frames: dict[str, Any], request: Any) -> dict[str, Any]:
    base = build_indicator_report(frames, request.fromDate, request.toDate, request.projects, request.projectLocations, request.years, request.quarters, request.months, request.communityTypes)
    months = list(request.months) or list(base.get("filterOptions", {}).get("months", []))
    if request.years:
        months = [month for month in months if month[:4] in request.years]
    if request.quarters:
        allowed = {match.group(1) for value in request.quarters if (match := re.search(r"q\s*([1-4])", value, re.I))}
        if allowed:
            months = [month for month in months if str((int(month[5:7]) - 1) // 3 + 1) in allowed]
    months = sorted(set(months) & set(master.months))
    if not months:
        raise ValueError("No common reporting month is available between the current filters and the master workbook.")

    platform: dict[tuple[str, str, str, str, str, str, str], int] = {}
    titles: dict[str, str] = {}
    for month in months:
        report = build_indicator_report(frames, request.fromDate, request.toDate, request.projects, request.projectLocations, request.years, [], [month], request.communityTypes)
        month_values, month_titles = _platform_values(report, month)
        platform.update(month_values)
        titles.update(month_titles)

    # Zero means no reported value in either source. Excluding it from each
    # source independently removes empty dimensional combinations without
    # hiding a nonzero value that exists on the other side.
    platform = {key: value for key, value in platform.items() if value != 0}
    master_values = {key: value for key, value in master.values.items() if key[1] in months and value != 0}
    if request.projects:
        allowed_projects = set(request.projects)
        master_values = {key: value for key, value in master_values.items() if key[2] in allowed_projects}
    if request.projectLocations:
        allowed_locations = set(request.projectLocations) | {"All project locations"}
        master_values = {key: value for key, value in master_values.items() if key[3] in allowed_locations}
    if request.communityTypes:
        allowed_populations = set()
        for label in request.communityTypes:
            normalized = _key(label)
            if normalized == "idp": allowed_populations.add("idp")
            elif "non syrian" in normalized: allowed_populations.add("non-syrian-refugee")
            elif "syrian" in normalized: allowed_populations.add("syrian-refugee")
        if allowed_populations:
            master_values = {key: value for key, value in master_values.items() if key[4] in allowed_populations}

    rows = []
    summary = {"matched": 0, "different": 0, "workbookOnly": 0, "platformOnly": 0, "missingWorkbook": 0}
    for key in sorted(set(platform) | set(master_values)):
        indicator_id, month, project, location, population, sex, age = key
        platform_present, workbook_present = key in platform, key in master_values
        platform_value, workbook_value = platform.get(key), master_values.get(key)
        if workbook_present and workbook_value is None:
            status = "missing-workbook"; summary["missingWorkbook"] += 1
        elif platform_present and workbook_present and platform_value == workbook_value:
            status = "matched"; summary["matched"] += 1
        elif platform_present and workbook_present:
            status = "different"; summary["different"] += 1
        elif workbook_present:
            status = "workbook-only"; summary["workbookOnly"] += 1
        else:
            status = "platform-only"; summary["platformOnly"] += 1
        rows.append({
            "indicatorId": indicator_id,
            "indicator": titles.get(indicator_id, indicator_id),
            "month": month,
            "project": project,
            "location": location,
            "population": population,
            "sex": sex,
            "ageGroup": age,
            "platformValue": platform_value,
            "workbookValue": workbook_value,
            "variance": (platform_value - workbook_value) if platform_value is not None and workbook_value is not None else None,
            "status": status,
        })
    summary["warnings"] = len(master.warnings)
    return {"metadata": master.metadata(), "months": months, "summary": summary, "rows": rows, "warnings": master.warnings}


def build_reconciliation_workbook(result: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Indicator Reporting Check"])
    summary_sheet.append(["Master workbook", result["metadata"]["filename"]])
    summary_sheet.append(["Worksheet", result["metadata"]["sheet"]])
    summary_sheet.append(["Months", ", ".join(result["months"])])
    summary_sheet.append([])
    summary_sheet.append(["Status", "Count"])
    labels = (("Matched", "matched"), ("Different", "different"), ("Workbook only", "workbookOnly"), ("Platform only", "platformOnly"), ("Missing workbook value", "missingWorkbook"), ("Warnings", "warnings"))
    for label, key in labels:
        summary_sheet.append([label, result["summary"].get(key, 0)])

    detail = workbook.create_sheet("Comparison")
    columns = ["Indicator", "Month", "Project", "Location", "Population", "Sex", "Age group", "Platform value", "Excel value", "Variance", "Status"]
    detail.append(columns)
    for row in result["rows"]:
        detail.append([safe_spreadsheet_value(value) for value in (row["indicator"], row["month"], row["project"], row["location"], row["population"], row["sex"], row["ageGroup"], row["platformValue"], row["workbookValue"], row["variance"], row["status"])])

    warning_sheet = workbook.create_sheet("Warnings")
    warning_sheet.append(["Warning"])
    for warning in result["warnings"]:
        warning_sheet.append([safe_spreadsheet_value(warning)])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="0072BC")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
