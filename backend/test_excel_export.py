from copy import copy
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

from backend.excel_export import append_frame
from backend.legal_platform import EXCEL_DATE_FORMAT, LegalStore, format_excel_dates


def test_date_parsing_is_shared_across_sheets_and_invalid_dates():
    workbook = Workbook()
    first = workbook.active
    second = workbook.create_sheet("Other")
    for sheet in (first, second):
        sheet.append(["Dates", "Identifier"])
        for value in ("01/10/2026", "2026-10-01", "31/02/2026", "invalid", date(2026, 10, 1)):
            sheet.append([value, "0001"])
        sheet.merge_cells("C1:D1")
    with patch("backend.legal_platform.pd.to_datetime", wraps=pd.to_datetime) as parse:
        format_excel_dates(workbook)
    assert parse.call_count == 3
    for sheet in workbook:
        assert sheet["A2"].value == sheet["A3"].value == sheet["A6"].value == date(2026, 10, 1)
        assert sheet["A4"].value == "31/02/2026"
        assert sheet["A5"].value == "invalid"
        assert sheet["B2"].value == "0001"
        assert sheet["A2"].number_format == EXCEL_DATE_FORMAT
        assert sheet.column_dimensions["A"].width == 22
    # A subsequent workbook must parse its own values, without retaining case data.
    another = Workbook()
    another.active.append(["01/10/2026"])
    with patch("backend.legal_platform.pd.to_datetime", wraps=pd.to_datetime) as parse:
        format_excel_dates(another)
    assert parse.call_count == 1


def test_direct_rows_match_pandas_values_and_header_styles():
    frame = pd.DataFrame({"ID": ["001", "002", "003"],
                          "Value": [float("inf"), float("-inf"), float("nan")],
                          "Text": ["'=SUM(A1:A2)", "Arabic عربي", ""],
                          "Other": [time(12, 30), timedelta(hours=12), Decimal("2.50")],
                          "Date": [datetime(2026, 10, 1)] * 3})
    previous = BytesIO()
    with pd.ExcelWriter(previous, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
        format_excel_dates(writer.book)
    workbook = Workbook()
    append_frame(workbook.active, frame)
    format_excel_dates(workbook)
    current = BytesIO()
    workbook.save(current)
    left = load_workbook(BytesIO(previous.getvalue())).active
    right = load_workbook(BytesIO(current.getvalue())).active
    for old_row, new_row in zip(left, right):
        for old, new in zip(old_row, new_row):
            assert old.value == new.value
            assert old.number_format == new.number_format
            assert copy(old.font) == copy(new.font)
            assert copy(old.border) == copy(new.border)
            assert copy(old.alignment) == copy(new.alignment)


def test_empty_flat_export_keeps_headers():
    workbook = Workbook()
    append_frame(workbook.active, pd.DataFrame(columns=["ID", "Date"]))
    assert workbook.active.max_row == 1
    assert [cell.value for cell in workbook.active[1]] == ["ID", "Date"]


def test_flat_export_preserves_a_trailing_blank_source_row():
    workbook = Workbook()
    append_frame(workbook.active, pd.DataFrame({"ID": ["001", None], "Date": [None, None]}))
    output = BytesIO()
    workbook.save(output)
    sheet = load_workbook(BytesIO(output.getvalue())).active
    assert sheet.max_row == 3
    assert sheet["A2"].value == "001"
    assert sheet["A3"].value is None


def test_hotline_parses_distinct_dates_once_across_columns():
    frame = pd.DataFrame({"Hotline ID": ["001", "002"],
                          "Contact Date": ["10/01/2026", "2026-10-02"],
                          "Follow-up Date": ["2026-10-02", "10/01/2026"]})
    store = LegalStore({"legalhotlines": frame}, "synthetic", [])
    with patch("backend.legal_platform.pd.to_datetime", wraps=pd.to_datetime) as parse:
        output = store.explorer_export("legalhotlines")
    assert parse.call_count == 1
    assert len(parse.call_args.args[0]) == 2
    sheet = load_workbook(BytesIO(output)).active
    assert sheet["B2"].value == sheet["C3"].value == datetime(2026, 10, 1)
    assert sheet["C2"].value == sheet["B3"].value == datetime(2026, 10, 2)
    assert frame["Contact Date"].tolist() == ["10/01/2026", "2026-10-02"]
