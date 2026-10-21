import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from backend import legal_platform
from backend.duplicate_exclusions import DuplicateExclusionRegistry
from backend.legal_platform import LegalStore, REVIEW_FINDING_ARABIC


RULE = "Beneficiary below 18 with legal counselling"


@pytest.fixture
def store(monkeypatch, request):
    class Today(date):
        @classmethod
        def today(cls):
            return cls(2026, 10, 21)

    monkeypatch.setattr(legal_platform, "date", Today)
    dob_header = getattr(request, "param", "Date of Birth")
    births = ["21/10/2026", "22/10/2008", "21/10/2008", "20/10/2000", "", "bad", "22/10/2026", "01/01/2012"]
    beneficiaries = pd.DataFrame({
        "Case ID": [f"B{i}" for i in range(8)],
        "Name (Filter Color Red)": [f"Beneficiary {i}" for i in range(8)],
        dob_header: births,
        "Age": [99, 99, 1, 1, 12, 12, 12, 99],
    })
    services = pd.DataFrame({
        "Service ID": [f"S{i}" for i in range(10)],
        "Beneficiary ID": [f"B{i}" for i in range(8)] + ["B1", "B1"],
        "Assessment ID": ["A1"] * 10,
        "Type of Service Provided": ["Legal Counselling"] * 7 + ["Legal Representation", "تمثيل; استشارة قانونية", "Legal Representation, Legal Counseling"],
        "Date of Service Provision": ["01/01/2020"] * 8 + ["01/02/2026", "01/03/2026"],
        "Service Status": ["Closed"] * 8 + ["In-Process", "Completed"],
        "Project": ["P1"] * 8 + ["P2", "P2"],
        "Lawyer": ["Lawyer A"] * 10,
        "Project Location": ["Erbil"] * 10,
    })
    payload = {
        "beneficiaries": beneficiaries,
        "assessments": pd.DataFrame({"Assessment ID": ["A1"], "Beneficiary ID": ["B1"]}),
        "legalservices": services,
    }
    return LegalStore.from_files({key: frame.to_csv(index=False).encode("utf-8") for key, frame in payload.items()}, "test")


@pytest.mark.parametrize("store", ["DoB", "Date of Birth", "Date of Birth / تاريخ الولادة"], indirect=True)
def test_current_age_and_linked_birth_dates(store):
    result = store.review("legalservices", rule=RULE)
    rows = {row["serviceId"]: row for row in result["rows"]}
    assert set(rows) == {"S0", "S1", "S8", "S9"}
    assert rows["S0"]["beneficiaryAge"] == 0
    assert rows["S1"]["beneficiaryAge"] == 17
    assert rows["S1"]["name"] == "Beneficiary 1"
    assert rows["S1"]["dateOfBirth"] == "22/10/2008"
    assert rows["S1"]["caseId"] == "B1"
    assert rows["S1"]["assessmentId"] == "A1"
    assert rows["S1"]["action"].startswith("Verify the beneficiary date of birth")


def test_filters_pagination_and_service_exclusion(store, tmp_path):
    result = store.review("legalservices", rule=RULE, page_size=1)
    assert result["total"] == 4
    assert result["ruleCounts"][RULE] == 4
    assert len(result["rows"]) == 1
    assert store.review("legalservices", rule=RULE, project="P2")["total"] == 2
    assert store.review("legalservices", rule=RULE, search="S8")["total"] == 1
    assert store.review("legalservices", rule=RULE, date="2026-02")["total"] == 1
    assert store.review("legalservices", rule=RULE, filters={"project": ["P1"]})["total"] == 2
    registry = DuplicateExclusionRegistry(tmp_path / "exclusions.json")
    registry.exclude_record("legalservices", RULE, "serviceId", "S8")
    store.set_review_exclusions(registry.exclusion_rows())
    assert {row["serviceId"] for row in store.review("legalservices", rule=RULE)["rows"]} == {"S0", "S1", "S9"}
    book = load_workbook(io.BytesIO(store.review_export("legalservices", selected_rules=[RULE], project="P2")))
    assert book.active.max_row == 2


def test_export_includes_age_evidence_and_obeys_filters(store):
    workbook = load_workbook(io.BytesIO(store.review_export("legalservices", selected_rules=[RULE], project="P2")))
    sheet = workbook.active
    records = list(sheet.values)
    headers = list(records[0])
    rows = [dict(zip(headers, values)) for values in records[1:]]
    assert len(rows) == 2
    assert {row["Service"] for row in rows} == {"S8", "S9"}
    assert all(str(row["Current age"]) == "17" for row in rows)
    assert all(row["نتيجة المراجعة"] == REVIEW_FINDING_ARABIC[RULE] for row in rows)
    assert all(row["Type of Service Provided"] for row in rows)
    assert all(row["Date of birth"] for row in rows)
    assert all(row["Lawyer"] == "Lawyer A" and row["Project location"] == "Erbil" for row in rows)
