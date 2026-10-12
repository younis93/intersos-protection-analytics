import io

import pandas as pd
import pytest
import asyncio
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

from backend.hotline import dashboard
from backend.legal_platform import LegalStore, versioned_dataset_name
from backend.test_legal_platform import required_payload


@pytest.fixture
def payload():
    result = required_payload()
    result['legalhotlines'] = pd.DataFrame({
        'Hotline ID': ['1', '2', '3', '4', '5'],
        'Contact Date تاريخ الاتصال': ['2/1/2026', '4/2/2026', '', 'invalid', '12/31/2025'],
        'Priority الاولوية': ['Low', 'Low منخفضة', '', 'High عالية', 'Medium متوسطة'],
        'Is the beneficiary detained هل المستفيد موقوف': ['Yes نعم', 'Yes', 'No لا', '', 'No'],
        'Has the beneficiary referred to the helpline? ': ['Yes نعم', '', 'No', '', 'Yes'],
        'Governorate of detention': ['Baghdad', 'Baghdad ', '', '', 'Erbil'],
        'Type of Request': ['Legal assistance', 'Information', '', '', 'Registration'],
    }).to_csv(index=False).encode('utf-8-sig')
    return result


def test_hotline_normalization_dates_and_source_preservation(payload):
    store = LegalStore.from_files(payload, 'test')
    result = dashboard(store.frames['legalhotlines'])
    assert result['total'] == 5
    assert result['missingContactDates'] == 2
    assert [r['count'] for r in result['trend']] == [1, 0, 1, 0, 1]
    assert [(r['detained'], r['notDetained']) for r in result['trend']] == [(0, 1), (0, 0), (1, 0), (0, 0), (1, 0)]
    assert result['filterOptions']['Priority'] == ['High', 'Low', 'Medium', 'Not recorded']
    assert result['filterOptions']['Hotline ID'] == ['1', '2', '3', '4', '5']
    assert result['map']['items'] == [
        {'label': 'Baghdad', 'count': 2, 'detained': 2, 'notDetained': 0, 'values': ['Baghdad']},
        {'label': 'Erbil', 'count': 1, 'detained': 0, 'notDetained': 1, 'values': ['Erbil']},
    ]
    rows = store.explorer('legalhotlines')['rows']
    assert rows[0]['Contact Date تاريخ الاتصال'] == '2/1/2026'
    assert rows[0]['Type of Request'] == 'Legal assistance'
    assert rows[1]['Priority الاولوية'] == 'Low منخفضة'
    exported_bytes = store.explorer_export('legalhotlines')
    exported = pd.read_excel(io.BytesIO(exported_bytes))
    assert exported.iloc[0]['Contact Date تاريخ الاتصال'] == pd.Timestamp('2026-02-01')
    assert exported.iloc[1]['Priority الاولوية'] == 'Low منخفضة'
    workbook = load_workbook(io.BytesIO(exported_bytes))
    date_column = next(cell.column for cell in workbook['Filtered data'][1] if str(cell.value).startswith('Contact Date'))
    assert workbook['Filtered data'].cell(2, date_column).number_format == '[$-en-US]yyyy-mmmm-dd'


@pytest.mark.parametrize('filters,expected', [
    ({'Priority': ['Low'], 'Governorate of detention': ['Baghdad']}, 2),
    ({'Contact Date': ['2026-02'], 'Priority': ['Low', 'High']}, 1),
    ({'Priority': ['Not recorded']}, 1),
    ({'Contact Date': ['2026-03']}, 0),
    ({'Has the beneficiary referred to the helpline?': ['Yes']}, 2),
])
def test_hotline_filters_match_table_and_exports(payload, filters, expected):
    store = LegalStore.from_files(payload, 'test')
    assert dashboard(store.frames['legalhotlines'], filters)['total'] == expected
    assert store.explorer('legalhotlines', filters=filters)['total'] == expected
    exported = pd.read_excel(io.BytesIO(store.explorer_export('legalhotlines', filters=filters)))
    assert len(exported) == expected


def test_hotline_search_pagination_and_date_sort(payload):
    store = LegalStore.from_files(payload, 'test')
    assert store.explorer('legalhotlines', search='Baghdad', filters={'Contact Date': ['2026-02']})['total'] == 1
    assert store.explorer('legalhotlines', search='Baghdad')['total'] == 2
    rows = store.explorer('legalhotlines', page_size=1, page=2, sort_column='Contact Date تاريخ الاتصال')['rows']
    assert rows[0]['Hotline ID'] == '1'
    exported = pd.read_excel(io.BytesIO(store.explorer_export('legalhotlines', search='Baghdad')))
    assert len(exported) == 2


def test_hotline_optional_and_folder_import(payload, tmp_path):
    assert versioned_dataset_name('legalhotlines (2).csv') == ('legalhotlines', 2)
    for name, raw in payload.items():
        (tmp_path / (name + '.csv')).write_bytes(raw)
    store = LegalStore.from_folder(tmp_path)
    assert store.metadata()['availability']['legalhotlines']
    assert len(store.frames['legalhotlines']) == 5
    with pytest.raises(ValueError, match='not loaded'):
        dashboard(None)


def test_hotline_only_store_is_ready_and_exposes_only_hotline(payload):
    store = LegalStore.from_files({'legalhotlines': payload['legalhotlines']}, 'Hotline only')
    metadata = store.metadata()
    assert metadata['ready'] is True
    assert metadata['availability']['legalhotlines'] is True
    assert metadata['availability']['beneficiaries'] is False
    assert metadata['overview'] is None
    assert dashboard(store.frames['legalhotlines'])['total'] == 5


def test_hotline_upload_endpoint_and_refresh(payload, monkeypatch):
    from backend import main
    monkeypatch.setattr(main, 'legal_store', LegalStore.from_files(required_payload(), 'test'))
    with pytest.raises(HTTPException) as error:
        main.legal_hotline_dashboard(main.LegalQuery(dataset='legalhotlines'))
    assert error.value.status_code == 400
    def upload(data):
        return asyncio.run(main.legal_upload([UploadFile(filename=name+'.csv', file=io.BytesIO(raw)) for name, raw in data.items()]))
    response = upload(payload)
    revision = response['revision']
    result = main.legal_hotline_dashboard(main.LegalQuery(dataset='legalhotlines', filters={'Priority': ['Low']}))
    assert result['total'] == 2
    response = upload(required_payload())
    assert response['revision'] != revision
    assert not response['availability']['legalhotlines']


@pytest.fixture
def detention_chart_source():
    return pd.DataFrame({
        'Contact Date': ['2026-02-01'] * 7,
        'Is the beneficiary detained': ['Yes', 'نعم', 'Yes نعم', 'No', '', 'No لا', None],
        'Governorate of detention': ['Baghdad', '', 'Baghdad', 'Erbil', 'Duhok', '', None],
        'Detaining Authority': ['Police', 'Court', '', 'Other', 'Unknown', '', None],
        'Priority': ['High'] * 7,
    })


def test_detention_charts_count_only_normalized_yes_and_use_detained_denominator(detention_chart_source):
    source = detention_chart_source
    original = source.copy(deep=True)
    result = dashboard(source)
    charts = {chart['id']: chart['rows'] for chart in result['charts']}
    assert charts['Governorate of detention'] == [
        {'label': 'Baghdad', 'count': 2, 'percent': 2/3},
        {'label': 'Not recorded', 'count': 1, 'percent': 1/3},
    ]
    assert {row['label']: (row['count'], row['percent']) for row in charts['Detaining Authority']} == {
        'Police': (1, 1/3), 'Court': (1, 1/3), 'Not recorded': (1, 1/3),
    }
    assert result['total'] == result['filteredRows'] == 7
    assert charts['Priority'] == [{'label': 'High', 'count': 7, 'percent': 1}]
    assert result['trend'][0]['count'] == 7
    assert result['trend'][0]['detained'] == 3
    assert 'Erbil' in result['filterOptions']['Governorate of detention']
    assert 'Other' in result['filterOptions']['Detaining Authority']
    assert next(item for item in result['map']['items'] if item['label'] == 'Erbil')['notDetained'] == 1
    pd.testing.assert_frame_equal(source, original)


@pytest.mark.parametrize('filters,expected_count', [
    ({'Governorate of detention': ['Baghdad']}, 2),
    ({'Detaining Authority': ['Court']}, 1),
    ({'Contact Date': ['2026-02'], 'Governorate of detention': ['Not recorded']}, 1),
    ({'Is the beneficiary detained': ['No']}, 0),
    ({'Is the beneficiary detained': ['Not recorded']}, 0),
    ({'Contact Date': ['2026-03']}, 0),
])
def test_detention_charts_respect_filters_and_return_empty_rows(detention_chart_source, filters, expected_count):
    result = dashboard(detention_chart_source, filters)
    for chart in result['charts']:
        if chart['id'] not in {'Governorate of detention', 'Detaining Authority'}:
            continue
        assert sum(row['count'] for row in chart['rows']) == expected_count
        if expected_count:
            assert sum(row['percent'] for row in chart['rows']) == pytest.approx(1)
        else:
            assert chart['rows'] == []
